from __future__ import annotations

from io import BytesIO
from pathlib import Path
from typing import cast
from xml.etree import ElementTree
from zipfile import ZipFile

from openpyxl import load_workbook as _openpyxl_load_workbook
from openpyxl.workbook.workbook import Workbook

_STYLES_PART = "xl/styles.xml"


def _member_name(name: str) -> str:
    """Map Windows ZIP separators to the OOXML forward-slash part names.

    Some Excel builds store workbook parts as ``xl\\workbook.xml``.  The ZIP
    specification and openpyxl both address those parts with forward slashes,
    so the in-memory package uses the canonical names without rewriting source.
    """

    return name.replace("\\", "/")


def _local_name(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def _namespace(tag: str) -> str | None:
    if not tag.startswith("{") or "}" not in tag:
        return None
    return tag[1:].split("}", 1)[0]


def _normalize_empty_fills(styles_xml: bytes) -> bytes | None:
    """Treat an empty OOXML fill as the semantically equivalent no-fill pattern.

    Excel and WPS tolerate ``<fill/>`` entries, but openpyxl rejects the entire
    workbook because a fill has no pattern or gradient child.  The normalization
    is applied only to an in-memory package used by RackCore; the selected source
    workbook is never rewritten while it is being opened or scanned.
    """

    root = ElementTree.fromstring(styles_xml)
    fills = next((item for item in root if _local_name(item.tag) == "fills"), None)
    if fills is None:
        return None

    namespace = _namespace(fills.tag)
    pattern_tag = f"{{{namespace}}}patternFill" if namespace else "patternFill"
    changed = False
    for fill in fills:
        if _local_name(fill.tag) != "fill" or len(fill):
            continue
        ElementTree.SubElement(fill, pattern_tag, {"patternType": "none"})
        changed = True

    if not changed:
        return None
    if namespace:
        ElementTree.register_namespace("", namespace)
    return cast(
        bytes,
        ElementTree.tostring(root, encoding="utf-8", xml_declaration=True),
    )


def _normalized_package(path: Path) -> BytesIO | None:
    with ZipFile(path) as source:
        members = source.infolist()
        names = [_member_name(info.filename) for info in members]
        needs_separator_normalization = any(
            info.filename != normalized
            for info, normalized in zip(members, names, strict=True)
        )
        try:
            styles_name = next(
                info.filename
                for info, normalized in zip(members, names, strict=True)
                if normalized == _STYLES_PART
            )
        except StopIteration:
            styles_name = None
        normalized_styles = (
            _normalize_empty_fills(source.read(styles_name))
            if styles_name is not None
            else None
        )
        if normalized_styles is None and not needs_separator_normalization:
            return None

        seen: set[str] = set()
        output = BytesIO()
        with ZipFile(output, "w") as target:
            target.comment = source.comment
            for member, normalized_name in zip(members, names, strict=True):
                if normalized_name in seen:
                    continue
                seen.add(normalized_name)
                payload = (
                    normalized_styles
                    if normalized_styles is not None and normalized_name == _STYLES_PART
                    else source.read(member.filename)
                )
                member.filename = normalized_name
                target.writestr(member, payload)
        output.seek(0)
        return output


def load_xlsx_workbook(
    path: Path,
    *,
    read_only: bool = False,
    data_only: bool = False,
) -> Workbook:
    """Load an editable XLSX, tolerating Excel/WPS empty-fill style records."""

    normalized = _normalized_package(path)
    source: Path | BytesIO = normalized if normalized is not None else path
    return _openpyxl_load_workbook(
        source,
        read_only=read_only,
        data_only=data_only,
    )
