from __future__ import annotations

import hashlib
import os
import tempfile
from collections import defaultdict
from dataclasses import dataclass, replace
from pathlib import Path
from typing import Literal, cast
from xml.etree import ElementTree
from zipfile import ZipFile

from openpyxl import Workbook
from openpyxl.cell.cell import Cell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet

from racktool.core.identity import normalize_path, sha256_file
from racktool.core.ooxml import load_xlsx_workbook
from racktool.core.position_table import position_metadata_digest, write_position_metadata
from racktool.core.project import project_error_conflicts, rescan_workbook
from racktool.core.source_export import append_position_sheet
from racktool.core.workbook import scan_workbook
from racktool.models.domain import Device, Placement, Rack
from racktool.models.project import RackProject
from racktool.models.workbook import SheetInfo

_RACK_SHEET = "机柜图"
_POSITION_SHEET = "设备位置表"
_POSITION_HEADERS = (
    "序号",
    "设备",
    "机柜",
    "U 位",
    "高度（U）",
    "机房",
    "状态",
    "设备 ID",
    "机柜 ID",
)
_RACK_BLOCK_COLUMNS = 3
_RACK_GUTTER_COLUMNS = 1
_FONT_NAME = "Arial"
_DEVICE_FILLS = (
    "D9EAF7",
    "E2F0D9",
    "FFF2CC",
    "FCE4D6",
    "E4DFEC",
    "DDEBF7",
    "F4CCCC",
)

_THIN_GRAY = Side(style="thin", color="AAB4C3")
_MEDIUM_BLUE = Side(style="medium", color="5B7699")
_GRID_BORDER = Border(
    left=_THIN_GRAY,
    right=_THIN_GRAY,
    top=_THIN_GRAY,
    bottom=_THIN_GRAY,
)
@dataclass(frozen=True, slots=True)
class _PositionRow:
    device_id: str
    display_text: str
    rack_id: str
    rack_name: str
    start_u: int | None
    end_u: int | None
    height_u: int | None
    sheet_name: str
    source_range: str
    status: str
    source_index: int
    rack_index: int


@dataclass(frozen=True, slots=True)
class _ExportManifest:
    rack_titles: tuple[tuple[str, str, Rack], ...]
    position_rows: tuple[_PositionRow, ...]
    sheet_names: tuple[str, ...]
    position_sheet: str
    source_sheets: tuple[SheetInfo, ...] = ()
    source_metadata: tuple[str, ...] = ()
    source_payload: bytes = b""
    position_metadata_digest: str = ""


@dataclass(frozen=True, slots=True)
class _RackBand:
    source_row: int
    racks: tuple[Rack, ...]


def _literal_text(cell: Cell, value: str) -> None:
    if len(value) > 32_767:
        raise ValueError("导出失败：有单元格文字超过 Excel 可保存的长度上限。")
    cell.value = value
    if value.startswith("="):
        cell.data_type = "s"


def _cell(sheet: Worksheet, row: int, column: int) -> Cell:
    return cast(Cell, sheet.cell(row, column))


def _normalise_display_text(value: str) -> str:
    return value.replace("\r\n", "\n").replace("\r", "\n")


def _device_fill(device: Device) -> PatternFill:
    seed = device.style_signature or device.device_id
    index = int(hashlib.sha256(seed.encode("utf-8")).hexdigest()[:8], 16)
    return PatternFill("solid", fgColor=_DEVICE_FILLS[index % len(_DEVICE_FILLS)])


def _normalise_destination(path: Path, project: RackProject) -> Path:
    candidate = path.expanduser()
    if not candidate.suffix:
        candidate = candidate.with_suffix(".xlsx")
    if candidate.suffix.casefold() != ".xlsx":
        raise ValueError("导出文件必须使用 .xlsx 扩展名。")
    destination = candidate.resolve()
    if project.source_workbook is not None:
        source = normalize_path(Path(project.source_workbook))
        if destination == source:
            raise ValueError("导出文件不能覆盖当前源工作簿，请选择另一个文件名。")
    return destination


def _ordered_racks(project: RackProject) -> list[Rack]:
    source_sheets: list[str] = []
    if project.source_workbook is not None:
        source = normalize_path(Path(project.source_workbook))
        if source.is_file():
            try:
                workbook = load_xlsx_workbook(source, read_only=True, data_only=False)
                try:
                    source_sheets = list(workbook.sheetnames)
                finally:
                    workbook.close()
            except Exception:  # noqa: BLE001 - source order is an optional export hint
                source_sheets = []
    if not source_sheets:
        source_sheets = sorted(
            {rack.source_sheet or "" for rack in project.racks},
            key=str.casefold,
        )
    sheet_order = {name: index for index, name in enumerate(source_sheets)}
    source_order = {rack.rack_id: index for index, rack in enumerate(project.racks)}
    return sorted(
        project.racks,
        key=lambda rack: (
            sheet_order.get(rack.source_sheet or "", 10**9),
            rack.bounds.min_row if rack.bounds is not None else 10**9,
            rack.bounds.min_col if rack.bounds is not None else 10**9,
            source_order[rack.rack_id],
        ),
    )


def _position_rows(project: RackProject) -> list[_PositionRow]:
    racks = {rack.rack_id: rack for rack in project.racks}
    rack_order = {rack.rack_id: index for index, rack in enumerate(_ordered_racks(project))}
    placements = {placement.device_id: placement for placement in project.placements}
    mappings = {
        mapping.device_id: mapping
        for mapping in project.mappings
        if mapping.mapping_kind == "device" and mapping.device_id is not None
    }
    rows: list[_PositionRow] = []
    for source_index, device in enumerate(project.devices):
        placement = placements.get(device.device_id)
        mapping = mappings.get(device.device_id)
        rack = racks.get(placement.rack_id) if placement is not None else None
        if placement is None:
            status = "未放置"
        elif placement.status == "active" and rack is not None and rack.status == "active":
            status = "在位"
        elif placement.status == "missing" or (rack is not None and rack.status == "missing"):
            status = "源文件中已缺失"
        else:
            status = placement.status
        start_u = (
            min(placement.start_u, placement.end_u) if placement is not None else None
        )
        end_u = max(placement.start_u, placement.end_u) if placement is not None else None
        rows.append(
            _PositionRow(
                device_id=device.device_id,
                display_text=_normalise_display_text(device.display_text),
                rack_id=placement.rack_id if placement is not None else "",
                rack_name=rack.rack_name if rack is not None else "",
                start_u=start_u,
                end_u=end_u,
                height_u=placement.height_u if placement is not None else None,
                sheet_name=(
                    mapping.sheet_name
                    if mapping is not None
                    else (rack.source_sheet or "" if rack is not None else "")
                ),
                source_range=mapping.source_range if mapping is not None else "",
                status=status,
                source_index=source_index,
                rack_index=rack_order.get(placement.rack_id, 10**9)
                if placement is not None
                else 10**9,
            )
        )
    rows.sort(
        key=lambda row: (
            row.rack_index,
            (row.start_u if row.start_u is not None else 10**9)
            if racks.get(row.rack_id) is not None and racks[row.rack_id].direction == "ascending"
            else -(row.end_u if row.end_u is not None else -1),
            row.source_index,
        )
    )
    return rows


def _u_label(row: _PositionRow) -> str:
    if row.start_u is None or row.end_u is None:
        return ""
    return f"{row.start_u}U" if row.start_u == row.end_u else f"{row.start_u}-{row.end_u}U"


def _unit_row(rack: Rack, body_start: int, unit: int) -> int:
    return body_start + (unit - 1 if rack.direction == "ascending" else rack.height_u - unit)


def _unique_sheet_name(workbook: Workbook, desired: str) -> str:
    existing = {name.casefold() for name in workbook.sheetnames}
    name = desired[:31]
    suffix = 1
    while name.casefold() in existing:
        suffix += 1
        tail = f" ({suffix})"
        name = desired[:31 - len(tail)] + tail
    return name


def _active_placements(project: RackProject) -> dict[str, list[tuple[Placement, Device]]]:
    devices = {device.device_id: device for device in project.devices}
    by_rack: dict[str, list[tuple[Placement, Device]]] = defaultdict(list)
    for placement in project.placements:
        device = devices.get(placement.device_id)
        if placement.status == "active" and device is not None:
            by_rack[placement.rack_id].append((placement, device))
    for items in by_rack.values():
        items.sort(
            key=lambda item: (
                -max(item[0].start_u, item[0].end_u),
                item[1].device_id,
            )
        )
    return by_rack


def _rack_bands(racks: list[Rack]) -> list[_RackBand]:
    source_order = {rack.rack_id: index for index, rack in enumerate(racks)}
    grouped: dict[int, list[Rack]] = {}
    for rack in racks:
        source_row = (
            rack.bounds.min_row
            if rack.bounds is not None
            else 10**9 + source_order[rack.rack_id]
        )
        grouped.setdefault(source_row, []).append(rack)
    return [
        _RackBand(
            source_row,
            tuple(
                sorted(
                    band,
                    key=lambda rack: (
                        rack.bounds.min_col if rack.bounds is not None else 10**9,
                        source_order[rack.rack_id],
                    ),
                )
            ),
        )
        for source_row, band in sorted(grouped.items())
    ]


def _diagram_column_count(rack_count: int) -> int:
    if rack_count <= 0:
        return 12
    return max(
        12,
        rack_count * (_RACK_BLOCK_COLUMNS + _RACK_GUTTER_COLUMNS)
        - _RACK_GUTTER_COLUMNS,
    )


def _configure_rack_columns(sheet: Worksheet, rack_count: int) -> None:
    for block_index in range(rack_count):
        start_column = block_index * (_RACK_BLOCK_COLUMNS + _RACK_GUTTER_COLUMNS) + 1
        sheet.column_dimensions[get_column_letter(start_column)].width = 5
        sheet.column_dimensions[get_column_letter(start_column + 1)].width = 30
        sheet.column_dimensions[get_column_letter(start_column + 2)].width = 5
        if block_index < rack_count - 1:
            sheet.column_dimensions[get_column_letter(start_column + 3)].width = 2


def _set_outer_border(
    sheet: Worksheet,
    *,
    min_row: int,
    max_row: int,
    min_column: int,
    max_column: int,
    side: Side,
) -> None:
    for row in range(min_row, max_row + 1):
        for column in range(min_column, max_column + 1):
            cell = sheet.cell(row, column)
            cell.border = Border(
                left=side if column == min_column else Side(),
                right=side if column == max_column else Side(),
                top=side if row == min_row else Side(),
                bottom=side if row == max_row else Side(),
            )


def _style_diagram_header(sheet: Worksheet, max_column: int) -> None:
    _literal_text(sheet["A1"], _RACK_SHEET)
    sheet["A1"].font = Font(name=_FONT_NAME, size=16, bold=True, color="172033")
    _literal_text(
        sheet["A2"],
        "每排机柜按源工作表中的行列位置排列；设备文字保留原换行；颜色仅用于区分设备。",
    )
    sheet["A2"].font = Font(name=_FONT_NAME, size=10, italic=True, color="667085")
    sheet.merge_cells(
        start_row=2,
        start_column=1,
        end_row=2,
        end_column=min(max_column, 20),
    )
    sheet.row_dimensions[1].height = 24
    sheet.row_dimensions[2].height = 20


def _style_rack_base(
    sheet: Worksheet,
    rack: Rack,
    *,
    start_column: int,
    title_row: int,
    max_height: int,
) -> str:
    end_column = start_column + _RACK_BLOCK_COLUMNS - 1
    sheet.merge_cells(
        start_row=title_row,
        start_column=start_column,
        end_row=title_row,
        end_column=end_column,
    )
    title = _cell(sheet, title_row, start_column)
    _literal_text(title, rack.rack_name)
    title.font = Font(name=_FONT_NAME, size=11, bold=True, color="172033")
    title.fill = PatternFill("solid", fgColor="F6C88F")
    title.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    _set_outer_border(
        sheet,
        min_row=title_row,
        max_row=title_row,
        min_column=start_column,
        max_column=end_column,
        side=_MEDIUM_BLUE,
    )
    sheet.row_dimensions[title_row].height = max(
        float(sheet.row_dimensions[title_row].height or 0), 24.0
    )

    body_start = title_row + 1
    for offset in range(max_height):
        row = body_start + offset
        for column in range(start_column, end_column + 1):
            cell = _cell(sheet, row, column)
            cell.border = _GRID_BORDER
            cell.font = Font(name=_FONT_NAME, size=8, color="172033")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.row_dimensions[row].height = max(
            float(sheet.row_dimensions[row].height or 0), 22.0
        )

    for unit in range(rack.height_u, 0, -1):
        row = _unit_row(rack, body_start, unit)
        for column in (start_column, end_column):
            cell = _cell(sheet, row, column)
            cell.value = unit
            cell.fill = PatternFill("solid", fgColor="FCE0B8")
            cell.font = Font(name=_FONT_NAME, size=8, color="172033")
            cell.alignment = Alignment(horizontal="center", vertical="center")
    if rack.height_u < max_height:
        for row in range(body_start + rack.height_u, body_start + max_height):
            for column in range(start_column, end_column + 1):
                _cell(sheet, row, column).fill = PatternFill("solid", fgColor="EEF1F5")
    return title.coordinate


def _render_rack_devices(
    sheet: Worksheet,
    rack: Rack,
    placements: list[tuple[Placement, Device]],
    *,
    start_column: int,
    body_start: int,
) -> None:
    for placement, device in placements:
        low = min(placement.start_u, placement.end_u)
        high = max(placement.start_u, placement.end_u)
        top_row, bottom_row = sorted((_unit_row(rack, body_start, low), _unit_row(rack, body_start, high)))
        left_column = start_column + 1
        right_column = left_column
        if top_row != bottom_row:
            sheet.merge_cells(
                start_row=top_row,
                start_column=left_column,
                end_row=bottom_row,
                end_column=right_column,
            )
        anchor = _cell(sheet, top_row, left_column)
        text = _normalise_display_text(device.display_text)
        _literal_text(anchor, text)
        anchor.fill = _device_fill(device)
        _set_outer_border(
            sheet,
            min_row=top_row,
            max_row=bottom_row,
            min_column=left_column,
            max_column=right_column,
            side=_MEDIUM_BLUE,
        )
        anchor.font = Font(name=_FONT_NAME, size=8, color="172033")
        anchor.alignment = Alignment(
            horizontal="left",
            vertical="center",
            wrap_text=True,
            shrink_to_fit=True,
        )
        if top_row == bottom_row:
            line_count = max(text.count("\n") + 1, 1)
            sheet.row_dimensions[top_row].height = max(
                float(sheet.row_dimensions[top_row].height or 0),
                min(72.0, line_count * 10.0 + 4.0),
            )


def _write_rack_sheet(
    workbook: Workbook,
    project: RackProject,
) -> tuple[tuple[str, str, Rack], ...]:
    active_racks = [rack for rack in _ordered_racks(project) if rack.status == "active"]
    placements = _active_placements(project)
    groups: dict[str, list[Rack]] = {}
    for rack in active_racks:
        groups.setdefault(rack.source_sheet or "未命名工作表", []).append(rack)
    bands_by_sheet = {
        sheet_name: _rack_bands(racks) for sheet_name, racks in groups.items()
    }
    if not groups:
        bands_by_sheet[_RACK_SHEET] = []
    title_cells: list[tuple[str, str, Rack]] = []
    for sheet_name, bands in bands_by_sheet.items():
        # Drafts have virtual source names, while imported sheets represent rooms.
        desired = sheet_name if project.source_workbook is not None else _RACK_SHEET
        sheet = workbook.create_sheet(_unique_sheet_name(workbook, desired))
        sheet.sheet_view.showGridLines = False
        sheet.sheet_properties.tabColor = "4472C4"
        max_racks_in_band = max((len(band.racks) for band in bands), default=0)
        diagram_max_column = _diagram_column_count(max_racks_in_band)
        _style_diagram_header(sheet, diagram_max_column)
        _configure_rack_columns(sheet, max_racks_in_band)
        current_row = 4
        sheet.merge_cells(
            start_row=current_row,
            start_column=1,
            end_row=current_row,
            end_column=diagram_max_column,
        )
        group_cell = _cell(sheet, current_row, 1)
        _literal_text(group_cell, f"机房：{sheet_name}")
        group_cell.font = Font(name=_FONT_NAME, size=11, bold=True, color="344054")
        group_cell.fill = PatternFill("solid", fgColor="E9EEF5")
        group_cell.alignment = Alignment(horizontal="left", vertical="center")
        group_cell.border = Border(bottom=Side(style="thin", color="AAB4C3"))
        sheet.row_dimensions[current_row].height = 22
        current_row += 2

        for band in bands:
            batch = list(band.racks)
            max_height = max(rack.height_u for rack in batch)
            title_row = current_row
            for block_index, rack in enumerate(batch):
                start_column = (
                    block_index * (_RACK_BLOCK_COLUMNS + _RACK_GUTTER_COLUMNS) + 1
                )
                coordinate = _style_rack_base(
                    sheet,
                    rack,
                    start_column=start_column,
                    title_row=title_row,
                    max_height=max_height,
                )
                title_cells.append((sheet.title, coordinate, rack))
                _render_rack_devices(
                    sheet,
                    rack,
                    placements.get(rack.rack_id, []),
                    start_column=start_column,
                    body_start=title_row + 1,
                )
            current_row += max_height + 3

        sheet.freeze_panes = "A4"
        sheet.print_area = f"A1:{get_column_letter(diagram_max_column)}{max(current_row - 1, 4)}"
        sheet.print_title_rows = "1:2"
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    return tuple(title_cells)


def _write_position_sheet(
    workbook: Workbook,
    project: RackProject,
    rows: list[_PositionRow],
) -> str:
    sheet = workbook.create_sheet(_unique_sheet_name(workbook, _POSITION_SHEET))
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.tabColor = "70AD47"
    header_row = 1
    for column, header in enumerate(_POSITION_HEADERS, start=1):
        cell = _cell(sheet, header_row, column)
        _literal_text(cell, header)
        cell.fill = PatternFill("solid", fgColor="344054")
        cell.font = Font(name=_FONT_NAME, size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(right=Side(style="thin", color="FFFFFF"))
    sheet.row_dimensions[header_row].height = 24

    for output_index, row in enumerate(rows, start=1):
        excel_row = header_row + output_index
        values: tuple[str | int | None, ...] = (
            output_index,
            row.display_text,
            row.rack_name,
            _u_label(row),
            row.height_u,
            row.sheet_name,
            row.status,
            row.device_id,
            row.rack_id,
        )
        for column, value in enumerate(values, start=1):
            cell = _cell(sheet, excel_row, column)
            if isinstance(value, str):
                _literal_text(cell, value)
            else:
                cell.value = value
            cell.font = Font(name=_FONT_NAME, size=10, color="172033")
            cell.alignment = Alignment(
                horizontal="left" if column in {2, 3, 6, 7} else "center",
                vertical="top" if column == 2 else "center",
                wrap_text=True,
            )
            cell.border = Border(bottom=Side(style="thin", color="E4E8EF"))
        line_count = max(row.display_text.count("\n") + 1, 1)
        sheet.row_dimensions[excel_row].height = min(90.0, max(20.0, line_count * 15.0))

    last_row = header_row + max(len(rows), 1)
    if rows:
        table_name = "RackToolDevicePositions"
        index = 1
        existing_names = {name.casefold() for name in workbook.defined_names}
        existing_names.update(name.casefold() for worksheet in workbook.worksheets for name in worksheet.tables)
        while table_name.casefold() in existing_names:
            index += 1
            table_name = f"RackToolDevicePositions{index}"
        table = Table(displayName=table_name, ref=f"A1:I{last_row}")
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        sheet.add_table(table)
    else:
        sheet.auto_filter.ref = "A1:I1"

    widths = (7, 44, 18, 14, 11, 22, 18, 25, 25)
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.column_dimensions["H"].hidden = True
    sheet.column_dimensions["I"].hidden = True
    sheet.freeze_panes = "A2"
    sheet.print_title_rows = "1:1"
    sheet.print_area = f"A1:G{last_row}"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    write_position_metadata(sheet, project)
    return cast(str, sheet.title)


def _source_metadata(workbook: Workbook) -> tuple[str, ...]:
    details = [str(workbook.defined_names)]
    for sheet in workbook.worksheets:
        details.append(repr((sheet.title, sheet.freeze_panes, str(sheet.print_area),
                             sheet.print_title_rows, sheet.print_title_cols)))
        for item in (sheet.sheet_format, sheet.sheet_properties, sheet.page_setup,
                     sheet.page_margins, sheet.data_validations):
            details.append(ElementTree.tostring(item.to_tree(), encoding="unicode"))
        for row in sheet.iter_rows():
            for cell in row:
                if cell.comment is not None:
                    details.append(repr((sheet.title, cell.coordinate, cell.comment.text, cell.comment.author)))
                if cell.hyperlink is not None:
                    link = cell.hyperlink
                    details.append(repr((sheet.title, cell.coordinate, link.target, link.location, link.tooltip)))
    return tuple(details)


def _source_workbook(project: RackProject) -> tuple[Workbook, tuple[SheetInfo, ...], tuple[str, ...], bytes]:
    if project.source_workbook is None or project.workbook_fingerprint is None:
        raise ValueError("保留原版面导出需要已绑定 Excel 的项目；新建项目请使用标准导出。")
    source = normalize_path(Path(project.source_workbook))
    payload = source.read_bytes()
    if hashlib.sha256(payload).hexdigest() != project.workbook_fingerprint:
        raise ValueError("源文件已变化，请重新扫描后再导出原版面。")
    with tempfile.TemporaryDirectory(prefix="racktool-source-export-") as directory:
        snapshot = Path(directory) / "source.xlsx"
        snapshot.write_bytes(payload)
        with ZipFile(snapshot) as archive:
            for member in archive.namelist():
                name = member.replace("\\", "/").lower()
                if any(part in name for part in (
                    "vbaproject", "xl/activex/", "xl/ctrlprops/",
                )):
                    raise ValueError("原版面含未支持的 Excel 对象，不能保证保留；请使用标准导出。")
        rescanned = rescan_workbook(snapshot, project)
        old_rows = sorted((replace(row, source_index=0, rack_index=0) for row in _position_rows(project)),
                          key=lambda row: row.device_id)
        new_rows = sorted((replace(row, source_index=0, rack_index=0) for row in _position_rows(rescanned.project)),
                          key=lambda row: row.device_id)
        if (not rescanned.accepted or new_rows != old_rows
                or sorted(rescanned.project.racks, key=lambda rack: rack.rack_id)
                != sorted(project.racks, key=lambda rack: rack.rack_id)):
            raise ValueError("项目与源文件不一致，请先同步或重新扫描，再导出原版面。")
        sheets = tuple(scan_workbook(snapshot).sheets)
        workbook = load_xlsx_workbook(snapshot)
    return workbook, sheets, _source_metadata(workbook), payload


def _build_workbook(
    project: RackProject, layout: Literal["standard", "source"] = "standard",
) -> tuple[Workbook, _ExportManifest]:
    source_sheets: tuple[SheetInfo, ...] = ()
    metadata: tuple[str, ...] = ()
    payload = b""
    if layout == "source":
        workbook, source_sheets, metadata, payload = _source_workbook(project)
    else:
        workbook = Workbook()
        if workbook.active is not None:
            workbook.remove(workbook.active)
    try:
        if layout == "standard":
            workbook.properties.creator = "RackTool"
            workbook.properties.title = "RackTool 机柜图与设备位置表"
            workbook.properties.subject = "RackTool 标准导出"
        rows = _position_rows(project)
        rack_titles = _write_rack_sheet(workbook, project) if layout == "standard" else ()
        position_sheet = _write_position_sheet(workbook, project, rows)
        return workbook, _ExportManifest(rack_titles, tuple(rows), tuple(workbook.sheetnames),
                                         position_sheet, source_sheets, metadata, payload,
                                         position_metadata_digest(workbook[position_sheet]))
    except Exception:
        workbook.close()
        raise


def _validate_export(path: Path, manifest: _ExportManifest) -> None:
    workbook = load_xlsx_workbook(path, read_only=False, data_only=False)
    try:
        if tuple(workbook.sheetnames) != manifest.sheet_names:
            raise ValueError("导出验证失败：工作表结构不完整。")
        position_sheet = workbook[manifest.position_sheet]
        if position_metadata_digest(position_sheet) != manifest.position_metadata_digest:
            raise ValueError("导出验证失败：设备位置表回导元数据发生变化。")
        headers = tuple(position_sheet.cell(1, column).value for column in range(1, 10))
        if headers != _POSITION_HEADERS:
            raise ValueError("导出验证失败：设备位置表标题不完整。")
        for index, row in enumerate(manifest.position_rows, 2):
            expected = (index - 1, row.display_text, row.rack_name or None, _u_label(row) or None, row.height_u,
                        row.sheet_name or None, row.status, row.device_id, row.rack_id or None)
            actual = tuple(position_sheet.cell(index, column).value for column in range(1, 10))
            if actual != expected:
                raise ValueError("导出验证失败：设备位置表内容在保存后发生变化。")
        for sheet_name, coordinate, rack in manifest.rack_titles:
            diagram_sheet = workbook[sheet_name]
            title = diagram_sheet[coordinate]
            if title.value != rack.rack_name:
                raise ValueError("导出验证失败：机柜名称在保存后发生变化。")
            for unit in range(1, rack.height_u + 1):
                row_number = _unit_row(rack, title.row + 1, unit)
                if any(diagram_sheet.cell(row_number, column).value != unit
                       for column in (title.column, title.column + _RACK_BLOCK_COLUMNS - 1)):
                    raise ValueError("导出验证失败：机柜 U 轴方向或标尺发生变化。")
            for row in manifest.position_rows:
                if row.rack_id != rack.rack_id or row.status != "在位":
                    continue
                assert row.start_u is not None and row.end_u is not None
                top, bottom = sorted((_unit_row(rack, title.row + 1, row.start_u),
                                      _unit_row(rack, title.row + 1, row.end_u)))
                anchor = diagram_sheet.cell(top, title.column + 1)
                if anchor.value != row.display_text:
                    raise ValueError("导出验证失败：机柜内设备文字或 U 位发生变化。")
                if bottom > top and not any(
                    merge.min_row == top and merge.max_row == bottom
                    and merge.min_col == merge.max_col == title.column + 1
                    for merge in diagram_sheet.merged_cells.ranges
                ):
                    raise ValueError("导出验证失败：设备 U 位合并范围发生变化。")
        for sheet in workbook.worksheets:
            if manifest.source_sheets and sheet.title != manifest.position_sheet:
                continue  # Original formulas are source content, not generated device text.
            if any(cell.data_type == "f" for row in sheet.iter_rows() for cell in row):
                raise ValueError("导出验证失败：导出表中出现了意外公式。")
        if manifest.source_sheets:
            changed_parts = {"xl/workbook.xml", "xl/_rels/workbook.xml.rels", "xl/styles.xml", "[Content_Types].xml"}
            from io import BytesIO

            with ZipFile(BytesIO(manifest.source_payload)) as original, ZipFile(path) as exported:
                for name in original.namelist():
                    canonical = name.replace("\\", "/")
                    if canonical not in changed_parts and original.read(name) != exported.read(canonical):
                        raise ValueError("导出验证失败：原工作簿部件发生变化。")
            actual_sheets = tuple(scan_workbook(path).sheets[:len(manifest.source_sheets)])
            if actual_sheets != manifest.source_sheets:
                raise ValueError("导出验证失败：原工作表的内容、样式、方向或尺寸发生变化。")
            workbook.remove(position_sheet)
            if _source_metadata(workbook) != manifest.source_metadata:
                raise ValueError("导出验证失败：原工作表的备注、链接或打印设置发生变化。")
    finally:
        workbook.close()


def export_project_workbook(
    project: RackProject,
    path: Path,
    *,
    overwrite: bool = False,
    layout: Literal["standard", "source"] = "standard",
) -> Path:
    """Export room sheets and compact positions, optionally preserving source layout."""
    if layout not in {"standard", "source"}:
        raise ValueError("导出版面必须是 standard 或 source。")
    errors = project_error_conflicts(project)
    if errors:
        raise ValueError("项目数据仍有未解决错误，请先核对分析结果并解决异常后再导出。")
    destination = _normalise_destination(path, project)
    if destination.exists():
        if destination.is_dir():
            raise IsADirectoryError(destination)
        if not overwrite:
            raise FileExistsError("导出文件已存在；请换一个文件名，或明确允许覆盖。")
    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook, manifest = _build_workbook(project, layout)
    with tempfile.NamedTemporaryFile(
        mode="wb",
        prefix=f".{destination.stem}.tmp-",
        suffix=".xlsx",
        dir=destination.parent,
        delete=False,
    ) as temp_handle:
        temp_path = Path(temp_handle.name)
    try:
        workbook.save(temp_path)
        workbook.close()
        if manifest.source_payload:
            append_position_sheet(manifest.source_payload, temp_path, manifest.position_sheet)
        _validate_export(temp_path, manifest)
        if layout == "source" and sha256_file(Path(project.source_workbook or "")) != project.workbook_fingerprint:
            raise ValueError("源文件在导出期间变化，请重新扫描后再试。")
        os.replace(temp_path, destination)
    except Exception:
        workbook.close()
        if temp_path.exists():
            temp_path.unlink()
        raise
    return destination
