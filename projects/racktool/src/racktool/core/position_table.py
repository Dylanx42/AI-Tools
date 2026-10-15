"""Validated round-trip of the exported position table through RackCore only."""

from __future__ import annotations

import base64
import hashlib
import json
import re
import zlib
from dataclasses import dataclass, replace
from pathlib import Path
from typing import Any, cast

from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.worksheet import Worksheet

from racktool.core.backup import create_backup, discard_backup
from racktool.core.identity import normalize_path, sha256_file
from racktool.core.ooxml import load_xlsx_workbook
from racktool.core.project import project_error_conflicts
from racktool.core.text import validate_excel_text
from racktool.models.project import IdentityConflict, RackProject
from racktool.persistence import load_project

_PREFIX = "_RackToolPositions_v1_"
_MAX_METADATA = 8 * 1024 * 1024
_HEADERS = ("序号", "设备", "机柜", "U 位", "高度（U）", "机房", "状态", "设备 ID", "机柜 ID")
_U_LABEL = re.compile(r"([1-9][0-9]*)(?:-([1-9][0-9]*))?U")


def _json(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))


def project_revision(project: RackProject) -> str:
    """Hash semantic project state independently of SQLite retrieval/list order."""
    payload = project.to_dict()
    for collection, key in (("racks", "rack_id"), ("devices", "device_id"),
                            ("placements", "placement_id"), ("mappings", "mapping_id")):
        payload[collection] = sorted(payload[collection], key=lambda item: item[key])
    return hashlib.sha256(_json(payload).encode()).hexdigest()


def _baseline(project: RackProject) -> dict[str, list[Any]]:
    # One shared representation prevents a second export/schema implementation.
    from racktool.core.export import _position_rows, _u_label

    return {row.device_id: [row.display_text, row.rack_name, _u_label(row), row.height_u,
                           row.sheet_name, row.status, row.device_id, row.rack_id]
            for row in _position_rows(project)}


def write_position_metadata(sheet: Worksheet, project: RackProject) -> None:
    """Hidden, sheet-scoped constants survive row sorting without adding visible cells."""
    payload = {"schema_version": 1, "project_id": project.project_id,
               "project_revision": project_revision(project), "baseline": _baseline(project)}
    raw = _json(payload).encode()
    if len(raw) > _MAX_METADATA:
        raise ValueError("设备位置表回导元数据超过支持的大小。")
    encoded = base64.b64encode(zlib.compress(raw)).decode("ascii")
    for index, offset in enumerate(range(0, len(encoded), 200)):
        sheet.defined_names.add(DefinedName(
            f"{_PREFIX}{index:05d}", attr_text='"' + encoded[offset:offset + 200] + '"',
            hidden=True,
        ))


def _metadata(sheet: Worksheet) -> dict[str, Any]:
    names = sorted(name for name in sheet.defined_names if name.startswith(_PREFIX))
    if not names or len(names) > _MAX_METADATA // 200:
        raise ValueError("位置表缺失受支持的导出元数据；请从当前项目重新导出。")
    if names != [f"{_PREFIX}{index:05d}" for index in range(len(names))]:
        raise ValueError("位置表导出元数据分块缺失或重复。")
    chunks = []
    for name in names:
        definition = sheet.defined_names[name]
        value = definition.attr_text
        if not definition.hidden or not isinstance(value, str) or not re.fullmatch(
            r'"[A-Za-z0-9+/=]{1,200}"', value,
        ):
            raise ValueError("位置表导出元数据格式无效。")
        chunks.append(value[1:-1])
    try:
        compressed = base64.b64decode("".join(chunks), validate=True)
        decompressor = zlib.decompressobj()
        raw = decompressor.decompress(compressed, _MAX_METADATA + 1)
        if len(raw) > _MAX_METADATA or not decompressor.eof or decompressor.unused_data:
            raise ValueError("元数据大小或压缩格式无效")
        payload = json.loads(raw)
    except (ValueError, zlib.error, UnicodeError) as error:
        raise ValueError("位置表导出元数据损坏。") from error
    if not isinstance(payload, dict) or payload.get("schema_version") != 1:
        raise ValueError("位置表导出元数据版本不受支持。")
    return payload


def position_metadata_digest(sheet: Worksheet) -> str:
    """Verify the serialized export retained the entire metadata payload."""
    return hashlib.sha256(_json(_metadata(sheet)).encode()).hexdigest()


@dataclass(frozen=True, slots=True)
class PositionTableChange:
    device_id: str
    rack_id: str
    start_u: int
    end_u: int
    display_text: str
    before: list[Any]
    after: list[Any]

    def to_dict(self) -> dict[str, Any]:
        return {"device_id": self.device_id, "rack_id": self.rack_id,
                "start_u": self.start_u, "end_u": self.end_u,
                "display_text": self.display_text, "before": self.before, "after": self.after}


@dataclass(frozen=True, slots=True)
class PositionTablePlan:
    project_id: str
    project_revision: str
    workbook_path: str
    workbook_sha256: str
    sheet_name: str | None = None
    changes: tuple[PositionTableChange, ...] = ()
    conflicts: tuple[IdentityConflict, ...] = ()
    preview_token: str | None = None

    @property
    def accepted(self) -> bool:
        return not self.conflicts

    def to_dict(self) -> dict[str, Any]:
        return {"accepted": self.accepted, "project_id": self.project_id,
                "project_revision": self.project_revision, "workbook_path": self.workbook_path,
                "workbook_sha256": self.workbook_sha256, "sheet_name": self.sheet_name,
                "changes": [change.to_dict() for change in self.changes],
                "conflicts": [conflict.to_dict() for conflict in self.conflicts],
                "preview_token": self.preview_token}


def _requests(changes: tuple[PositionTableChange, ...]) -> list[Any]:
    from racktool.core.sync import EditRequest

    return [EditRequest(change.device_id, change.rack_id, change.start_u,
                        change.end_u, change.display_text) for change in changes]


def _draft_candidate(project: RackProject, changes: tuple[PositionTableChange, ...]) -> RackProject:
    from racktool.core.editing import _device_range, _validate_draft

    _validate_draft(project)
    by_id = {change.device_id: change for change in changes}
    racks = {rack.rack_id: rack for rack in project.racks}
    devices = [replace(device, display_text=by_id[device.device_id].display_text)
               if device.device_id in by_id else device for device in project.devices]
    placements = [replace(placement, rack_id=by_id[placement.device_id].rack_id,
                          start_u=by_id[placement.device_id].start_u,
                          end_u=by_id[placement.device_id].end_u)
                  if placement.device_id in by_id else placement for placement in project.placements]
    mappings = []
    for mapping in project.mappings:
        change = by_id.get(mapping.device_id or "")
        if change is None:
            mappings.append(mapping)
        else:
            rack = racks[change.rack_id]
            mappings.append(replace(mapping, rack_id=rack.rack_id, sheet_name=rack.source_sheet or "",
                                    source_range=_device_range(rack, change.start_u, change.end_u)))
    candidate = replace(project, devices=devices, placements=placements, mappings=mappings)
    _validate_draft(candidate)
    return candidate


def _text(value: Any, label: str) -> str:
    if not isinstance(value, str):
        raise TypeError(f"{label}必须是文字。")
    validate_excel_text(value)
    return value.replace("\r\n", "\n").replace("\r", "\n")


def _changes(sheet: Worksheet, project: RackProject, baseline: dict[str, list[Any]]) -> tuple[PositionTableChange, ...]:
    if tuple(sheet.cell(1, col).value for col in range(1, 10)) != _HEADERS:
        raise ValueError("位置表标题被修改或不是当前导出格式。")
    seen: set[str] = set()
    changes = []
    for row in sheet.iter_rows(min_row=2, max_col=9):
        if all(cell.value is None for cell in row):
            continue
        if any(cell.data_type == "f" for cell in row):
            raise ValueError(f"位置表第 {row[0].row} 行含公式；回导仅接受明确值。")
        values = [cell.value for cell in row]
        device_id = _text(values[7], "设备 ID")
        if device_id in seen:
            raise ValueError("位置表有重复设备 ID。")
        if device_id not in baseline:
            raise ValueError("位置表有未知设备 ID；添加设备请使用添加入口。")
        seen.add(device_id)
        original = baseline[device_id]
        after: list[Any] = ["" if value is None else value for value in values[1:9]]
        after[3] = values[4]
        if after[5] != original[5] or after[7] != original[7]:
            raise ValueError("位置表的状态或原机柜 ID 被修改，不能回导。")
        if isinstance(after[3], bool):
            raise TypeError("高度（U）必须是整数，不能是布尔值。")
        text = _text(after[0], "设备")
        after[0] = text
        if after == original:
            continue
        if original[5] != "在位":
            raise ValueError("仅支持编辑在位设备；缺失或未放置设备请先解决项目异常。")
        rack_name = _text(after[1], "机柜")
        room = _text(after[4], "机房")
        racks = [rack for rack in project.racks if rack.status == "active"
                 and rack.rack_name == rack_name and (rack.source_sheet or "") == room]
        if len(racks) != 1:
            raise ValueError("机房与机柜名无法定位唯一目标机柜。")
        u_label = _text(after[2], "U 位")
        match = _U_LABEL.fullmatch(u_label)
        if match is None:
            raise ValueError("U 位必须使用 21U 或 21-22U 格式。")
        low, high = int(match[1]), int(match[2] or match[1])
        if high < low or low < 1 or high > racks[0].height_u:
            raise ValueError("U 位范围顺序无效或超出目标机柜。")
        if isinstance(after[3], bool) or not isinstance(after[3], (int, float)) or after[3] != high - low + 1:
            raise ValueError("高度（U）必须是与 U 位范围一致的整数；调整高度请同步修改此列。")
        changes.append(PositionTableChange(device_id, racks[0].rack_id, low, high, text, original, after))
    if seen != set(baseline):
        raise ValueError("位置表缺少原设备行；删行不能作为删除设备操作。")
    return tuple(sorted(changes, key=lambda change: change.device_id))


def preview_position_table(database: Path, workbook: Path, sheet: str | None = None) -> PositionTablePlan:
    """Validate all rows, project revision and the complete final edit plan without writing."""
    project = load_project(database)
    revision = project_revision(project)
    path = normalize_path(workbook)
    plan = PositionTablePlan(project.project_id, revision, str(path), sha256_file(path))
    opened = None
    try:
        if path.suffix.lower() != ".xlsx":
            raise ValueError("设备位置表必须是 .xlsx 文件。")
        if project_error_conflicts(project):
            raise ValueError("当前项目仍有未解决错误，不能回导位置表。")
        if project.source_workbook is not None:
            source = Path(project.source_workbook)
            if not source.is_file() or sha256_file(source) != project.workbook_fingerprint:
                raise ValueError("项目源文件已变化或缺失，请重新扫描并导出后再回导。")
        opened = load_xlsx_workbook(path, data_only=False)
        candidates = [worksheet for worksheet in opened.worksheets
                      if any(name.startswith(_PREFIX) for name in worksheet.defined_names)]
        if sheet is not None:
            if sheet not in opened.sheetnames:
                raise ValueError("指定设备位置表工作表不存在。")
            if not isinstance(opened[sheet], Worksheet):
                raise ValueError("指定工作表不是单元格工作表。")
            selected = cast(Worksheet, opened[sheet])
        elif len(candidates) == 1:
            selected = candidates[0]
        else:
            raise ValueError("找不到唯一受支持的设备位置表；请指定工作表或重新导出。")
        plan = replace(plan, sheet_name=selected.title)
        metadata = _metadata(selected)
        if metadata.get("project_id") != project.project_id:
            raise ValueError("位置表属于另一个项目。")
        if metadata.get("project_revision") != revision:
            raise ValueError("位置表导出后项目已变化；请重新导出并基于当前版本编辑。")
        baseline = _baseline(project)
        if metadata.get("baseline") != baseline:
            raise ValueError("位置表原行基线不匹配或已被修改。")
        changes = _changes(selected, project, baseline)
        if project.source_workbook is None:
            _draft_candidate(project, changes)
        elif changes:
            from racktool.core.sync import plan_device_edits

            edits = plan_device_edits(project, _requests(changes))
            if edits.conflicts:
                return replace(plan, changes=changes, conflicts=edits.conflicts)
        if project_revision(load_project(database)) != revision or sha256_file(path) != plan.workbook_sha256:
            raise ValueError("项目或位置表在预览期间变化，请重新预览。")
        token = hashlib.sha256(_json([revision, plan.workbook_sha256, selected.title,
                                     [change.to_dict() for change in changes]]).encode()).hexdigest()
        return replace(plan, changes=changes, preview_token=token)
    except (ValueError, TypeError, KeyError) as error:
        return replace(plan, conflicts=(IdentityConflict(
            "position-table-invalid", "error", str(error),
        ),))
    finally:
        if opened is not None:
            opened.close()


@dataclass(frozen=True, slots=True)
class PositionTableResult:
    status: str
    plan: PositionTablePlan
    project: RackProject | None = None
    backup_path: str | None = None
    message: str = ""

    def to_dict(self) -> dict[str, Any]:
        return {"status": self.status, "plan": self.plan.to_dict(),
                "project": self.project.to_dict() if self.project is not None else None,
                "backup_path": self.backup_path, "message": self.message}


def apply_position_table(
    database: Path, workbook: Path, sheet: str | None = None, *,
    expected_preview_token: str | None = None,
) -> PositionTableResult:
    """Revalidate, then publish the complete batch once; never perform row-by-row commits."""
    database_path = normalize_path(database)
    plan = preview_position_table(database_path, workbook, sheet)
    if expected_preview_token is None:
        return PositionTableResult("rejected", plan, message="请先预览并提供 preview_token 后再提交。")
    if not plan.accepted:
        return PositionTableResult("rejected", plan, message="设备位置表校验失败，没有写入。")
    if expected_preview_token is not None and plan.preview_token != expected_preview_token:
        return PositionTableResult("rejected", plan, message="项目或位置表已变化，请重新预览。")
    project = load_project(database_path)
    if project_revision(project) != plan.project_revision:
        return PositionTableResult("rejected", plan, message="项目已变化，请重新预览。")
    if not plan.changes:
        return PositionTableResult("unchanged", plan, project, message="位置表没有设备修改。")
    if project.source_workbook is not None:
        from racktool.core.service import commit_write_plan
        from racktool.core.sync import plan_device_edits

        edits = plan_device_edits(project, _requests(plan.changes))
        if sha256_file(Path(plan.workbook_path)) != plan.workbook_sha256:
            return PositionTableResult("rejected", plan, message="位置表在提交前变化，请重新预览。")
        result = commit_write_plan(Path(project.source_workbook), database_path, edits)
        return PositionTableResult(result.status, plan, result.project, result.backup_path, result.message)
    from racktool.core.editing import _publish_draft

    candidate = _draft_candidate(project, plan.changes)
    previous_hash = sha256_file(database_path)
    if project_revision(load_project(database_path)) != plan.project_revision:
        return PositionTableResult("rejected", plan, message="项目在提交前变化，请重新预览。")
    backup = create_backup(database_path)
    try:
        if sha256_file(Path(plan.workbook_path)) != plan.workbook_sha256:
            raise ValueError("位置表在提交前变化，请重新预览。")
        persisted = _publish_draft(database_path, candidate, previous_hash)
    except Exception:
        discard_backup(backup)
        raise
    return PositionTableResult("applied", plan, persisted, str(backup), "设备位置表全部修改已一次提交。")
