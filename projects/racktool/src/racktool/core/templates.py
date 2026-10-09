"""Structured empty rack templates; XLSX rendering stays in the shared exporter."""

from __future__ import annotations

import hashlib

from openpyxl.utils import get_column_letter

from racktool.models.domain import CellRange, Rack
from racktool.models.project import RackProject, SourceMapping


def create_empty_rack_project(count: int = 18, height_u: int = 42) -> RackProject:
    """Create A/B rack rows without reading or modifying an existing workbook."""
    if not 1 <= count <= 99 or not 1 <= height_u <= 100:
        raise ValueError("每排机柜数量应为 1–99，高度应为 1–100 U。")
    fingerprint = hashlib.sha256(f"empty-racks:A:B:{count}:{height_u}".encode()).hexdigest()
    sheet_name = "两排空机柜"
    racks: list[Rack] = []
    mappings: list[SourceMapping] = []
    # Virtual source coordinates encode the row ordering, not the export layout.
    for band, prefix in enumerate(("A", "B")):
        title_row = 1 + band * (height_u + 3)
        for number in range(1, count + 1):
            column = 1 + (number - 1) * 4
            name = f"{prefix}{number:02d}"
            left, right = get_column_letter(column), get_column_letter(column + 2)
            title_range = f"{left}{title_row}:{right}{title_row}"
            rack = Rack(
                rack_id=f"empty-{name}",
                rack_name=name,
                height_u=height_u,
                source_sheet=sheet_name,
                bounds=CellRange(
                    min_row=title_row,
                    max_row=title_row + height_u,
                    min_col=column,
                    max_col=column + 2,
                    a1=f"{left}{title_row}:{right}{title_row + height_u}",
                ),
                start_row=title_row + 1,
                end_row=title_row + height_u,
                left_axis_column=column,
                right_axis_column=column + 2,
                device_columns=[column + 1],
                direction="descending",
                u_to_row={u: title_row + height_u + 1 - u for u in range(1, height_u + 1)},
                title_range=title_range,
            )
            racks.append(rack)
            mappings.append(SourceMapping(
                mapping_id=f"title-{name}",
                workbook_fingerprint=fingerprint,
                sheet_name=sheet_name,
                source_range=title_range,
                mapping_kind="rack_title",
                rack_id=rack.rack_id,
            ))
    return RackProject(
        project_id=f"empty-racks-{count}-{height_u}",
        source_workbook=None,
        workbook_fingerprint=fingerprint,
        layout_fingerprint=None,
        profile_id=None,
        racks=racks,
        mappings=mappings,
    )
