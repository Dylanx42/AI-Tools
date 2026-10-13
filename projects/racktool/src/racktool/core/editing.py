"""Edit newly created, unbound projects; source XLSX writes remain in Safe Sync."""

from __future__ import annotations

import hashlib
import os
import tempfile
from collections.abc import Callable
from dataclasses import replace
from pathlib import Path

from openpyxl.utils import get_column_letter

from racktool.core.identity import new_entity_id, normalize_path, sha256_file
from racktool.core.project import project_error_conflicts
from racktool.core.templates import _virtual_rack, create_empty_rack_project
from racktool.core.text import validate_excel_text
from racktool.models.domain import Device, Placement, Rack
from racktool.models.project import RackProject, SourceMapping
from racktool.persistence import load_project, save_project


def _validate_draft(project: RackProject) -> None:
    if (
        project.source_workbook is not None or project.profile_id is not None
        or project.metadata.get("project_mode") != "draft"
        or project.metadata.get("draft_schema_version") != 1
    ):
        raise ValueError("此编辑入口仅支持新建项目；已绑定 Excel 的项目请使用安全同步。")
    errors = project_error_conflicts(project)
    if errors:
        raise ValueError("; ".join(f"{item.code}: {item.message}" for item in errors))


def _database_path(database: Path) -> Path:
    path = normalize_path(database)
    if path.suffix.lower() != ".sqlite":
        raise ValueError("新建项目编辑必须使用 .sqlite 项目文件。")
    return path


def _publish_draft(
    database: Path, project: RackProject, previous_sha256: str | None,
) -> RackProject:
    _validate_draft(project)
    database.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary_name = tempfile.mkstemp(
        prefix=f".{database.name}.draft-", suffix=".sqlite", dir=database.parent,
    )
    os.close(descriptor)
    temporary = Path(temporary_name)
    try:
        save_project(temporary, project)
        persisted = load_project(temporary)
        _validate_draft(persisted)
        if previous_sha256 is None:
            if database.exists():
                raise FileExistsError("项目文件已存在，不覆盖已有项目。")
        elif not database.is_file() or sha256_file(database) != previous_sha256:
            raise ValueError("项目在编辑期间被其他程序修改，请重新读取后再操作。")
        os.replace(temporary, database)
        return persisted
    finally:
        temporary.unlink(missing_ok=True)


def create_draft_project(database: Path, count: int = 18, height_u: int = 42) -> RackProject:
    """Create a persistent A/B rack project with independent, stable identities."""
    path = _database_path(database)
    if path.exists():
        raise FileExistsError("项目文件已存在，不覆盖已有项目。")
    template = create_empty_rack_project(count, height_u)
    project_id = new_entity_id("project")
    fingerprint = hashlib.sha256(project_id.encode()).hexdigest()
    rack_ids = {rack.rack_id: new_entity_id("rack") for rack in template.racks}
    sheet_name = "新建机柜项目"
    project = replace(
        template, project_id=project_id, workbook_fingerprint=fingerprint,
        racks=[replace(rack, rack_id=rack_ids[rack.rack_id], source_sheet=sheet_name)
               for rack in template.racks],
        mappings=[replace(
            mapping, mapping_id=new_entity_id("map"), workbook_fingerprint=fingerprint,
            rack_id=rack_ids[mapping.rack_id or ""], sheet_name=sheet_name,
        ) for mapping in template.mappings],
        metadata={"project_mode": "draft", "draft_schema_version": 1},
    )
    return _publish_draft(path, project, None)


def _edit_draft(database: Path, edit: Callable[[RackProject], RackProject]) -> RackProject:
    path = _database_path(database)
    revision = sha256_file(path)
    project = load_project(path)
    _validate_draft(project)
    return _publish_draft(path, edit(project), revision)


def _text(value: str) -> str:
    validate_excel_text(value)
    return value


def _rack(project: RackProject, reference: str) -> Rack:
    matches = [rack for rack in project.racks if rack.status == "active" and (
        rack.rack_id == reference or rack.rack_name.casefold() == reference.casefold()
    )]
    if len(matches) != 1:
        raise ValueError("找不到唯一的目标机柜，请使用有效的机柜 ID 或名称。")
    return matches[0]


def _device_range(rack: Rack, start_u: int, end_u: int) -> str:
    low, high = sorted((start_u, end_u))
    if low < 1 or high > rack.height_u or any(u not in rack.u_to_row for u in range(low, high + 1)):
        raise ValueError("目标 U 位超出机柜范围。")
    if len(rack.device_columns) != 1:
        raise ValueError("新建项目机柜必须只有一列设备区。")
    rows = [rack.u_to_row[u] for u in range(low, high + 1)]
    column = get_column_letter(rack.device_columns[0])
    first, last = f"{column}{min(rows)}", f"{column}{max(rows)}"
    return first if first == last else f"{first}:{last}"


def add_draft_rack(
    database: Path, rack_name: str, *, height_u: int = 42, row: int = 1,
) -> RackProject:
    """Append a cabinet to an existing band or the next new band (row is one-based)."""
    name = _text(rack_name).strip()
    if not 1 <= height_u <= 100:
        raise ValueError("机柜高度应为 1–100 U。")

    def edit(project: RackProject) -> RackProject:
        if any(rack.rack_name.casefold() == name.casefold() for rack in project.racks):
            raise ValueError("机柜名称已存在。")
        bounds = [rack.bounds for rack in project.racks if rack.bounds is not None]
        title_rows = sorted({bound.min_row for bound in bounds})
        if not 1 <= row <= len(title_rows) + 1:
            raise ValueError("机柜排号必须是已有排，或紧接其后的新一排。")
        if row == len(title_rows) + 1:
            title_row = max((bound.max_row for bound in bounds), default=-2) + 3
            column = 1
        else:
            title_row = title_rows[row - 1]
            column = max(bound.max_col for bound in bounds if bound.min_row == title_row) + 2
        if column + 2 > 16_384 or title_row + height_u > 1_048_576:
            raise ValueError("机柜布局超过 Excel 的行列上限。")
        rack = _virtual_rack(name, new_entity_id("rack"), "新建机柜项目", column, title_row, height_u)
        mapping = SourceMapping(
            mapping_id=new_entity_id("map"), workbook_fingerprint=project.workbook_fingerprint or "",
            sheet_name=rack.source_sheet or "", source_range=rack.title_range or "",
            mapping_kind="rack_title", rack_id=rack.rack_id,
        )
        return replace(project, racks=[*project.racks, rack], mappings=[*project.mappings, mapping])

    return _edit_draft(database, edit)


def add_draft_device(
    database: Path, rack_reference: str, start_u: int, end_u: int, display_text: str,
) -> RackProject:
    text = _text(display_text)

    def edit(project: RackProject) -> RackProject:
        rack = _rack(project, rack_reference)
        source_range = _device_range(rack, start_u, end_u)
        device = Device(device_id=new_entity_id("device"), display_text=text)
        placement = Placement(
            placement_id=new_entity_id("placement"), device_id=device.device_id,
            rack_id=rack.rack_id, start_u=min(start_u, end_u), end_u=max(start_u, end_u),
        )
        mapping = SourceMapping(
            mapping_id=new_entity_id("map"), workbook_fingerprint=project.workbook_fingerprint or "",
            sheet_name=rack.source_sheet or "", source_range=source_range, mapping_kind="device",
            device_id=device.device_id, rack_id=rack.rack_id,
        )
        return replace(
            project, devices=[*project.devices, device], placements=[*project.placements, placement],
            mappings=[*project.mappings, mapping],
        )

    return _edit_draft(database, edit)


def edit_draft_device(
    database: Path, device_id: str, *, display_text: str | None = None,
    rack_reference: str | None = None, start_u: int | None = None, end_u: int | None = None,
) -> RackProject:
    if display_text is None and rack_reference is None and start_u is None and end_u is None:
        raise ValueError("请提供要修改的设备文字或位置。")
    if (start_u is None) != (end_u is None):
        raise ValueError("修改 U 位时必须同时提供起始 U 和结束 U。")
    if display_text is not None:
        _text(display_text)

    def edit(project: RackProject) -> RackProject:
        if not any(device.device_id == device_id for device in project.devices):
            raise ValueError("找不到指定设备。")
        placements = [p for p in project.placements if p.device_id == device_id and p.status == "active"]
        if len(placements) != 1:
            raise ValueError("设备必须有唯一的有效位置。")
        placement = placements[0]
        rack = _rack(project, rack_reference or placement.rack_id)
        if start_u is not None and end_u is not None:
            low, high = sorted((start_u, end_u))
        else:
            low, high = placement.start_u, placement.end_u
        source_range = _device_range(rack, low, high)
        return replace(
            project,
            devices=[replace(device, display_text=display_text) if (
                device.device_id == device_id and display_text is not None
            ) else device for device in project.devices],
            placements=[replace(p, rack_id=rack.rack_id, start_u=low, end_u=high)
                        if p.placement_id == placement.placement_id else p for p in project.placements],
            mappings=[replace(mapping, rack_id=rack.rack_id, source_range=source_range,
                              sheet_name=rack.source_sheet or "")
                      if mapping.mapping_kind == "device" and mapping.device_id == device_id
                      else mapping for mapping in project.mappings],
        )

    return _edit_draft(database, edit)
