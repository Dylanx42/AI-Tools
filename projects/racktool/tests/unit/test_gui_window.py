from __future__ import annotations

import os
import sys
from dataclasses import replace
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
pytest.importorskip("PySide6")

from racktool.gui import window as gui_window
from racktool.gui.session import GuiSession
from racktool.gui.window import CockpitWindow, create_app
from racktool.models.domain import Device
from racktool.models.project import IdentityConflict


def _fill_descending_axis(
    sheet: Worksheet,
    column: int,
    start_row: int,
    height: int,
) -> None:
    for offset, u_number in enumerate(range(height, 0, -1)):
        sheet.cell(start_row + offset, column, u_number)


def _make_layout(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "机柜"
    sheet.merge_cells("A1:C1")
    sheet["A1"] = "RACK-A"
    _fill_descending_axis(sheet, 1, 2, 12)
    _fill_descending_axis(sheet, 3, 2, 12)
    sheet["B2"] = "设备 A"
    sheet["B5"] = "设备 B"
    workbook.save(path)


@pytest.fixture(scope="module", autouse=True)
def qt_application() -> object:
    return create_app([])


def test_gui_entrypoint_starts_with_no_workbook_by_default(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    opened: list[Path | None] = []
    monkeypatch.setattr(sys, "argv", ["racktool-gui"])
    monkeypatch.setattr(
        gui_window,
        "launch",
        lambda workbook=None: opened.append(workbook) or 0,
    )

    with pytest.raises(SystemExit) as exit_info:
        gui_window.launch_main()

    assert exit_info.value.code == 0
    assert opened == [None]


def test_window_uses_navigable_pages_and_caps_visible_devices(tmp_path: Path) -> None:
    path = tmp_path / "layout.xlsx"
    _make_layout(path)
    session = GuiSession.open_workbook(path)
    session.project = replace(
        session.project,
        devices=[
            Device(device_id=f"device-{index}", display_text=f"设备 {index}")
            for index in range(110)
        ],
        placements=[],
        mappings=[],
    )
    cockpit = CockpitWindow()
    try:
        cockpit.load_session(session)

        assert cockpit.content_stack.count() == 4
        assert [button.text().strip()[-2:] for button in cockpit.nav_buttons] == [
            "总览",
            "机柜",
            "设备",
            "异常",
        ]
        assert cockpit.device_table.rowCount() == 100
        headers = [
            cockpit.device_table.horizontalHeaderItem(index).text()
            for index in range(cockpit.device_table.columnCount())
        ]
        assert headers == ["设备", "机柜", "U 位", "高度", "状态"]
        assert "ID" not in "".join(headers)
        assert "仅显示前 100 条" in cockpit.device_count_label.text()
    finally:
        cockpit.widget().close()


def test_rack_scene_only_renders_the_selected_rack(tmp_path: Path) -> None:
    path = tmp_path / "layout.xlsx"
    _make_layout(path)
    session = GuiSession.open_workbook(path)
    first_rack = session.project.racks[0]
    second_rack = replace(first_rack, rack_id="rack-b", rack_name="RACK-B")
    placements = list(session.project.placements)
    placements[1] = replace(placements[1], rack_id=second_rack.rack_id)
    session.project = replace(
        session.project,
        racks=[first_rack, second_rack],
        placements=placements,
    )
    cockpit = CockpitWindow()
    try:
        cockpit.load_session(session)

        assert cockpit.rack_list.count() == 2
        assert cockpit._rack_scene_device_count == 1
        assert cockpit._rack_scene_device_count < len(session.project.placements)

        cockpit.rack_list.setCurrentRow(1)
        assert cockpit.rack_title.text() == "RACK-B"
        assert cockpit._rack_scene_device_count == 1
    finally:
        cockpit.widget().close()


def test_move_drawer_stages_without_writing_source(tmp_path: Path) -> None:
    path = tmp_path / "layout.xlsx"
    _make_layout(path)
    before = path.read_bytes()
    session = GuiSession.open_workbook(path)
    device = next(
        item for item in session.project.devices if item.display_text == "设备 B"
    )
    cockpit = CockpitWindow()
    try:
        cockpit.load_session(session)
        cockpit._selected_device = device.device_id
        cockpit._refresh_device_detail()
        cockpit._open_move_drawer()
        cockpit.target_start_u.setValue(4)
        cockpit._update_move_preview()

        assert cockpit.drawer_add_button.isEnabled()
        cockpit._stage_drawer_move()

        assert path.read_bytes() == before
        assert len(session.pending_moves) == 1
        assert cockpit.pending_table.rowCount() == 1
        assert cockpit.sync_button.isEnabled()
        assert not cockpit.move_drawer.isVisible()
    finally:
        cockpit.widget().close()


def test_global_sync_is_the_only_device_move_write_control() -> None:
    source = Path(__file__).parents[2] / "src" / "racktool" / "gui" / "window.py"
    text = source.read_text(encoding="utf-8")

    assert text.count("apply_pending_moves()") == 1
    assert 'setObjectName("globalSyncButton")' in text
    assert "openpyxl" not in text


def test_export_button_creates_two_sheet_excel_workbook(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    path = tmp_path / "layout.xlsx"
    output = tmp_path / "gui-export.xlsx"
    _make_layout(path)
    session = GuiSession.open_workbook(path)
    cockpit = CockpitWindow()
    try:
        cockpit.load_session(session)
        monkeypatch.setattr(
            cockpit.QtWidgets.QFileDialog,
            "getSaveFileName",
            lambda *_args, **_kwargs: (str(output), "Excel 工作簿 (*.xlsx)"),
        )

        cockpit.export_button.click()

        assert cockpit.export_button.text() == "⇩ 导出表格"
        assert output.is_file()
        workbook = load_workbook(output, read_only=True, data_only=False)
        try:
            assert workbook.sheetnames == ["机柜图", "设备位置表"]
        finally:
            workbook.close()
        assert session.status_message == "已导出 gui-export.xlsx"
    finally:
        cockpit.widget().close()


def _make_sorted_layout(path: Path, *, multiline: bool = False) -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "机柜"
    sheet.merge_cells("A1:C1")
    sheet["A1"] = "RACK-B"
    _fill_descending_axis(sheet, 1, 2, 12)
    _fill_descending_axis(sheet, 3, 2, 12)
    sheet["B2"] = "交换机\n核心\n管理口" if multiline else "设备 A"
    sheet["B5"] = "设备 B"
    sheet.merge_cells("E1:G1")
    sheet["E1"] = "RACK-A"
    _fill_descending_axis(sheet, 5, 2, 12)
    _fill_descending_axis(sheet, 7, 2, 12)
    sheet["F2"] = "设备 C"
    sheet["F3"] = "设备 D"
    sheet["F4"] = "设备 E"
    workbook.save(path)


def test_issue_entry_opens_grouped_chinese_exceptions(tmp_path: Path) -> None:
    path = tmp_path / "layout.xlsx"
    _make_layout(path)
    session = GuiSession.open_workbook(path)
    session.project = replace(
        session.project,
        conflicts=[
            IdentityConflict(
                code="unresolved-u-axis",
                severity="warning",
                message="U axis at column 9 has no title",
                evidence=["机柜!I2:I13"],
            ),
            IdentityConflict(
                code="unresolved-u-axis",
                severity="warning",
                message="U axis at column 12 has no title",
                evidence=["机柜!L2:L13"],
            ),
        ],
    )
    cockpit = CockpitWindow()
    try:
        cockpit.load_session(session)
        cockpit.issue_entry_button.click()

        assert cockpit.content_stack.currentIndex() == CockpitWindow.PAGE_EXCEPTIONS
        headers = [
            cockpit.exception_table.horizontalHeaderItem(index).text()
            for index in range(cockpit.exception_table.columnCount())
        ]
        assert headers == ["级别", "问题", "具体位置", "如何处理"]
        visible = " ".join(
            cockpit.exception_table.item(0, column).text()
            for column in range(cockpit.exception_table.columnCount())
        )
        assert "unresolved-u-axis" not in visible
        assert "U axis" not in visible
        assert "重新扫描" in visible
        assert "2 处" in visible
        assert "机柜 · I2:I13" in visible
        assert "机柜 · L2:L13" in visible
        assert cockpit.exception_rescan_button.text() == "修正 Excel 后重新扫描"
        assert "unresolved-u-axis" in cockpit.exception_technical.toPlainText()
        assert cockpit.overview_issues.minimumHeight() >= 90
        assert cockpit.overview_issues.sizeHint().height() >= 90
    finally:
        cockpit.widget().close()


def test_sidebar_rack_list_can_stretch_and_sort(tmp_path: Path) -> None:
    path = tmp_path / "layout.xlsx"
    _make_sorted_layout(path)
    session = GuiSession.open_workbook(path)
    cockpit = CockpitWindow()
    try:
        cockpit.load_session(session)
        names = [
            cockpit.rack_list.item(index).text().splitlines()[0]
            for index in range(cockpit.rack_list.count())
        ]
        cockpit.rack_sort_box.setCurrentIndex(1)
        sorted_names = [
            cockpit.rack_list.item(index).text().splitlines()[0]
            for index in range(cockpit.rack_list.count())
        ]

        assert cockpit.sidebar.minimumWidth() < cockpit.sidebar.maximumWidth()
        assert cockpit.nav_splitter.objectName() == "navSplitter"
        assert cockpit.nav_splitter.handleWidth() >= 6
        assert (
            cockpit.rack_list.horizontalScrollBarPolicy()
            != cockpit.QtCore.Qt.ScrollBarPolicy.ScrollBarAlwaysOff
        )
        assert [cockpit.rack_sort_box.itemText(index) for index in range(3)] == [
            "按源表位置排序",
            "按名称排序",
            "按占用率排序",
        ]
        assert any("RACK-B" in name for name in names)
        assert sorted_names[0].startswith("●  RACK-A")
        assert all("..." not in name and "…" not in name for name in sorted_names)
        first_item = cockpit.rack_list.item(0)
        widest_line = max(first_item.text().splitlines(), key=len)
        assert first_item.sizeHint().width() >= (
            cockpit.rack_list.fontMetrics().horizontalAdvance(widest_line) + 28
        )
    finally:
        cockpit.widget().close()


def test_multiline_text_reaches_preview_table_detail_and_pending(tmp_path: Path) -> None:
    path = tmp_path / "layout.xlsx"
    _make_sorted_layout(path, multiline=True)
    before = path.read_bytes()
    session = GuiSession.open_workbook(path)
    device = next(
        row for row in session.device_rows() if "交换机" in str(row["display_text"])
    )
    cockpit = CockpitWindow()
    try:
        cockpit.load_session(session)
        cockpit._selected_device = str(device["device_id"])
        cockpit._selected_rack = str(device["rack_id"])
        cockpit._render_selected_rack()
        cockpit._refresh_device_detail()
        cockpit._refresh_device_table()

        matching = [
            cockpit.device_table.item(row, 0).text()
            for row in range(cockpit.device_table.rowCount())
            if cockpit.device_table.item(row, 0) is not None
            and "交换机" in cockpit.device_table.item(row, 0).text()
        ]
        scene_text = " ".join(
            item.toPlainText()
            for item in cockpit.rack_scene.items()
            if hasattr(item, "toPlainText")
        )
        assert "交换机\n核心\n管理口" in matching[0]
        assert "交换机\n核心\n管理口" in cockpit.device_name.text()
        assert "交换机\n核心\n管理口" in scene_text

        cockpit._open_move_drawer()
        cockpit.target_start_u.setValue(4)
        cockpit._update_move_preview()
        cockpit._stage_drawer_move()
        pending_text = cockpit.pending_table.item(0, 0).text()
        assert "交换机\n核心\n管理口" in pending_text
        assert path.read_bytes() == before
    finally:
        cockpit.widget().close()


def test_device_filters_and_overview_sheet_view(tmp_path: Path) -> None:
    path = tmp_path / "layout.xlsx"
    _make_sorted_layout(path, multiline=True)
    session = GuiSession.open_workbook(path)
    cockpit = CockpitWindow()
    try:
        cockpit.load_session(session)
        rack_a = next(row for row in session.rack_rows() if row["rack_name"] == "RACK-A")
        rack_index = cockpit.device_rack_filter.findData(str(rack_a["rack_id"]))
        cockpit.device_rack_filter.setCurrentIndex(rack_index)

        labels = [
            cockpit.device_table.item(row, 0).text()
            for row in range(cockpit.device_table.rowCount())
        ]
        assert cockpit.device_table.rowCount() == 3
        assert all("设备 A" not in label and "交换机" not in label for label in labels)
        assert cockpit.overview_sheet_box.currentText() == "机柜"
        assert cockpit.overview_view.objectName() == "overviewSheetView"
        assert len(cockpit.overview_scene.items()) > 2
        scene_text = " ".join(
            item.toPlainText()
            for item in cockpit.overview_scene.items()
            if hasattr(item, "toPlainText")
        )
        assert "RACK-A" in scene_text
        assert "交换机" in scene_text
        assert "核心" in scene_text
    finally:
        cockpit.widget().close()


def test_unexpected_sync_error_is_caught_without_losing_pending_move(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    path = tmp_path / "layout.xlsx"
    _make_layout(path)
    before = path.read_bytes()
    session = GuiSession.open_workbook(path)
    device = next(item for item in session.project.devices if item.display_text == "设备 B")
    rack_id = session.project.racks[0].rack_id
    session.stage_move(device.device_id, rack_id, 4, 4)
    cockpit = CockpitWindow()
    shown: dict[str, object] = {}

    def fail_sync() -> object:
        raise OSError("raw internal failure")

    try:
        cockpit.load_session(session)
        monkeypatch.setattr(session, "apply_pending_moves", fail_sync)
        monkeypatch.setattr(
            cockpit.QtWidgets.QMessageBox,
            "question",
            lambda *_args, **_kwargs: cockpit.QtWidgets.QMessageBox.StandardButton.Yes,
        )
        monkeypatch.setattr(
            cockpit,
            "_show_error",
            lambda title, error: shown.update(title=title, error=error),
        )

        cockpit._sync_pending()

        assert shown["title"] == "同步未完成"
        assert isinstance(shown["error"], OSError)
        assert len(session.pending_moves) == 1
        assert path.read_bytes() == before
    finally:
        cockpit.widget().close()
