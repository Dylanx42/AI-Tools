"""Reusable behavioral checks used by pytest and the downloadable corpus runner."""

from __future__ import annotations

import json
import shutil
from collections import Counter
from copy import deepcopy
from dataclasses import asdict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

from racktool.core import (
    add_draft_device,
    add_draft_rack,
    analyze_workbook,
    create_draft_project,
    edit_draft_device,
    export_project_workbook,
    scan_workbook,
)
from racktool.core.identity import sha256_file
from racktool.core.project import project_error_conflicts
from racktool.core.service import commit_write_plan, import_project, rescan_project
from racktool.core.sync import plan_device_add, plan_device_move, plan_device_rename
from racktool.models.project import RackProject
from racktool.persistence import load_project
from racktool.profiles import load_profile
from tests.layout_corpus import DraftCase, LayoutCase, build_layout, device_range


def _check_project(project: RackProject, facts: dict[str, Any]) -> None:
    assert not project_error_conflicts(project)
    assert len(project.racks) == len(facts["racks"])
    assert {(rack.rack_name, rack.height_u) for rack in project.racks} == {
        (rack["name"], rack["height"]) for rack in facts["racks"]
    }
    names = {rack.rack_id: rack.rack_name for rack in project.racks}
    texts = {device.device_id: device.display_text for device in project.devices}
    assert len(project.devices) == len(project.placements) == len(facts["devices"])
    assert Counter((texts[p.device_id], names[p.rack_id], p.start_u, p.end_u) for p in project.placements) == Counter(
        (device["text"], device["rack"], device["low"], device["high"]) for device in facts["devices"]
    )
    mappings = {mapping.device_id: mapping for mapping in project.mappings if mapping.mapping_kind == "device"}
    for placement in project.placements:
        expected = next(device for device in facts["devices"] if (
            device["text"] == texts[placement.device_id] and device["rack"] == names[placement.rack_id]
            and device["low"] == placement.start_u and device["high"] == placement.end_u
        ))
        if "source_range" in expected:
            assert mappings[placement.device_id].source_range == expected["source_range"]


def _check_export(path: Path, facts: dict[str, Any]) -> None:
    analysis = analyze_workbook(path)
    sheets = [sheet for sheet in analysis.sheets if sheet.racks]
    racks = [rack for sheet in sheets for rack in sheet.racks]
    assert len(racks) == len(facts["racks"])
    assert {(rack.rack_name, rack.height_u) for rack in racks} == {
        (rack["name"], rack["height"]) for rack in facts["racks"]
    }
    direction = facts.get("case", {}).get("direction", "descending")
    assert all(rack.direction == direction for rack in racks)
    names = {rack.candidate_id: rack.rack_name for rack in racks}
    texts = {device.candidate_id: device.display_text for sheet in sheets for device in sheet.devices}
    placements = [p for sheet in sheets for p in sheet.placements]
    assert len(texts) == len(placements) == len(facts["devices"])
    assert Counter((texts[p.device_candidate_id], names[p.rack_candidate_id], p.start_u, p.end_u)
                   for p in placements) == Counter(
        (device["text"], device["rack"], device["low"], device["high"]) for device in facts["devices"]
    )
    assert all(len(rack.device_columns) == 1 for rack in racks)
    workbook = load_workbook(path)
    try:
        assert all(cell.data_type != "f" for sheet in workbook.worksheets for row in sheet.iter_rows() for cell in row)
    finally:
        workbook.close()


def _cell_set(ranges: list[str]) -> set[tuple[int, int]]:
    cells: set[tuple[int, int]] = set()
    for source_range in ranges:
        min_col, min_row, max_col, max_row = range_boundaries(source_range)
        assert None not in (min_col, min_row, max_col, max_row)
        cells.update((row, column) for row in range(min_row, max_row + 1) for column in range(min_col, max_col + 1))
    return cells


def _check_unrelated(before: Path, after: Path, touched: list[str]) -> None:
    from openpyxl.utils.cell import coordinate_to_tuple

    allowed = _cell_set(touched)
    before_info, after_info = scan_workbook(before), scan_workbook(after)
    assert [sheet.name for sheet in before_info.sheets] == [sheet.name for sheet in after_info.sheets]
    for original, current in zip(before_info.sheets, after_info.sheets, strict=True):
        assert original.state == current.state
        assert original.row_heights == current.row_heights
        assert original.column_widths == current.column_widths
        assert {cell.coordinate: cell for cell in original.cells if (
            original.name != "Racks" or coordinate_to_tuple(cell.coordinate) not in allowed
        )} == {cell.coordinate: cell for cell in current.cells if (
            current.name != "Racks" or coordinate_to_tuple(cell.coordinate) not in allowed
        )}
    workbook = load_workbook(after)
    try:
        notes = workbook["Preserved"]
        assert notes.sheet_state == "hidden"
        assert notes["A2"].comment.text == "SYNTHETIC COMMENT"
        assert notes["A3"].hyperlink.target == "https://example.com"
        assert notes["D3"].value == "=1+2" and notes["D3"].data_type == "f"
        assert workbook.defined_names["SyntheticKeep"].attr_text == "'Preserved'!$A$1"
    finally:
        workbook.close()


def run_layout_case(case: LayoutCase, directory: Path) -> dict[str, Any]:
    facts = build_layout(case, directory)
    expected = deepcopy(facts)
    source, edited, database = directory / "source.xlsx", directory / "edited.xlsx", directory / "project.sqlite"
    shutil.copy2(source, edited)
    profile_path = directory / "profile.yaml"
    profile = load_profile(profile_path) if profile_path.exists() else None
    project = import_project(edited, database, profile=profile)
    _check_project(project, expected)
    original_ids = {device.device_id for device in project.devices}
    touched: list[str] = []
    for operation in ("move", "add", "rename"):
        before_workbook, before_database = edited.read_bytes(), database.read_bytes()
        target = facts["racks"][-1]
        rack_id = next(rack.rack_id for rack in project.racks if rack.rack_name == target["name"])
        if operation == "move":
            original = expected["devices"][0]
            device_id = next(device.device_id for device in project.devices if device.display_text == original["text"])
            plan = plan_device_move(project, device_id, rack_id, 6, 7)
            touched.extend([original["source_range"], device_range(target, 6, 7)])
            original.update(rack=target["name"], low=6, high=7, source_range=device_range(target, 6, 7))
        elif operation == "add":
            text = "=SYNTHETIC ADD\n" + case.case_id
            plan = plan_device_add(project, rack_id, 9, 10, text)
            source_range = device_range(target, 9, 10)
            touched.append(source_range)
            expected["devices"].append({"text": text, "rack": target["name"], "low": 9, "high": 10, "source_range": source_range})
        else:
            original = expected["devices"][1]
            device_id = next(device.device_id for device in project.devices if device.display_text == original["text"])
            text = "=SYNTHETIC RENAME\n" + case.case_id
            plan = plan_device_rename(project, device_id, text)
            touched.append(original["source_range"])
            original["text"] = text
        assert not plan.conflicts, [conflict.to_dict() for conflict in plan.conflicts]
        assert len(plan.actions) == 1
        action = plan.actions[0]
        expected_range = expected["devices"][-1]["source_range"] if operation == "add" else (
            expected["devices"][0 if operation == "move" else 1]["source_range"]
        )
        assert action.new_range == expected_range
        assert edited.read_bytes() == before_workbook and database.read_bytes() == before_database
        result = commit_write_plan(edited, database, plan)
        assert result.status == "applied", result.errors
        assert result.project is not None and result.backup_path is not None
        assert Path(result.backup_path).read_bytes() == before_workbook
        project = load_project(database)
        assert original_ids.issubset(device.device_id for device in project.devices)
        _check_project(project, expected)
    for low, high in ((4, 4), (case.height, case.height + 1)):
        before_workbook, before_database = edited.read_bytes(), database.read_bytes()
        rack_id = next(rack.rack_id for rack in project.racks if rack.rack_name == facts["racks"][0]["name"])
        rejected = plan_device_add(project, rack_id, low, high, "SYNTHETIC INVALID")
        assert rejected.conflicts
        assert commit_write_plan(edited, database, rejected).status == "rejected"
        assert edited.read_bytes() == before_workbook and database.read_bytes() == before_database
    rescanned = rescan_project(edited, database)
    assert rescanned.accepted
    assert {device.device_id for device in rescanned.project.devices} == {device.device_id for device in project.devices}
    _check_project(rescanned.project, expected)
    _check_unrelated(source, edited, touched)
    output = export_project_workbook(rescanned.project, directory / "standard-export.xlsx")
    _check_export(output, expected)
    (directory / "expected-edited.json").write_text(json.dumps(expected, ensure_ascii=False, indent=2), encoding="utf-8")
    return {"case_id": case.case_id, "kind": "layout", "passed": True, "family": case.family,
            "height": case.height, "direction": case.direction, "label": case.label,
            "racks": len(project.racks), "devices": len(project.devices), "committed_edits": 3,
            "rejected_edits": 2, "source_sha256": sha256_file(source), "edited_sha256": sha256_file(edited)}


def run_draft_case(case: DraftCase, directory: Path) -> dict[str, Any]:
    directory.mkdir(parents=True, exist_ok=True)
    database = directory / "project.sqlite"
    project = create_draft_project(database, case.count, case.height)
    initial_ids = {rack.rack_id for rack in project.racks}
    facts: dict[str, Any] = {"case": asdict(case), "synthetic": True, "racks": [
        {"name": f"{prefix}{number:02d}", "height": case.height} for prefix in "AB" for number in range(1, case.count + 1)
    ], "devices": []}
    extra_height = 47 if case.height != 47 else 42
    add_draft_rack(database, "C01", height_u=extra_height, row=3)
    facts["racks"].append({"name": "C01", "height": extra_height})
    fill_end = max(1, case.height * case.density // 100)
    add_draft_device(database, "A01", 1, fill_end, "SYNTHETIC OCCUPANCY")
    facts["devices"].append({"text": "SYNTHETIC OCCUPANCY", "rack": "A01", "low": 1, "high": fill_end})
    text = ("=" if case.text_mode == "formula-text" else "") + "SYNTHETIC DUPLICATE\nMODEL"
    project = add_draft_device(database, "A01", case.height - 4, case.height, text)
    device_id = next(device.device_id for device in project.devices if device.display_text == text)
    facts["devices"].append({"text": text, "rack": "A01", "low": case.height - 4, "high": case.height})
    project = add_draft_device(database, "B01", 3, 4, text)
    facts["devices"].append({"text": text, "rack": "B01", "low": 3, "high": 4})
    _check_project(project, facts)
    assert len([device for device in project.devices if device.display_text == text]) == 2
    changed_text = ("=" if case.text_mode == "formula-text" else "") + "SYNTHETIC EDITED\nNEW MODEL"
    project = edit_draft_device(database, device_id, display_text=changed_text)
    facts["devices"][1]["text"] = changed_text
    _check_project(project, facts)
    project = edit_draft_device(database, device_id, rack_reference="C01", start_u=6, end_u=8)
    facts["devices"][1].update(rack="C01", low=6, high=8)
    assert any(device.device_id == device_id for device in project.devices)
    assert initial_ids.issubset(rack.rack_id for rack in project.racks)
    _check_project(project, facts)
    for rack, low, high in (("A01", 1, 3), ("C01", extra_height, extra_height + 1)):
        before = database.read_bytes()
        try:
            edit_draft_device(database, device_id, rack_reference=rack, start_u=low, end_u=high)
        except ValueError:
            pass
        else:
            raise AssertionError("Invalid draft edit was accepted")
        assert database.read_bytes() == before
    project = load_project(database)
    _check_project(project, facts)
    output = export_project_workbook(project, directory / "edited.xlsx")
    _check_export(output, facts)
    (directory / "expected-edited.json").write_text(json.dumps(facts, ensure_ascii=False, indent=2), encoding="utf-8")
    return {"case_id": case.case_id, "kind": "draft", "passed": True, "count": case.count,
            "height": case.height, "density": case.density, "text_mode": case.text_mode,
            "racks": len(project.racks), "devices": len(project.devices), "committed_edits": 6,
            "rejected_edits": 2, "edited_sha256": sha256_file(output)}
