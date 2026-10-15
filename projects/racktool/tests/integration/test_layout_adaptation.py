from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from racktool.cli.layout_cli import handle_layout_cli, register_layout_cli
from racktool.core.analyzer import analyze_workbook
from racktool.core.service import import_project, rescan_project
from racktool.profiles.adaptation import check_layout, inspect_layout, propose_layout
from racktool.profiles.loader import load_profile
from racktool.profiles.schema import ProfileError


def _workbook(path: Path, *, second_room: bool = False) -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Room A"
    sheet.merge_cells("G2:H2")
    sheet["G2"] = "Rack secret 192.0.2.1"
    for offset, u in enumerate(range(12, 0, -1)):
        sheet.cell(4 + offset, 6, u)
        sheet.cell(4 + offset, 9, u)
    sheet.merge_cells("G4:H5")
    sheet["G4"] = "Do not execute this private device text"
    if second_room:
        other = workbook.create_sheet("Room B")
        other.merge_cells("C1:D1")
        other["C1"] = "Another rack"
        for offset, u in enumerate(range(1, 13)):
            other.cell(4 + offset, 2, u)
            other.cell(4 + offset, 5, u)
        other.merge_cells("C4:D5")
        other["C4"] = "Second device"
    workbook.save(path)


def _hints() -> dict[str, object]:
    return {
        "title_range": "G2:H2", "left_axis": "F", "right_axis": "I",
        "start_row": 4, "end_row": 15, "device_range": "G4:H15", "direction": "descending",
    }


def test_unknown_fixed_layout_proposal_import_and_rescan(tmp_path: Path) -> None:
    source = tmp_path / "unknown.xlsx"
    _workbook(source)
    original = source.read_bytes()
    digest = hashlib.sha256(original).hexdigest()
    assert not analyze_workbook(source).sheets[0].racks
    inspection = inspect_layout(source)
    assert inspection["source_sha256"] == digest
    assert inspection["status"] == "review_required"
    assert inspection["sheets"][0]["profile_candidate"] is None
    assert inspection["sheets"][0]["u_axes"]
    profile_path = tmp_path / "reviewed.yaml"
    result = propose_layout(
        source, profile_path, sheet_name="Room A", profile_id="reviewed-fixed",
        hints=_hints(), expected_sha256=digest,
    )
    assert result["status"] == "ready_for_import"
    assert result["dry_run"]["dry_run"] is True
    assert result["dry_run"]["selected_sheets"] == ("Room A",)
    yaml_text = profile_path.read_text()
    assert "secret" not in yaml_text
    assert "192.0.2.1" not in yaml_text
    assert "private device text" not in yaml_text
    database = tmp_path / "project.sqlite"
    project = import_project(source, database, profile=load_profile(profile_path))
    assert len(project.devices) == 1
    before_id = project.devices[0].device_id
    rescanned = rescan_project(source, database)
    assert rescanned.project.devices[0].device_id == before_id
    assert source.read_bytes() == original
    assert check_layout(source, profile_path, expected_sha256=digest)["status"] == "ready_for_import"


def test_bad_hints_and_stale_hash_do_not_force_application(tmp_path: Path) -> None:
    source = tmp_path / "unknown.xlsx"
    _workbook(source)
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    profile_path = tmp_path / "wrong.yaml"
    result = propose_layout(
        source, profile_path, sheet_name="Room A", profile_id="wrong-direction",
        hints={**_hints(), "direction": "ascending"},
    )
    assert result["status"] == "review_required"
    assert result["dry_run"]["status"] == "rejected"
    assert result["dry_run"]["selected_sheets"] == ()
    assert result["review_reasons"]
    workbook = load_workbook(source)
    workbook["Room A"]["Z1"] = "Changed after inspection"
    workbook.save(source)
    workbook.close()
    with pytest.raises(ProfileError, match="SHA256 changed"):
        check_layout(source, profile_path, expected_sha256=digest)
    with pytest.raises(ProfileError, match="complete geometry"):
        propose_layout(
            source, tmp_path / "partial.yaml", sheet_name="Room A", profile_id="partial",
            hints={"title_range": "G2:H2"},
        )
    with pytest.raises(FileExistsError):
        propose_layout(
            source, profile_path, sheet_name="Room A", profile_id="exists", hints=_hints(),
        )


def test_heterogeneous_sheets_have_separate_persisted_rules(tmp_path: Path) -> None:
    source = tmp_path / "rooms.xlsx"
    _workbook(source, second_room=True)
    inspection = inspect_layout(source)
    assert inspection["batch_contract"]["combined_heterogeneous_project_supported"] is False
    first = tmp_path / "room-a.yaml"
    second = tmp_path / "room-b.yaml"
    propose_layout(source, first, sheet_name="Room A", profile_id="room-a", hints=_hints())
    proposal = propose_layout(
        source, second, sheet_name="Room B", profile_id="room-b", hints={
            "title_range": "C1:D1", "left_axis": "B", "right_axis": "E",
            "start_row": 4, "end_row": 15, "device_range": "C4:D15", "direction": "ascending",
        },
    )
    assert proposal["status"] == "ready_for_import"
    for room, rule in (("Room A", first), ("Room B", second)):
        database = tmp_path / (rule.stem + ".sqlite")
        project = import_project(source, database, profile=load_profile(rule))
        assert {rack.source_sheet for rack in project.racks} == {room}
        rescanned = rescan_project(source, database)
        assert {rack.source_sheet for rack in rescanned.project.racks} == {room}


def test_cli_json_and_review_exit_status(
    tmp_path: Path, capsys: pytest.CaptureFixture[str],
) -> None:
    source = tmp_path / "unknown.xlsx"
    _workbook(source)
    parser = argparse.ArgumentParser()
    register_layout_cli(parser.add_subparsers(dest="command", required=True))
    assert handle_layout_cli(parser.parse_args(["layout", "inspect", str(source)])) == 0
    assert json.loads(capsys.readouterr().out)["status"] == "review_required"
    args = parser.parse_args([
        "layout", "propose", str(source), str(tmp_path / "invalid.yaml"),
        "--sheet", "Room A", "--profile-id", "invalid-geometry",
        "--title-range", "G2:H2", "--left-axis", "F", "--right-axis", "I",
        "--start-row", "4", "--end-row", "15", "--device-range", "G4:H15",
        "--direction", "ascending",
    ])
    assert handle_layout_cli(args) == 1
    assert json.loads(capsys.readouterr().out)["status"] == "review_required"


def test_structural_suggestion_is_scoped_and_never_automatically_approved(tmp_path: Path) -> None:
    source = tmp_path / "standard.xlsx"
    _workbook(source)
    workbook = load_workbook(source)
    sheet = workbook["Room A"]
    sheet.unmerge_cells("G2:H2")
    sheet["G2"] = None
    sheet.merge_cells("F3:I3")
    sheet["F3"] = "Standard title"
    workbook.save(source)
    workbook.close()
    inspection = inspect_layout(source)
    assert inspection["status"] == "review_required"
    candidate = inspection["sheets"][0]["profile_candidate"]
    assert candidate["status"] == "ready_for_import"
    assert candidate["human_confirmation_required"] is True
    assert candidate["selected_analysis"]["sheets"][0]["devices"]
    assert "Standard title" not in json.dumps(candidate["profile"])
    output = tmp_path / "structural.yaml"
    proposal = propose_layout(source, output, sheet_name="Room A", profile_id="structural")
    assert proposal["status"] == "ready_for_import"
    assert load_profile(output).rack.title_mode == "merged_cell_above_u_axis"


def test_low_confidence_fixed_layout_requires_review_despite_matching(tmp_path: Path) -> None:
    source = tmp_path / "short-axis.xlsx"
    _workbook(source)
    workbook = load_workbook(source)
    for row in range(4, 16):
        workbook["Room A"].cell(row, 6, None).value = 12 - row if row < 12 else None
        workbook["Room A"].cell(row, 9, None).value = 12 - row if row < 12 else None
    workbook.save(source)
    workbook.close()
    result = propose_layout(
        source, tmp_path / "short.yaml", sheet_name="Room A", profile_id="short-axis",
        hints={**_hints(), "end_row": 11, "device_range": "G4:H11"},
    )
    assert result["dry_run"]["status"] == "applied"
    assert result["status"] == "review_required"
    assert "low-confidence-rack-candidates" in " ".join(result["review_reasons"])
    assert result["selected_analysis"] is not None


def test_existing_geometry_issues_are_not_silently_ignored(tmp_path: Path) -> None:
    source = tmp_path / "axis-gap.xlsx"
    _workbook(source)
    workbook = load_workbook(source)
    workbook["Room A"]["F7"] = None
    workbook["Room A"]["I7"] = None
    workbook.save(source)
    workbook.close()
    result = propose_layout(
        source, tmp_path / "gap.yaml", sheet_name="Room A", profile_id="gap-axis", hints=_hints(),
    )
    assert result["status"] == "review_required"
    assert result["review_reasons"]
    assert not result["dry_run"]["selected_sheets"]


def test_layout_commands_are_registered_at_top_level(
    tmp_path: Path, capsys: pytest.CaptureFixture[str],
) -> None:
    from racktool.cli.main import main

    source = tmp_path / "source.xlsx"
    _workbook(source)
    assert main(["layout", "inspect", str(source)]) == 0
    inspection = json.loads(capsys.readouterr().out)
    digest = inspection["source_sha256"]
    profile = tmp_path / "reviewed.yaml"
    assert main([
        "layout", "propose", str(source), str(profile), "--sheet", "Room A",
        "--profile-id", "reviewed-top-level", "--title-range", "G2:H2",
        "--left-axis", "F", "--right-axis", "I", "--start-row", "4", "--end-row", "15",
        "--device-range", "G4:H15", "--direction", "descending", "--expected-sha256", digest,
    ]) == 0
    assert json.loads(capsys.readouterr().out)["status"] == "ready_for_import"
    assert main(["layout", "check", str(source), str(profile), "--expected-sha256", digest]) == 0
    assert json.loads(capsys.readouterr().out)["status"] == "ready_for_import"
    with pytest.raises(SystemExit) as rejected:
        main(["layout", "check", str(source), str(profile), "--expected-sha256", "0" * 64])
    assert rejected.value.code == 2
    assert "SHA256 changed" in capsys.readouterr().err
