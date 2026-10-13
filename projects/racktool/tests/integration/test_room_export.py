from collections import Counter
from dataclasses import replace
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import PatternFill
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

from racktool.cli.main import main
from racktool.core import (
    analyze_workbook,
    commit_write_plan,
    create_draft_project,
    export_project_workbook,
    import_project,
    plan_device_add,
    plan_device_rename,
)
from racktool.core.identity import sha256_file
from racktool.core.workbook import scan_workbook
from racktool.models.domain import Device


def make_rooms(path: Path):
    workbook = Workbook()
    first = workbook.active
    assert first is not None
    first.title = "Room-Z"
    second = workbook.create_sheet("Room-A")
    for sheet, name, ascending in ((first, "A01", False), (second, "B01", True)):
        sheet.merge_cells("C3:E3")
        sheet["C3"] = name
        sheet["C3"].fill = PatternFill("solid", fgColor="FF9900")
        for offset, unit in enumerate(range(1, 13) if ascending else range(12, 0, -1), 4):
            sheet.cell(offset, 3, unit)
            sheet.cell(offset, 5, unit)
            sheet.row_dimensions[offset].height = 26
        sheet.column_dimensions["D"].width = 42
        low, high = (6, 7) if ascending else (9, 10)
        top = 4 + (low - 1 if ascending else 12 - high)
        sheet.merge_cells(start_row=top, start_column=4, end_row=top + 1, end_column=4)
        sheet.cell(top, 4, "Server\n" + name)
        sheet.freeze_panes = "D4"
        sheet.print_area = "C3:E15"
        sheet.print_title_rows = "3:3"
        sheet.page_setup.orientation = "landscape"
    notes = workbook.create_sheet("Notes")
    notes["A1"] = "=1+2"
    notes["A2"].comment = Comment("original comment", "owner")
    notes["A3"] = "Link"
    notes["A3"].hyperlink = "https://example.com"
    notes.sheet_state = "hidden"
    validation = DataValidation(type="list", formula1='"A,B"')
    notes.add_data_validation(validation)
    validation.add("B1")
    workbook.defined_names.add(DefinedName("KeepMe", attr_text="'Notes'!$A$1"))
    workbook.save(path)
    workbook.close()
    return import_project(path, path.with_suffix(".sqlite"))


def test_standard_export_keeps_axis_direction(tmp_path: Path) -> None:
    project = make_rooms(tmp_path / "input.xlsx")
    output = export_project_workbook(project, tmp_path / "export.xlsx")
    analysis = analyze_workbook(output)
    assert Counter((r.rack_name, r.direction) for s in analysis.sheets for r in s.racks) == Counter({
        ("A01", "descending"): 1, ("B01", "ascending"): 1,
    })
    assert Counter((d.display_text, p.start_u, p.end_u)
                   for s in analysis.sheets for p in s.placements
                   for d in s.devices if d.candidate_id == p.device_candidate_id) == Counter({
        ("Server\nA01", 9, 10): 1, ("Server\nB01", 6, 7): 1,
    })


def test_standard_export_has_room_tabs_and_compact_positions(tmp_path: Path) -> None:
    project = make_rooms(tmp_path / "input.xlsx")
    output = export_project_workbook(project, tmp_path / "export.xlsx")
    workbook = load_workbook(output)
    try:
        assert workbook.sheetnames == ["Room-Z", "Room-A", "设备位置表"]
        sheet = workbook["设备位置表"]
        assert [c.value for c in sheet[1]] == ["序号", "设备", "机柜", "U 位", "高度（U）", "机房", "状态", "设备 ID", "机柜 ID"]
        assert sheet["D2"].value == "9-10U" and sheet["F2"].value == "Room-Z"
        assert sheet["D3"].value == "6-7U" and sheet["F3"].value == "Room-A"
        assert sheet.freeze_panes == "A2"
        assert sheet.column_dimensions["H"].hidden and sheet.column_dimensions["I"].hidden
        assert next(iter(sheet.tables.values())).ref == "A1:I3"
    finally:
        workbook.close()


def test_source_export_keeps_original_layout_and_metadata(tmp_path: Path) -> None:
    source = tmp_path / "input.xlsx"
    project = make_rooms(source)
    original = source.read_bytes()
    output = export_project_workbook(project, tmp_path / "export.xlsx", layout="source")
    assert source.read_bytes() == original
    changed_parts = {"xl/workbook.xml", "xl/_rels/workbook.xml.rels", "xl/styles.xml", "[Content_Types].xml"}
    with ZipFile(BytesIO(original)) as before, ZipFile(output) as after:
        assert all(before.read(name) == after.read(name) for name in before.namelist() if name not in changed_parts)
    assert scan_workbook(output).sheets[:3] == scan_workbook(source).sheets
    workbook = load_workbook(output)
    try:
        assert workbook.sheetnames == ["Room-Z", "Room-A", "Notes", "设备位置表"]
        assert workbook["Room-Z"].print_area == "'Room-Z'!$C$3:$E$15"
        assert workbook["Room-A"].freeze_panes == "D4"
        assert workbook["Notes"]["A1"].data_type == "f"
        assert workbook["Notes"]["A2"].comment.text == "original comment"
        assert workbook["Notes"]["A3"].hyperlink.target == "https://example.com"
        assert workbook["Notes"].data_validations.count == 1
        assert "KeepMe" in workbook.defined_names
    finally:
        workbook.close()


def test_source_export_does_not_overwrite_existing_position_sheet(tmp_path: Path) -> None:
    source = tmp_path / "input.xlsx"
    make_rooms(source)
    workbook = load_workbook(source)
    workbook.create_sheet("设备位置表")["A1"] = "Original business table"
    workbook.save(source)
    workbook.close()
    project = import_project(source, tmp_path / "reimport.sqlite")
    output = export_project_workbook(project, tmp_path / "export.xlsx", layout="source")
    workbook = load_workbook(output)
    try:
        assert workbook["设备位置表"]["A1"].value == "Original business table"
        assert workbook["设备位置表 (2)"]["A1"].value == "序号"
    finally:
        workbook.close()


@pytest.mark.parametrize("change", ["stale", "missing"])
def test_source_export_refuses_stale_or_missing_source(tmp_path: Path, change: str) -> None:
    source = tmp_path / "input.xlsx"
    project = make_rooms(source)
    if change == "missing":
        source.unlink()
    else:
        source.write_bytes(source.read_bytes() + b"external change")
    target = tmp_path / "existing.xlsx"
    target.write_bytes(b"Keep existing export")
    with pytest.raises((ValueError, FileNotFoundError)):
        export_project_workbook(project, target, overwrite=True, layout="source")
    assert target.read_bytes() == b"Keep existing export"


def test_source_export_external_change_before_publication_preserves_target(tmp_path: Path, monkeypatch) -> None:
    from racktool.core import export as exporter

    source = tmp_path / "input.xlsx"
    project = make_rooms(source)
    target = tmp_path / "existing.xlsx"
    target.write_bytes(b"Keep existing export")
    verify = exporter._validate_export

    def change_source(path, manifest):
        verify(path, manifest)
        source.write_bytes(source.read_bytes() + b"external change")

    monkeypatch.setattr(exporter, "_validate_export", change_source)
    with pytest.raises(ValueError, match="导出期间变化"):
        export_project_workbook(project, target, overwrite=True, layout="source")
    assert target.read_bytes() == b"Keep existing export"
    assert not list(tmp_path.glob(".existing.tmp-*"))


def test_source_export_cli(tmp_path: Path, capsys) -> None:
    source = tmp_path / "input.xlsx"
    make_rooms(source)
    original = sha256_file(source)
    assert main(["export", str(source.with_suffix(".sqlite")), str(tmp_path / "output.xlsx"), "--layout", "source"]) == 0
    assert '"status": "exported"' in capsys.readouterr().out
    assert sha256_file(source) == original


def test_source_export_rejects_uncommitted_model_and_unbound_draft(tmp_path: Path) -> None:
    project = make_rooms(tmp_path / "input.xlsx")
    changed = replace(project, devices=[replace(d, display_text="new text") for d in project.devices])
    with pytest.raises(ValueError, match="项目与源文件不一致"):
        export_project_workbook(changed, tmp_path / "invalid.xlsx", layout="source")
    draft = create_draft_project(tmp_path / "draft.sqlite")
    with pytest.raises(ValueError, match="已绑定"):
        export_project_workbook(draft, tmp_path / "invalid.xlsx", layout="source")
    assert not (tmp_path / "invalid.xlsx").exists()


def test_source_export_retains_original_style_extensions(tmp_path: Path) -> None:
    source = tmp_path / "input.xlsx"
    make_rooms(source)
    output = BytesIO()
    extension = b'<extLst><ext uri="fixture"><t:keep xmlns:t="urn:test">KEEP</t:keep></ext></extLst>'
    with ZipFile(source) as before, ZipFile(output, "w") as after:
        for name in before.namelist():
            data = before.read(name)
            if name == "xl/styles.xml":
                data = data.replace(b"</styleSheet>", extension + b"</styleSheet>")
            after.writestr(name, data)
    source.write_bytes(output.getvalue())
    project = import_project(source, tmp_path / "with-extension.sqlite")
    exported = export_project_workbook(project, tmp_path / "result.xlsx", layout="source")
    with ZipFile(exported) as archive:
        assert extension in archive.read("xl/styles.xml")


def test_source_export_after_safe_add_and_rename_has_current_positions(tmp_path: Path) -> None:
    source = tmp_path / "input.xlsx"
    project = make_rooms(source)
    database = source.with_suffix(".sqlite")
    rack = next(r for r in project.racks if r.rack_name == "B01")
    plan = plan_device_add(project, rack.rack_id, 1, 2, "NEW")
    assert not plan.conflicts
    result = commit_write_plan(source, database, plan)
    assert result.status == "applied" and result.project is not None
    new_id = plan.actions[0].device_id
    rename = plan_device_rename(result.project, new_id, "=NEW\nRENAMED")
    result = commit_write_plan(source, database, rename)
    assert result.status == "applied" and result.project is not None
    exported = export_project_workbook(result.project, tmp_path / "result.xlsx", layout="source")
    workbook = load_workbook(exported)
    try:
        assert workbook["Room-A"]["D4"].value == "=NEW\nRENAMED"
        assert workbook["Room-A"]["D4"].data_type == "s"
        rows = list(workbook["设备位置表"].iter_rows(min_row=2, values_only=True))
        row = next(row for row in rows if row[7] == new_id)
        assert row[1:7] == ("=NEW\nRENAMED", "B01", "1-2U", 2, "Room-A", "在位")
    finally:
        workbook.close()


def test_standard_export_rejects_direction_corruption_before_publication(tmp_path: Path, monkeypatch) -> None:
    from racktool.core import export as exporter

    project = make_rooms(tmp_path / "input.xlsx")
    build = exporter._build_workbook

    def flip_axis(project, layout):
        workbook, manifest = build(project, layout)
        sheet = workbook["Room-A"]
        title = next(cell for row in sheet.iter_rows() for cell in row if cell.value == "B01")
        for offset, unit in enumerate(range(12, 0, -1), 1):
            sheet.cell(title.row + offset, title.column, unit)
        return workbook, manifest

    monkeypatch.setattr(exporter, "_build_workbook", flip_axis)
    target = tmp_path / "existing.xlsx"
    target.write_bytes(b"Keep existing export")
    with pytest.raises(ValueError, match="U 轴方向"):
        export_project_workbook(project, target, overwrite=True)
    assert target.read_bytes() == b"Keep existing export"


def test_export_keeps_unplaced_device_with_blank_position(tmp_path: Path) -> None:
    project = make_rooms(tmp_path / "input.xlsx")
    project = replace(project, devices=[*project.devices, Device("UNPLACED", "Unplaced device")])
    output = export_project_workbook(project, tmp_path / "result.xlsx")
    workbook = load_workbook(output)
    try:
        row = next(row for row in workbook["设备位置表"].iter_rows(min_row=2, values_only=True)
                   if row[7] == "UNPLACED")
        assert row[2:7] == (None, None, None, None, "未放置")
    finally:
        workbook.close()
