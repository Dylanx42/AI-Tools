from pathlib import Path

import pytest
from openpyxl import load_workbook

from racktool.core import analyze_workbook, create_empty_rack_project, export_project_workbook
from racktool.core.project import project_error_conflicts


@pytest.mark.parametrize("height", [42, 47])
def test_empty_template_round_trips_both_rows_with_one_cell_per_u(
    tmp_path: Path, height: int,
) -> None:
    project = create_empty_rack_project(height_u=height)
    assert not project_error_conflicts(project)
    assert project.source_workbook is None
    output = export_project_workbook(project, tmp_path / "empty.xlsx")

    analysis = analyze_workbook(output)
    diagram = next(sheet for sheet in analysis.sheets if sheet.name == "机柜图")
    expected_names = {f"{prefix}{number:02d}" for prefix in "AB" for number in range(1, 19)}
    assert len(diagram.racks) == 36
    assert {rack.rack_name for rack in diagram.racks} == expected_names
    assert all(not sheet.devices and not sheet.placements and not sheet.issues
               for sheet in analysis.sheets)
    workbook = load_workbook(output)
    try:
        assert workbook.sheetnames == ["机柜图", "设备位置表"]
        sheet = workbook["机柜图"]
        for rack in diagram.racks:
            assert rack.height_u == height
            assert set(rack.u_to_row) == set(range(1, height + 1))
            assert len(rack.device_columns) == 1
            for u, row in rack.u_to_row.items():
                device = sheet.cell(row, rack.device_columns[0])
                assert device.value is None
                assert all(device.coordinate not in merged for merged in sheet.merged_cells.ranges)
                assert rack.left_axis_column is not None and rack.right_axis_column is not None
                assert sheet.cell(row, rack.left_axis_column).value == u
                assert sheet.cell(row, rack.right_axis_column).value == u
        assert len({rack.start_row for rack in diagram.racks if rack.rack_name.startswith("A")}) == 1
        assert len({rack.start_row for rack in diagram.racks if rack.rack_name.startswith("B")}) == 1
        assert len({rack.start_row for rack in diagram.racks}) == 2
    finally:
        workbook.close()


@pytest.mark.parametrize("count,height", [(0, 42), (100, 42), (18, 0), (18, 101)])
def test_empty_template_rejects_invalid_dimensions(count: int, height: int) -> None:
    with pytest.raises(ValueError, match="机柜数量"):
        create_empty_rack_project(count, height)
