from __future__ import annotations

from dataclasses import replace
from pathlib import Path
from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

from racktool.core.project import rescan_workbook
from racktool.core.service import commit_write_plan, import_project
from racktool.core.sync import (
    EditRequest,
    apply_writeback,
    plan_device_edit,
    plan_device_edits,
    plan_device_move,
)
from racktool.persistence import load_project


def _setup(tmp_path: Path, *, text: str = "原设备\n第二行", target_width: int = 1, source_width: int = 1):
    source, database = tmp_path / "rooms.xlsx", tmp_path / "project.sqlite"
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, ascending, width in (("一号机房", False, source_width), ("二号机房", True, target_width)):
        sheet = workbook.create_sheet(name)
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width+2)
        sheet.cell(1, 1, "A01" if not ascending else "B01")
        for row in range(2, 14):
            for column in (1, width+2):
                sheet.cell(row, column, row-1 if ascending else 14-row)
        for row in range(2, 14):
            for column in range(2, width+2):
                sheet.cell(row, column).border = Border(left=Side(style="thin"), right=Side(style="thin"))
        sheet.column_dimensions["B"].width = 24 if ascending else 16
        sheet.row_dimensions[5].height = 22
        sheet["Z1"] = "机房原备注"
    sheet = workbook["一号机房"]
    sheet.merge_cells(start_row=4, start_column=2, end_row=5, end_column=source_width + 1)  # descending 9-10U
    sheet["B4"] = text
    sheet["B4"].font = Font(name="Arial", bold=True, color="FF0000")
    sheet["B4"].fill = PatternFill("solid", fgColor="00FF00")
    sheet["B4"].alignment = Alignment(wrap_text=True, horizontal="center")
    sheet["B4"].border = Border(top=Side(style="medium"), bottom=Side(style="medium"),
                                 left=Side(style="medium"), right=Side(style="medium"))
    workbook.save(source)
    workbook.close()
    project = import_project(source, database)
    device = next(d for d in project.devices if d.display_text == text)
    old_rack = next(r for r in project.racks if r.source_sheet == "一号机房")
    rack = next(r for r in project.racks if r.source_sheet == "二号机房")
    return source, database, project, device, old_rack, rack


@pytest.mark.parametrize("units, expected_range", [((2, 4), "B3:B5"), ((6, 6), "B7"), ((11, 12), "B12:B13")])
def test_cross_room_resize_preserves_identity_text_and_common_style(tmp_path, units, expected_range):
    source, database, project, device, _, rack = _setup(tmp_path)
    original = source.read_bytes()
    placement_id = next(p.placement_id for p in project.placements if p.device_id == device.device_id)
    mapping_id = next(m.mapping_id for m in project.mappings if m.device_id == device.device_id)
    plan = plan_device_edit(project, device.device_id, rack.rack_id, *units)
    assert not plan.conflicts
    assert source.read_bytes() == original
    assert plan.actions[0].source_sheet_name == "一号机房"
    assert plan.actions[0].sheet_name == "二号机房"
    assert plan.actions[0].new_range == expected_range
    result = commit_write_plan(source, database, plan)
    assert result.status == "applied", result.errors
    assert Path(result.backup_path).read_bytes() == original
    persisted = load_project(database)
    assert {d.device_id for d in persisted.devices} == {device.device_id}
    placement = next(p for p in persisted.placements if p.device_id == device.device_id)
    mapping = next(m for m in persisted.mappings if m.device_id == device.device_id)
    assert (placement.placement_id, placement.rack_id, placement.start_u, placement.end_u) == (placement_id, rack.rack_id, *units)
    assert (mapping.mapping_id, mapping.sheet_name, mapping.source_range) == (mapping_id, "二号机房", expected_range)
    rescanned = rescan_workbook(source, persisted)
    assert rescanned.accepted and not rescanned.created_device_ids
    workbook = load_workbook(source)
    try:
        old, target = workbook["一号机房"], workbook["二号机房"]
        assert old["B4"].value is None and old["B5"].value is None
        assert old["B4"].font.bold and old["B4"].fill.fgColor.rgb == "0000FF00"
        assert old["B5"].border.bottom.style == "medium"
        assert "B4:B5" not in {str(m) for m in old.merged_cells.ranges}
        anchor = target[expected_range.split(":")[0]]
        assert anchor.value == "原设备\n第二行"
        assert anchor.font.bold and anchor.fill.fgColor.rgb == "0000FF00"
        assert anchor.alignment.wrap_text
        assert anchor.border.top.style == anchor.border.bottom.style == "medium"
        assert old["Z1"].value == target["Z1"].value == "机房原备注"
        assert target.column_dimensions["B"].width == 24 and target.row_dimensions[5].height == 22
    finally:
        workbook.close()


@pytest.mark.parametrize("units", [(9, 11), (10, 10), (8, 12)])
def test_resize_overlapping_own_source_and_then_shrink(tmp_path, units):
    source, database, project, device, rack, _ = _setup(tmp_path)
    plan = plan_device_edit(project, device.device_id, rack.rack_id, *units)
    assert not plan.conflicts
    result = commit_write_plan(source, database, plan)
    assert result.status == "applied", result.errors
    plan = plan_device_edit(result.project, device.device_id, rack.rack_id, 10, 10)
    result = commit_write_plan(source, database, plan)
    assert result.status == "applied", result.errors
    workbook = load_workbook(source)
    assert workbook["一号机房"]["B4"].value == "原设备\n第二行"
    assert all(workbook["一号机房"].cell(row, 2).value is None for row in range(2, 14) if row != 4)
    assert not any(str(m).startswith("B") for m in workbook["一号机房"].merged_cells.ranges)
    workbook.close()


@pytest.mark.parametrize("text,data_type", [("=SUM(1,2)", "f"), ("literal", "s")])
def test_resize_preserves_original_data_type_unless_text_explicitly_changed(tmp_path, text, data_type):
    source, database, project, device, _, rack = _setup(tmp_path, text=text)
    result = commit_write_plan(source, database, plan_device_edit(project, device.device_id, rack.rack_id, 2, 4))
    assert result.status == "applied", result.errors
    workbook = load_workbook(source, data_only=False)
    assert workbook["二号机房"]["B3"].data_type == data_type
    assert workbook["二号机房"]["B3"].value == text
    workbook.close()
    result = commit_write_plan(source, database, plan_device_edit(result.project, device.device_id, rack.rack_id, 5, 6, display_text="=literal\nchanged"))
    assert result.status == "applied", result.errors
    workbook = load_workbook(source, data_only=False)
    assert workbook["二号机房"]["B6"].data_type == "s"
    assert workbook["二号机房"]["B6"].value == "=literal\nchanged"
    workbook.close()


def test_plain_move_supports_cross_sheet_but_still_refuses_height_change(tmp_path):
    source, _, project, device, _, rack = _setup(tmp_path)
    changed = plan_device_move(project, device.device_id, rack.rack_id, 2, 4)
    assert any(c.code == "device-height-change" for c in changed.conflicts)
    plan = plan_device_move(project, device.device_id, rack.rack_id, 2, 3)
    assert not plan.conflicts
    assert apply_writeback(source, project, plan).status == "applied"


@pytest.mark.parametrize("side", ["source", "target"])
@pytest.mark.parametrize("kind", ["comment", "hyperlink", "validation", "defined-name", "merge", "conditional"])
def test_edit_refuses_unknown_metadata_on_both_sheets(tmp_path, side, kind):
    source, database, _, _, _, _ = _setup(tmp_path)
    workbook = load_workbook(source)
    sheet = workbook["一号机房" if side == "source" else "二号机房"]
    coordinate = "B4" if side == "source" else "B3"
    if kind == "comment":
        sheet[coordinate].comment = Comment("未知", "author")
    elif kind == "hyperlink":
        sheet[coordinate].hyperlink = "https://example.com"
    elif kind == "validation":
        validation = DataValidation(type="whole", formula1=1)
        validation.add(coordinate)
        sheet.add_data_validation(validation)
    elif kind == "defined-name":
        workbook.defined_names.add(DefinedName("BusinessCell", attr_text=f"'{sheet.title}'!${coordinate[0]}${coordinate[1:]}"))
    elif kind == "merge":
        if side == "source":
            sheet.unmerge_cells("B4:B5")
            sheet.merge_cells("B4:B6")
        else:
            sheet.merge_cells("B3:B4")
    else:
        from openpyxl.formatting.rule import ColorScaleRule
        sheet.conditional_formatting.add(coordinate, ColorScaleRule(start_type="min", start_color="FF0000", end_type="max", end_color="00FF00"))
    workbook.save(source)
    workbook.close()
    project = import_project(source, database)
    device = next(d for d in project.devices if d.display_text == "原设备\n第二行")
    rack = next(r for r in project.racks if r.source_sheet == "二号机房")
    original = source.read_bytes()
    plan = plan_device_edit(project, device.device_id, rack.rack_id, 2, 4)
    if kind == "merge" and side == "source":
        # This is now the device's actual three-U range, so it is a supported resize.
        assert not plan.conflicts
    else:
        assert plan.conflicts and not plan.actions
    assert source.read_bytes() == original


def test_edits_batch_rejects_whole_batch_and_swap_not_partial_writes(tmp_path):
    source, _, project, device, old, target = _setup(tmp_path)
    original = source.read_bytes()
    duplicate = plan_device_edits(project, [EditRequest(device.device_id, target.rack_id, 2, 3), EditRequest(device.device_id, old.rack_id, 3, 3)])
    assert duplicate.conflicts and not duplicate.actions
    invalid = plan_device_edits(project, [EditRequest(device.device_id, target.rack_id, 0, 999)])
    assert invalid.conflicts and not invalid.actions
    assert source.read_bytes() == original


def test_independent_cross_sheet_batch_is_one_safe_transaction(tmp_path):
    source, database, _, _, _, _ = _setup(tmp_path)
    workbook = load_workbook(source)
    workbook["二号机房"]["B9"] = "另一设备"
    workbook.save(source)
    workbook.close()
    project = import_project(source, database)
    first, second = project.devices
    if first.display_text == "另一设备":
        first, second = second, first
    old = next(r for r in project.racks if r.source_sheet == "一号机房")
    target = next(r for r in project.racks if r.source_sheet == "二号机房")
    plan = plan_device_edits(project, [EditRequest(first.device_id, target.rack_id, 2, 4), EditRequest(second.device_id, old.rack_id, 3, 4, "批量改名")])
    assert len(plan.actions) == 2 and not plan.conflicts
    result = commit_write_plan(source, database, plan)
    assert result.status == "applied", result.errors
    assert {d.device_id for d in result.project.devices} == {first.device_id, second.device_id}
    swap = plan_device_edits(result.project, [EditRequest(first.device_id, old.rack_id, 3, 4), EditRequest(second.device_id, target.rack_id, 2, 4)])
    assert swap.conflicts and not swap.actions


def test_forged_source_sheet_and_stale_source_are_refused(tmp_path):
    source, _, project, device, _, target = _setup(tmp_path)
    plan = plan_device_edit(project, device.device_id, target.rack_id, 2, 4)
    original = source.read_bytes()
    forged = replace(plan, actions=(replace(plan.actions[0], source_sheet_name="二号机房"),))
    assert apply_writeback(source, project, forged).status == "rejected"
    assert source.read_bytes() == original
    source.write_bytes(original + b"external")
    assert apply_writeback(source, project, plan).status == "rejected"


def test_persistence_failure_rolls_back_cross_room_resize(tmp_path, monkeypatch):
    source, database, project, device, _, target = _setup(tmp_path)
    original, original_database = source.read_bytes(), database.read_bytes()
    plan = plan_device_edit(project, device.device_id, target.rack_id, 2, 4)
    def fail(*_):
        raise RuntimeError("simulated database failure")
    monkeypatch.setattr("racktool.core.service.save_project", fail)
    result = commit_write_plan(source, database, plan)
    assert result.status == "failed"
    assert source.read_bytes() == original and database.read_bytes() == original_database


def test_worksheet_extensions_are_refused_before_writing(tmp_path):
    source, database, _, _, _, _ = _setup(tmp_path)
    with ZipFile(source) as package:
        parts = [(info, package.read(info.filename)) for info in package.infolist()]
    with ZipFile(source, "w") as package:
        for info, data in parts:
            if info.filename == "xl/worksheets/sheet1.xml":
                data = data.replace(b"</worksheet>", b'<extLst><ext uri="unknown" /></extLst></worksheet>')
            package.writestr(info, data)
    project = import_project(source, database)
    device = project.devices[0]
    target = next(r for r in project.racks if r.source_sheet == "二号机房")
    original = source.read_bytes()
    plan = plan_device_edit(project, device.device_id, target.rack_id, 2, 4)
    assert any(c.code == "unsupported-workbook-object" for c in plan.conflicts)
    assert source.read_bytes() == original


def test_cross_room_preserves_multicolumn_lane_and_refuses_narrower_target(tmp_path):
    source, database, project, device, _, target = _setup(tmp_path, source_width=2, target_width=2)
    plan = plan_device_edit(project, device.device_id, target.rack_id, 2, 4)
    assert not plan.conflicts
    assert plan.actions[0].new_range == "B3:C5"
    assert commit_write_plan(source, database, plan).status == "applied"
    next_dir = tmp_path / "narrow"
    next_dir.mkdir()
    source, _, project, device, _, target = _setup(next_dir, source_width=2, target_width=1)
    original = source.read_bytes()
    plan = plan_device_edit(project, device.device_id, target.rack_id, 2, 4)
    assert any(c.code == "device-width-or-layout-conflict" and "列数" in c.message for c in plan.conflicts)
    assert not plan.actions and source.read_bytes() == original


def test_single_lane_into_wider_rack_keeps_lane_and_column_dimensions(tmp_path):
    source, database, project, device, _, target = _setup(tmp_path, target_width=2)
    plan = plan_device_edit(project, device.device_id, target.rack_id, 2, 4)
    assert plan.actions[0].new_range == "B3:B5"
    result = commit_write_plan(source, database, plan)
    assert result.status == "applied", result.errors
    workbook = load_workbook(source)
    assert workbook["二号机房"]["C3"].value is None
    assert workbook["二号机房"].column_dimensions["B"].width == 24
    workbook.close()


def test_sqlite_only_state_change_rejects_prepared_plan(tmp_path):
    from racktool.persistence import save_project
    source, database, project, device, _, target = _setup(tmp_path)
    plan = plan_device_edit(project, device.device_id, target.rack_id, 2, 4)
    assert plan.project_revision
    save_project(database, replace(project, metadata={**project.metadata, "external": "changed"}))
    original = source.read_bytes()
    result = commit_write_plan(source, database, plan)
    assert result.status == "rejected"
    assert any(c.code == "stale-write-plan" for c in result.plan.conflicts)
    assert source.read_bytes() == original


def test_write_style_validation_failure_protects_source_and_database(tmp_path, monkeypatch):
    from racktool.core import sync
    source, database, project, device, _, target = _setup(tmp_path)
    original, original_db = source.read_bytes(), database.read_bytes()
    plan = plan_device_edit(project, device.device_id, target.rack_id, 2, 4)
    apply = sync._apply_action
    def corrupt(workbook, action):
        apply(workbook, action)
        workbook[action.sheet_name][action.new_range.split(":")[0]].font = Font(name="Wrong")
    monkeypatch.setattr(sync, "_apply_action", corrupt)
    result = commit_write_plan(source, database, plan)
    assert result.status == "failed" and any("style" in error for error in result.errors)
    assert source.read_bytes() == original and database.read_bytes() == original_db


@pytest.mark.parametrize("units", [(0, 1), (-1, 3), (2, 10**100)])
def test_invalid_height_bounds_refused_without_traversing_huge_u_range(tmp_path, units):
    _source, _, project, device, _, target = _setup(tmp_path)
    plan = plan_device_edit(project, device.device_id, target.rack_id, *units)
    assert any(c.code == "u-out-of-bounds" for c in plan.conflicts)
    assert not plan.actions


def test_expected_sha256_import_uses_snapshot_and_binds_real_source(tmp_path):
    from racktool.core.identity import sha256_file
    source, _, _, _, _, _ = _setup(tmp_path)
    database = tmp_path / "new.sqlite"
    digest = sha256_file(source)
    project = import_project(source, database, expected_sha256=digest)
    assert project.source_workbook == str(source.resolve()) and project.workbook_fingerprint == digest
    assert not list(tmp_path.glob(".*.tmp-*"))
    with pytest.raises(ValueError, match="SHA256"):
        import_project(source, tmp_path / "wrong.sqlite", expected_sha256="0" * 64)
    assert not (tmp_path / "wrong.sqlite").exists()


@pytest.mark.parametrize("existing", [False, True])
def test_expected_sha256_import_source_race_restores_database(tmp_path, monkeypatch, existing):
    from racktool.core import service
    from racktool.core.identity import sha256_file
    source, previous_db, _, _, _, _ = _setup(tmp_path)
    database = previous_db if existing else tmp_path / "new.sqlite"
    original_database = database.read_bytes() if existing else None
    digest = sha256_file(source)
    save = service.save_project
    def external_change(path, project):
        save(path, project)
        source.write_bytes(source.read_bytes() + b"external-update")
    monkeypatch.setattr(service, "save_project", external_change)
    with pytest.raises(ValueError, match="changed during"):
        import_project(source, database, expected_sha256=digest)
    if existing:
        assert database.read_bytes() == original_database
    else:
        assert not database.exists()
    assert source.read_bytes().endswith(b"external-update")


@pytest.mark.parametrize("guarded", [False, True])
def test_import_database_cannot_overwrite_source_workbook(tmp_path, guarded):
    from racktool.core.identity import sha256_file
    source, _, _, _, _, _ = _setup(tmp_path)
    original = source.read_bytes()
    with pytest.raises(ValueError, match="must not overwrite"):
        import_project(source, source, expected_sha256=sha256_file(source) if guarded else None)
    assert source.read_bytes() == original


def test_sync_patches_only_cells_and_merges_preserving_workbook_style_extensions(tmp_path):
    from racktool.core.sync_ooxml import _spans
    source, database, _, _, _, _ = _setup(tmp_path)
    workbook_extension = b'<extLst><ext uri="racktool-opaque-workbook"><keep:opaque value="unchanged-workbook"/></ext></extLst>'
    styles_extension = b'<extLst><ext uri="racktool-opaque-styles"><keep:opaque value="unchanged-styles"/></ext></extLst>'
    with ZipFile(source) as package:
        parts = [(info, package.read(info)) for info in package.infolist()]
    with ZipFile(source, "w") as package:
        for info, data in parts:
            if info.filename == "xl/workbook.xml":
                data = data.replace(b'<workbook ', b'<workbook xmlns:keep="urn:racktool:opaque" ', 1)
                data = data.replace(b'</workbook>', workbook_extension + b'</workbook>')
            elif info.filename == "xl/styles.xml":
                data = data.replace(b'<styleSheet ', b'<styleSheet xmlns:keep="urn:racktool:opaque" ', 1)
                data = data.replace(b'</styleSheet>', styles_extension + b'</styleSheet>')
            package.writestr(info, data)
        package.writestr("docProps/preserved-opaque.xml", b'<opaque xmlns="urn:fixture" value="keep-this-byte-for-byte"/>')
    project = import_project(source, database)
    device = project.devices[0]
    target = next(r for r in project.racks if r.source_sheet == "二号机房")
    with ZipFile(source) as package:
        before = {name: package.read(name) for name in package.namelist()}
    plan = plan_device_edit(project, device.device_id, target.rack_id, 2, 4, display_text="新设备\n下一行")
    assert not plan.conflicts
    result = commit_write_plan(source, database, plan)
    assert result.status == "applied", result.errors
    with ZipFile(source) as package:
        after = {name: package.read(name) for name in package.namelist()}
    assert before.keys() == after.keys()
    changed_parts = {"xl/styles.xml", "xl/worksheets/sheet1.xml", "xl/worksheets/sheet2.xml"}
    assert all(raw == after[name] for name, raw in before.items() if name not in changed_parts)
    assert workbook_extension in after["xl/workbook.xml"]
    assert styles_extension in after["xl/styles.xml"]
    for tag in ("font", "fill", "border", "xf"):
        old_elements = [match.group() for match in _spans(before["xl/styles.xml"], tag)]
        new_elements = [match.group() for match in _spans(after["xl/styles.xml"], tag)]
        # cellStyleXfs and cellXfs XF lists have separate preserved index prefixes.
        if tag == "xf":
            for container in ("cellStyleXfs", "cellXfs"):
                old_elements = [match.group() for match in _spans(_spans(before["xl/styles.xml"], container)[0].group(), tag)]
                new_elements = [match.group() for match in _spans(_spans(after["xl/styles.xml"], container)[0].group(), tag)]
                assert new_elements[:len(old_elements)] == old_elements
        else:
            assert new_elements[:len(old_elements)] == old_elements
    for name in ("xl/worksheets/sheet1.xml", "xl/worksheets/sheet2.xml"):
        original_note = next(match.group() for match in _spans(before[name], "c") if b'r="Z1"' in match.group())
        assert original_note in after[name]
        original_header = next(match.group() for match in _spans(before[name], "row") if b'r="1"' in match.group()[:25])
        assert original_header in after[name]
    rescanned = rescan_workbook(source, result.project)
    assert rescanned.accepted and not rescanned.created_device_ids


def test_sync_rejects_generated_style_prefix_reindex_without_touching_source(tmp_path, monkeypatch):
    from racktool.core import sync
    source, database, project, device, _, target = _setup(tmp_path)
    original, original_db = source.read_bytes(), database.read_bytes()
    plan = plan_device_edit(project, device.device_id, target.rack_id, 2, 4)
    apply = sync._apply_action
    def corrupt(workbook, action):
        apply(workbook, action)
        workbook._fonts[0].name = "Cannot replace original indexed font"
    monkeypatch.setattr(sync, "_apply_action", corrupt)
    result = commit_write_plan(source, database, plan)
    assert result.status == "failed" and any("original fonts index" in error for error in result.errors)
    assert source.read_bytes() == original and database.read_bytes() == original_db


def test_sync_can_add_styles_to_legal_workbook_without_styles_part(tmp_path):
    import runpy
    fixture = runpy.run_path(str(Path(__file__).resolve().parents[2] / "scripts" / "verify_skill_workflows.py"))
    source, database = tmp_path / "no-styles.xlsx", tmp_path / "project.sqlite"
    fixture["make_workbook"](source)
    with ZipFile(source) as package:
        original_workbook = package.read("xl/workbook.xml")
        original_sheets = {name: package.read(name) for name in package.namelist() if name.startswith("xl/worksheets/")}
        assert "xl/styles.xml" not in package.namelist()
    project = import_project(source, database)
    target = next(r for r in project.racks if r.rack_name == "B01")
    source_rack = next(r for r in project.racks if r.rack_name == "A01")
    device_id = next(p.device_id for p in project.placements if p.rack_id == source_rack.rack_id)
    plan = plan_device_edit(project, device_id, target.rack_id, 2, 4)
    assert not plan.conflicts
    result = commit_write_plan(source, database, plan)
    assert result.status == "applied", result.errors
    with ZipFile(source) as package:
        assert package.read("xl/workbook.xml") == original_workbook
        assert "xl/styles.xml" in package.namelist()
        assert b"/styles" in package.read("xl/_rels/workbook.xml.rels")
        assert b"styles+xml" in package.read("[Content_Types].xml")
        assert all(package.read(name) != data for name, data in original_sheets.items())
    workbook = load_workbook(source)
    assert workbook["Room-A"]["D5"].value == "SYNTHETIC-SERVER"
    assert "D5:D7" in {str(merged) for merged in workbook["Room-A"].merged_cells.ranges}
    workbook.close()


def test_base_style_false_apply_switches_remain_raw_during_cross_room_edit(tmp_path):
    import re

    from racktool.core.sync_ooxml import _spans
    source, database, _, _, _, _ = _setup(tmp_path)
    with ZipFile(source) as package:
        parts = [(info, package.read(info)) for info in package.infolist()]
    with ZipFile(source, "w") as package:
        for info, data in parts:
            if info.filename == "xl/styles.xml":
                container = _spans(data, "cellStyleXfs")[0]
                replacement = re.sub(rb"<xf\b", b'<xf applyFont="0" applyFill="0" applyBorder="0" applyNumberFormat="0"', container.group(), count=1)
                data = data[:container.start()] + replacement + data[container.end():]
                original_base = replacement
            package.writestr(info, data)
    project = import_project(source, database)
    device = project.devices[0]
    target = next(r for r in project.racks if r.source_sheet == "二号机房")
    plan = plan_device_edit(project, device.device_id, target.rack_id, 2, 4)
    result = commit_write_plan(source, database, plan)
    assert result.status == "applied", result.errors
    with ZipFile(source) as package:
        assert _spans(package.read("xl/styles.xml"), "cellStyleXfs")[0].group() == original_base
    assert rescan_workbook(source, result.project).accepted
