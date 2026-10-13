from __future__ import annotations

import json
from dataclasses import replace
from pathlib import Path

import pytest
from openpyxl import load_workbook

from racktool.cli.main import main
from racktool.core import (
    add_draft_device,
    add_draft_rack,
    analyze_workbook,
    create_draft_project,
    edit_draft_device,
    export_project_workbook,
)
from racktool.core import editing as editing_module
from racktool.core.project import project_error_conflicts
from racktool.core.service import import_project, load_project_state


def test_new_projects_have_independent_ids_and_never_replace_existing_files(tmp_path: Path) -> None:
    database = tmp_path / "first.sqlite"
    first = create_draft_project(database, count=1, height_u=12)
    second = create_draft_project(tmp_path / "second.sqlite", count=1, height_u=12)
    assert first.source_workbook is None
    assert first.project_id != second.project_id
    assert {rack.rack_id for rack in first.racks}.isdisjoint(rack.rack_id for rack in second.racks)
    original = database.read_bytes()
    with pytest.raises(FileExistsError):
        create_draft_project(database)
    assert database.read_bytes() == original


def test_cli_create_add_edit_and_export_preserves_id_and_round_trips_positions(
    tmp_path: Path, capsys: pytest.CaptureFixture[str],
) -> None:
    database = tmp_path / "project.sqlite"
    assert main(["project", "create", str(database), "--count", "1", "--height", "12"]) == 0
    initial = json.loads(capsys.readouterr().out)
    assert len(initial["racks"]) == 2
    assert main(["project", "add-rack", str(database), "C01", "--row", "3", "--height", "47"]) == 0
    capsys.readouterr()
    assert main(["project", "add-rack", str(database), "A02", "--height", "47"]) == 0
    capsys.readouterr()
    assert main(["project", "add-device", str(database), "--rack", "A01", "--start-u", "11",
                 "--end-u", "12", "--text", "SWITCH\nMODEL"]) == 0
    added = json.loads(capsys.readouterr().out)
    device_id = added["devices"][0]["device_id"]
    placement_id = added["placements"][0]["placement_id"]
    mapping_id = next(mapping["mapping_id"] for mapping in added["mappings"]
                      if mapping["device_id"] == device_id)
    assert main(["project", "edit-device", str(database), device_id, "--rack", "C01",
                 "--start-u", "30", "--end-u", "31", "--text", "=DISPLAY TEXT\nNEW MODEL"]) == 0
    edited = json.loads(capsys.readouterr().out)
    assert edited["devices"][0]["device_id"] == device_id
    assert edited["placements"][0]["placement_id"] == placement_id
    assert next(mapping["mapping_id"] for mapping in edited["mappings"]
                if mapping["device_id"] == device_id) == mapping_id
    assert {rack["rack_id"] for rack in initial["racks"]}.issubset(
        rack["rack_id"] for rack in edited["racks"]
    )
    assert main(["project", "show", str(database)]) == 0
    assert json.loads(capsys.readouterr().out) == edited

    output = tmp_path / "draft-export.xlsx"
    assert main(["export", str(database), str(output)]) == 0
    capsys.readouterr()
    assert not project_error_conflicts(load_project_state(database))
    analysis = analyze_workbook(output)
    sheet = next(item for item in analysis.sheets if item.name == "机柜图")
    assert {(rack.rack_name, rack.height_u) for rack in sheet.racks} == {
        ("A01", 12), ("A02", 47), ("B01", 12), ("C01", 47),
    }
    rack_names = {rack.candidate_id: rack.rack_name for rack in sheet.racks}
    assert len(sheet.devices) == len(sheet.placements) == 1
    assert sheet.devices[0].display_text == "=DISPLAY TEXT\nNEW MODEL"
    placement = sheet.placements[0]
    assert (rack_names[placement.rack_candidate_id], placement.start_u, placement.end_u) == ("C01", 30, 31)
    workbook = load_workbook(output)
    try:
        assert all(cell.data_type != "f" for sheet in workbook.worksheets
                   for row in sheet.iter_rows() for cell in row)
    finally:
        workbook.close()


def test_editing_rejects_overlap_without_changing_saved_project(tmp_path: Path) -> None:
    database = tmp_path / "project.sqlite"
    create_draft_project(database, count=1, height_u=12)
    first = add_draft_device(database, "A01", 11, 12, "SAME DISPLAY TEXT")
    first_id = first.devices[0].device_id
    second = add_draft_device(database, "B01", 11, 12, "SAME DISPLAY TEXT")
    second_id = next(device.device_id for device in second.devices if device.device_id != first_id)
    assert first_id != second_id
    original = database.read_bytes()
    with pytest.raises(ValueError, match="overlapping-active-placements"):
        add_draft_device(database, "A01", 10, 11, "CONFLICT")
    with pytest.raises(ValueError, match="overlapping-active-placements"):
        edit_draft_device(database, second_id, rack_reference="A01")
    assert database.read_bytes() == original


@pytest.mark.parametrize("start,end", [(0, 1), (1, 0), (12, 13), (-1, 1)])
def test_invalid_u_ranges_are_rejected_for_add_and_edit_without_writes(
    tmp_path: Path, start: int, end: int,
) -> None:
    database = tmp_path / "project.sqlite"
    create_draft_project(database, count=1, height_u=12)
    project = add_draft_device(database, "A01", 1, 1, "DEVICE")
    original = database.read_bytes()
    with pytest.raises(ValueError, match="超出机柜"):
        add_draft_device(database, "B01", start, end, "DEVICE")
    with pytest.raises(ValueError, match="超出机柜"):
        edit_draft_device(database, project.devices[0].device_id, start_u=start, end_u=end)
    assert database.read_bytes() == original


def test_duplicate_rack_missing_target_and_empty_text_are_rejected(tmp_path: Path) -> None:
    database = tmp_path / "project.sqlite"
    create_draft_project(database, count=1, height_u=12)
    project = add_draft_device(database, "A01", 1, 1, "DEVICE")
    original = database.read_bytes()
    with pytest.raises(ValueError, match="名称已存在"):
        add_draft_rack(database, "a01")
    with pytest.raises(ValueError, match="排号"):
        add_draft_rack(database, "NEW", row=99)
    with pytest.raises(ValueError, match="目标机柜"):
        add_draft_device(database, "MISSING", 1, 1, "DEVICE")
    with pytest.raises(ValueError, match="不能为空"):
        edit_draft_device(database, project.devices[0].device_id, display_text=" \n")
    with pytest.raises(ValueError, match="同时提供"):
        edit_draft_device(database, project.devices[0].device_id, start_u=2)
    assert database.read_bytes() == original


def test_imported_project_edit_is_refused_without_touching_source_or_database(tmp_path: Path) -> None:
    draft = tmp_path / "draft.sqlite"
    create_draft_project(draft, count=1, height_u=12)
    project = add_draft_device(draft, "A01", 1, 1, "DEVICE")
    source = export_project_workbook(project, tmp_path / "source.xlsx")
    imported_database = tmp_path / "imported.sqlite"
    imported = import_project(source, imported_database)
    source_before, database_before = source.read_bytes(), imported_database.read_bytes()
    with pytest.raises(ValueError, match="已绑定 Excel"):
        add_draft_rack(imported_database, "NEW")
    with pytest.raises(ValueError, match="已绑定 Excel"):
        add_draft_device(imported_database, "A01", 2, 2, "NEW DEVICE")
    with pytest.raises(ValueError, match="已绑定 Excel"):
        edit_draft_device(imported_database, imported.devices[0].device_id, display_text="EDITED")
    assert source.read_bytes() == source_before
    assert imported_database.read_bytes() == database_before


def test_save_failure_and_external_revision_change_preserve_saved_state(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    database = tmp_path / "project.sqlite"
    create_draft_project(database, count=1, height_u=12)
    project = add_draft_device(database, "A01", 1, 1, "BEFORE")
    original = database.read_bytes()
    real_save = editing_module.save_project

    def fail_save(path: Path, _project: object) -> None:
        path.write_bytes(b"failed temporary save")
        raise OSError("forced save failure")

    monkeypatch.setattr(editing_module, "save_project", fail_save)
    with pytest.raises(OSError, match="forced save failure"):
        edit_draft_device(database, project.devices[0].device_id, display_text="NOT SAVED")
    assert database.read_bytes() == original
    assert not list(tmp_path.glob(".project.sqlite.draft-*"))

    def external_write(path: Path, updated):
        real_save(path, updated)
        real_save(database, replace(project, devices=[replace(project.devices[0], display_text="EXTERNAL")]))

    monkeypatch.setattr(editing_module, "save_project", external_write)
    with pytest.raises(ValueError, match="其他程序修改"):
        edit_draft_device(database, project.devices[0].device_id, display_text="NOT SAVED")
    assert load_project_state(database).devices[0].display_text == "EXTERNAL"
    assert not list(tmp_path.glob(".project.sqlite.draft-*"))
