from pathlib import Path

import pytest
from openpyxl import load_workbook

from racktool.core import (
    add_draft_device,
    add_draft_rack,
    create_draft_project,
    edit_draft_device,
    export_project_workbook,
)
from tests.layout_corpus import INVALID_TEXT


@pytest.mark.parametrize("text", INVALID_TEXT, ids=["nul", "vertical-tab", "xml-noncharacter", "lone-surrogate"])
def test_draft_rejects_unexportable_text_before_database_changes(tmp_path: Path, text: str) -> None:
    database = tmp_path / "draft.sqlite"
    create_draft_project(database, count=1, height_u=12)
    project = add_draft_device(database, "A01", 1, 1, "SYNTHETIC")
    before = database.read_bytes()
    with pytest.raises(ValueError, match="字符"):
        add_draft_rack(database, text)
    with pytest.raises(ValueError, match="字符"):
        add_draft_device(database, "B01", 2, 2, text)
    with pytest.raises(ValueError, match="字符"):
        edit_draft_device(database, project.devices[0].device_id, display_text=text)
    assert database.read_bytes() == before


@pytest.mark.parametrize("text", ["中文设备💻\n型号&<机柜>", "  空格 / TAB\t设备  ", "=SUM(1,2)\n备注"])
def test_valid_unicode_symbols_and_literal_formula_text_round_trip(tmp_path: Path, text: str) -> None:
    database = tmp_path / "draft.sqlite"
    create_draft_project(database, count=1, height_u=12)
    project = add_draft_device(database, "A01", 1, 2, text)
    output = export_project_workbook(project, tmp_path / "output.xlsx")
    workbook = load_workbook(output)
    try:
        assert any(cell.value == text and cell.data_type == "s" for sheet in workbook.worksheets
                   for row in sheet.iter_rows() for cell in row)
    finally:
        workbook.close()
