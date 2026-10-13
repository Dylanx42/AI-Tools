from pathlib import Path

import pytest
from openpyxl import load_workbook

from racktool.core.service import commit_write_plan, import_project
from racktool.core.sync import plan_device_move
from tests.layout_corpus import LayoutCase, build_layout


@pytest.mark.parametrize("data_type", ["s", "f"], ids=["literal-equals-text", "existing-formula"])
def test_move_preserves_original_formula_or_literal_cell_type(tmp_path: Path, data_type: str) -> None:
    facts = build_layout(LayoutCase("paired-one", "ascending", 12, "fullwidth-u"), tmp_path)
    source = tmp_path / "source.xlsx"
    original = facts["devices"][1]
    anchor = original["source_range"].split(":")[0]
    workbook = load_workbook(source)
    sheet = workbook["Racks"]
    sheet[anchor] = "=1+2" if data_type == "f" else "=SYNTHETIC TEXT\nMODEL"
    sheet[anchor].data_type = data_type
    workbook.save(source)
    workbook.close()
    database = tmp_path / "project.sqlite"
    project = import_project(source, database)
    expected_text = "=1+2" if data_type == "f" else "=SYNTHETIC TEXT\nMODEL"
    device = next(device for device in project.devices if device.display_text == expected_text)
    rack_id = next(rack.rack_id for rack in project.racks if rack.rack_name == original["rack"])
    plan = plan_device_move(project, device.device_id, rack_id, 6, 6)
    assert not plan.conflicts
    result = commit_write_plan(source, database, plan)
    assert result.status == "applied", result.errors
    workbook = load_workbook(source)
    try:
        assert workbook["Racks"][plan.actions[0].new_range].data_type == data_type
        assert workbook["Racks"][plan.actions[0].new_range].value == device.display_text
    finally:
        workbook.close()
