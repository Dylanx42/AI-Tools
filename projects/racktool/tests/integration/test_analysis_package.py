from __future__ import annotations

import hashlib
import json
from pathlib import Path
from zipfile import ZipFile

import pytest
from openpyxl import Workbook

from racktool.cli.main import main
from racktool.core import analysis_package as package_module
from racktool.core import export_analysis_package


def _make_layout(path: Path, text: str = "SYNTHETIC-DEVICE\nMODEL") -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Synthetic rack"
    sheet.merge_cells("A1:C1")
    sheet["A1"] = "SYNTHETIC-RACK"
    for row, u in enumerate(range(12, 0, -1), start=2):
        sheet.cell(row, 1, u)
        sheet.cell(row, 3, u)
    sheet.merge_cells("B2:B3")
    sheet["B2"] = text
    sheet.row_dimensions[2].height = 25
    sheet.column_dimensions["B"].width = 30
    hidden = workbook.create_sheet("Synthetic hidden")
    hidden.sheet_state = "hidden"
    hidden["A1"] = "SYNTHETIC-HIDDEN-TEXT"
    workbook.create_sheet("Synthetic empty")
    workbook.save(path)
    workbook.close()


def test_package_contains_traceable_structure_and_review_without_source(tmp_path: Path) -> None:
    source = tmp_path / "layout.xlsx"
    _make_layout(source)
    original = source.read_bytes()
    output = export_analysis_package(source, tmp_path / "layout-analysis")

    assert output.suffix == ".zip"
    assert source.read_bytes() == original
    with ZipFile(output) as archive:
        assert set(archive.namelist()) == {
            "manifest.json", "workbook_structure.json", "analysis.json", "review.json", "README.txt",
        }
        manifest = json.loads(archive.read("manifest.json"))
        assert manifest["schema_version"] == 1
        assert manifest["source"] == {
            "name": "layout.xlsx", "sha256": hashlib.sha256(original).hexdigest(), "included": False,
        }
        assert manifest["privacy"]["redacted"] is False
        for name, entry in manifest["files"].items():
            contents = archive.read(name)
            assert entry["sha256"] == hashlib.sha256(contents).hexdigest()
            assert entry["bytes"] == len(contents)
            assert str(tmp_path).encode() not in contents
        structure = json.loads(archive.read("workbook_structure.json"))
        assert structure["sheets"][0]["merged_ranges"] == ["A1:C1", "B2:B3"]
        assert structure["sheets"][0]["row_heights"]["2"] == 25
        assert structure["sheets"][0]["column_widths"]["B"] == 30
        assert structure["sheets"][1]["cells"][0]["value"] == "SYNTHETIC-HIDDEN-TEXT"
        analysis = json.loads(archive.read("analysis.json"))
        assert analysis["sheets"][0]["placements"][0]["height_u"] == 2
        assert analysis["sheets"][0]["devices"][0]["evidence"]
        review = json.loads(archive.read("review.json"))
        assert review["human_confirmation_required"] is True
        assert review["sheets"][0]["racks"] == 1
        assert review["sheets"][1]["state"] == "hidden"
        assert review["sheets"][1]["review_reasons"] == ["nonempty-sheet-without-rack-candidates"]
        assert review["sheets"][2]["review_reasons"] == []


@pytest.mark.parametrize("include_source", [False, True])
def test_package_cli_is_deterministic_and_source_copy_is_opt_in(
    tmp_path: Path, capsys: pytest.CaptureFixture[str], include_source: bool,
) -> None:
    source = tmp_path / "layout.xlsx"
    _make_layout(source)
    first, second = tmp_path / "first.zip", tmp_path / "second.zip"
    options = ["--include-source"] if include_source else []
    for output in (first, second):
        assert main(["analysis-package", str(source), str(output), *options]) == 0
        result = json.loads(capsys.readouterr().out)
        assert result["source_included"] is include_source
    assert first.read_bytes() == second.read_bytes()
    with ZipFile(first) as archive:
        assert ("source.xlsx" in archive.namelist()) is include_source
        if include_source:
            assert archive.read("source.xlsx") == source.read_bytes()


def test_package_refuses_overwrite_and_keeps_target_after_validation_failure(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    source = tmp_path / "layout.xlsx"
    _make_layout(source)
    target = tmp_path / "existing.zip"
    target.write_bytes(b"keep existing package")
    with pytest.raises(FileExistsError):
        export_analysis_package(source, target)
    with pytest.raises(ValueError, match="覆盖源工作簿"):
        export_analysis_package(source, source, overwrite=True)

    def fail(*_args: object) -> None:
        raise ValueError("forced package validation failure")

    monkeypatch.setattr(package_module, "_validate_package", fail)
    with pytest.raises(ValueError, match="forced package validation failure"):
        export_analysis_package(source, target, overwrite=True)
    assert target.read_bytes() == b"keep existing package"
    assert not list(tmp_path.glob(".existing.tmp-*"))


def test_package_snapshot_stays_coherent_if_source_is_changed_externally(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    source = tmp_path / "layout.xlsx"
    _make_layout(source, "BEFORE")
    original = source.read_bytes()
    scan = package_module.scan_workbook

    def externally_change_source(snapshot: Path):
        result = scan(snapshot)
        _make_layout(source, "AFTER")
        return result

    monkeypatch.setattr(package_module, "scan_workbook", externally_change_source)
    output = export_analysis_package(source, tmp_path / "snapshot.zip", include_source=True)
    with ZipFile(output) as archive:
        assert archive.read("source.xlsx") == original
        analysis = json.loads(archive.read("analysis.json"))
        assert analysis["sheets"][0]["devices"][0]["display_text"] == "BEFORE"
        assert json.loads(archive.read("manifest.json"))["source"]["sha256"] == hashlib.sha256(original).hexdigest()


def test_low_confidence_review_hint_does_not_change_parser_candidates(tmp_path: Path) -> None:
    source = tmp_path / "layout.xlsx"
    _make_layout(source)
    first = export_analysis_package(source, tmp_path / "default.zip")
    second = export_analysis_package(source, tmp_path / "review-all.zip", review_threshold=1)
    with ZipFile(first) as default, ZipFile(second) as reviewed:
        assert default.read("analysis.json") == reviewed.read("analysis.json")
        summary = json.loads(reviewed.read("review.json"))
        assert "low-confidence-candidates" in summary["sheets"][0]["review_reasons"]
        assert summary["sheets"][0]["low_confidence_candidates"]
