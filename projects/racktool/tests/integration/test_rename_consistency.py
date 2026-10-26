"""Synthetic rename regressions; no private workbook or Office session is used."""
from __future__ import annotations

import json
from pathlib import Path
from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

from racktool.cli.main import main
from racktool.core.conversation import apply_conversation_edit, preview_conversation_edit
from racktool.core.export import export_project_workbook
from racktool.core.position_table import apply_position_table, preview_position_table
from racktool.core.service import import_project, load_project_state
from racktool.core.sync import EditRequest, plan_device_edit, plan_device_edits, plan_device_rename


def _bound(tmp_path: Path, *, chart_sheet: str | None = "Other", metadata: str = ""):
    source, database = tmp_path / "synthetic.xlsx", tmp_path / "project.sqlite"
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, rack in (("Room", "A01"), ("Other", "B01")):
        sheet = workbook.create_sheet(name)
        sheet.merge_cells("A1:C1")
        sheet["A1"] = rack
        for row in range(2, 14):
            sheet.cell(row, 1, 14 - row)
            sheet.cell(row, 3, 14 - row)
        sheet["Z1"] = "Unrelated note"
    sheet = workbook["Room"]
    sheet.merge_cells("B4:B5")  # 9-10U
    sheet["B4"] = "Synthetic device"
    sheet["B4"].font = Font(bold=True, color="FF0000")
    sheet["B4"].fill = PatternFill("solid", fgColor="00FF00")
    sheet["B4"].alignment = Alignment(wrap_text=True)
    sheet.column_dimensions["B"].width = 24
    sheet.row_dimensions[4].height = 22
    if metadata == "comment":
        sheet["B4"].comment = Comment("Protected note", "Synthetic")
    elif metadata == "hyperlink":
        sheet["B4"].hyperlink = "https://example.com/device"
    elif metadata == "validation":
        validation = DataValidation(type="list", formula1='"a,b"')
        validation.add("B4:B5")
        sheet.add_data_validation(validation)
    elif metadata == "defined-name":
        workbook.defined_names.add(DefinedName("Protected", attr_text="'Room'!$B$4:$B$5"))
    if chart_sheet is not None:
        chart = BarChart()
        chart.add_data(Reference(workbook[chart_sheet], min_col=1, min_row=2, max_row=4))
        workbook[chart_sheet].add_chart(chart, "Z20")
    workbook.save(source)
    workbook.close()
    project = import_project(source, database)
    assert len(project.devices) == 1
    return source, database, project


def _request(project, **changes):
    return {"operation": "edit", "selector": {"device_id": project.devices[0].device_id},
            "text": "=renamed\nSynthetic", "write_scope": "working-copy", **changes}


def _cli(capsys, arguments, *, expected=0):
    code = main([str(argument) for argument in arguments])
    output = json.loads(capsys.readouterr().out)
    assert code == expected, output
    return output


def _preview_cli(tmp_path, database, request, capsys, *, expected=0):
    request_path, plan_path = tmp_path / "request.json", tmp_path / "plan.json"
    request_path.write_text(json.dumps(request, ensure_ascii=False), encoding="utf-8")
    return _cli(capsys, ["conversation", "preview", database, "--request", request_path,
                         "--plan", plan_path], expected=expected)


@pytest.mark.parametrize("chart_sheet", ["Room", "Other"])
@pytest.mark.parametrize("entry", ["conversation-implicit", "conversation-explicit", "sync-edit", "position-table"])
def test_rename_entrances_preserve_identity_and_unrelated_package(tmp_path, capsys, entry, chart_sheet):
    source, database, project = _bound(tmp_path, chart_sheet=chart_sheet)
    device, placement = project.devices[0], project.placements[0]
    mapping = next(item for item in project.mappings if item.device_id == device.device_id)
    before = source.read_bytes(), database.read_bytes()
    with ZipFile(source) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    assert any(name.startswith("xl/charts/") for name in parts)
    assert any(name.startswith("xl/drawings/") for name in parts)
    text = _request(project)["text"]
    expected = plan_device_rename(project, device.device_id, text)
    assert not expected.conflicts
    # Reversed endpoints describe the same occupied U range.
    edit = plan_device_edit(project, device.device_id, placement.rack_id, 10, 9, display_text=text)
    assert edit == expected
    assert plan_device_edits(project, [EditRequest(device.device_id, placement.rack_id, 9, 10, text)]) == expected
    if entry.startswith("conversation"):
        request = _request(project)
        if entry == "conversation-explicit":
            request.update(target={"rack_id": placement.rack_id}, start_u=10, end_u=9)
        preview = _preview_cli(tmp_path, database, request, capsys)
        assert preview["summary"]["write_plan"] == expected.to_dict()
        assert (source.read_bytes(), database.read_bytes()) == before
        result = _cli(capsys, ["conversation", "apply", database, tmp_path / "plan.json",
                               "--token", preview["token"], "--commit"])
    elif entry == "sync-edit":
        args = ["sync", "edit", source, database, device.device_id, placement.rack_id,
                "9", "10", "--text", text]
        assert _cli(capsys, args) == expected.to_dict()
        assert (source.read_bytes(), database.read_bytes()) == before
        result = _cli(capsys, [*args, "--commit"])
    else:
        positions = export_project_workbook(project, tmp_path / "positions.xlsx")
        workbook = load_workbook(positions)
        sheet = next(sheet for sheet in workbook.worksheets if sheet["H1"].value == "设备 ID")
        sheet["B2"] = text
        sheet["B2"].data_type = "s"
        workbook.save(positions)
        workbook.close()
        preview = preview_position_table(database, positions)
        assert preview.accepted, preview.to_dict()
        assert (source.read_bytes(), database.read_bytes()) == before
        result = apply_position_table(database, positions, expected_preview_token=preview.preview_token).to_dict()
    assert result["status"] == "applied", result
    assert Path(result["backup_path"]).read_bytes() == before[0]
    persisted = load_project_state(database)
    assert [(item.device_id, item.display_text) for item in persisted.devices] == [(device.device_id, text)]
    assert persisted.placements == project.placements
    current_mapping = next(item for item in persisted.mappings if item.device_id == device.device_id)
    assert (current_mapping.mapping_id, current_mapping.source_range, current_mapping.sheet_name) == (
        mapping.mapping_id, mapping.source_range, mapping.sheet_name,
    )
    # Every other OOXML part, including chart, drawing, relations and styles, stays byte-identical.
    with ZipFile(source) as archive:
        assert set(archive.namelist()) == set(parts)
        assert all(archive.read(name) == data for name, data in parts.items()
                   if name != "xl/worksheets/sheet1.xml")
    workbook = load_workbook(source)
    try:
        sheet = workbook["Room"]
        assert sheet["B4"].value == text and sheet["B4"].data_type == "s"
        assert {str(merged) for merged in sheet.merged_cells.ranges} == {"A1:C1", "B4:B5"}
        assert sheet["B4"].font.bold and sheet["B4"].fill.fgColor.rgb == "0000FF00"
        assert sheet["B4"].alignment.wrap_text
        assert sheet.column_dimensions["B"].width == 24 and sheet.row_dimensions[4].height == 22
        assert sheet["Z1"].value == workbook["Other"]["Z1"].value == "Unrelated note"
    finally:
        workbook.close()


@pytest.mark.parametrize("room, units", [("Room", (8, 9)), ("Room", (9, 11)),
                                        ("Other", (9, 10)), ("Other", (1, 3))])
def test_geometry_changes_still_refuse_with_structured_evidence(tmp_path, capsys, room, units):
    source, database, project = _bound(tmp_path)
    before = source.read_bytes(), database.read_bytes()
    rack = next(rack for rack in project.racks if rack.source_sheet == room)
    plan = plan_device_edit(project, project.devices[0].device_id, rack.rack_id, *units,
                            display_text="renamed")
    assert any(item.code == "unsupported-workbook-object" for item in plan.conflicts)
    result = _preview_cli(tmp_path, database, _request(project, target={"rack_id": rack.rack_id},
                         start_u=units[0], end_u=units[1], text="renamed"), capsys, expected=2)
    assert result["status"] == "rejected"
    assert "unsupported-workbook-object" in result["message"]
    assert result["conflicts"] == [item.to_dict() for item in plan.conflicts]
    assert any("xl/charts/" in item for conflict in result["conflicts"] for item in conflict["evidence"])
    assert not (tmp_path / "plan.json").exists()
    assert not database.with_name(f".{database.name}.conversation-key").exists()
    assert (source.read_bytes(), database.read_bytes()) == before


@pytest.mark.parametrize("metadata", ["comment", "hyperlink", "validation", "defined-name"])
def test_pure_rename_keeps_action_range_guards_and_conflict_details(tmp_path, capsys, metadata):
    source, database, project = _bound(tmp_path, metadata=metadata)
    before = source.read_bytes(), database.read_bytes()
    expected = plan_device_rename(project, project.devices[0].device_id, _request(project)["text"])
    assert expected.conflicts and all(item.evidence for item in expected.conflicts)
    result = _preview_cli(tmp_path, database, _request(project), capsys, expected=2)
    assert result["conflicts"] == [item.to_dict() for item in expected.conflicts]
    assert not (tmp_path / "plan.json").exists()
    assert (source.read_bytes(), database.read_bytes()) == before


def test_unchanged_text_is_a_noop_even_with_unrelated_chart(tmp_path):
    source, database, project = _bound(tmp_path)
    before = source.read_bytes(), database.read_bytes()
    plan_path = tmp_path / "plan.json"
    preview = preview_conversation_edit(database, _request(project, text=project.devices[0].display_text), plan_path)
    assert preview["summary"]["write_plan"]["actions"] == []
    result = apply_conversation_edit(database, plan_path, preview["token"], commit=True)
    assert result["status"] == "applied" and result["backup_path"] is None
    assert (source.read_bytes(), database.read_bytes()) == before


def test_rename_database_failure_rolls_back_both_files(tmp_path, monkeypatch):
    from racktool.core import service

    source, database, project = _bound(tmp_path)
    before = source.read_bytes(), database.read_bytes()
    plan_path = tmp_path / "plan.json"
    preview = preview_conversation_edit(database, _request(project), plan_path)

    def fail_persistence(*args, **kwargs):
        raise OSError("synthetic persistence failure")

    monkeypatch.setattr(service, "save_project", fail_persistence)
    result = apply_conversation_edit(database, plan_path, preview["token"], commit=True)
    assert result["status"] == "failed"
    assert "synthetic persistence failure" in result["errors"]
    assert (source.read_bytes(), database.read_bytes()) == before


def test_old_generic_edit_preview_requires_new_review(tmp_path, monkeypatch):
    from racktool.core import conversation
    from racktool.core.sync import _plan_device_change

    source, database, project = _bound(tmp_path, chart_sheet=None)
    before = source.read_bytes(), database.read_bytes()
    plan_path = tmp_path / "old-plan.json"

    def legacy_edit(*args, **kwargs):
        return _plan_device_change(*args, allow_resize=True, **kwargs)

    # Sign a real preview from the previous generic planner, then simulate upgrading Core.
    with monkeypatch.context() as legacy:
        legacy.setattr(conversation, "plan_device_edit", legacy_edit)
        preview = preview_conversation_edit(database, _request(project), plan_path)
    assert preview["summary"]["write_plan"]["actions"][0]["operation"] == "edit"
    with pytest.raises(ValueError, match="重新预览"):
        apply_conversation_edit(database, plan_path, preview["token"], commit=True)
    assert (source.read_bytes(), database.read_bytes()) == before
