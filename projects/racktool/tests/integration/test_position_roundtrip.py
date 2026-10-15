"""Synthetic position-table acceptance; no fixture claims to be a real Golden."""

from __future__ import annotations

from dataclasses import replace
from pathlib import Path
from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill
from openpyxl.workbook.defined_name import DefinedName

from racktool.core import add_draft_device, create_draft_project, export_project_workbook
from racktool.core.identity import sha256_file
from racktool.core.position_table import (
    apply_position_table,
    preview_position_table,
    project_revision,
)
from racktool.core.service import import_project
from racktool.persistence import load_project, save_project


def _bound(tmp_path: Path):
    source = tmp_path / "working.xlsx"
    database = tmp_path / "project.sqlite"
    workbook = Workbook()
    workbook.remove(workbook.active)
    for room, name, ascending in (("Room-Z", "A01", False), ("Room-A", "B01", True)):
        sheet = workbook.create_sheet(room)
        sheet.merge_cells("C3:E3")
        sheet["C3"] = name
        sheet["C3"].fill = PatternFill("solid", fgColor="FF9900")
        units = range(1, 13) if ascending else range(12, 0, -1)
        for row, unit in enumerate(units, 4):
            sheet.cell(row, 3, unit)
            sheet.cell(row, 5, unit)
        low, high = (6, 7) if ascending else (9, 10)
        top = 4 + (low - 1 if ascending else 12 - high)
        sheet.merge_cells(start_row=top, end_row=top + 1, start_column=4, end_column=4)
        sheet.cell(top, 4, "Synthetic " + name)
        sheet.defined_names.add(DefinedName("LocalOriginal", attr_text=f"'{room}'!$C$3"))
    workbook.defined_names.add(DefinedName("GlobalOriginal", attr_text="'Room-Z'!$C$3"))
    workbook.save(source)
    workbook.close()
    return source, database, import_project(source, database)


def _draft(tmp_path: Path):
    database = tmp_path / "draft.sqlite"
    create_draft_project(database, count=1, height_u=12)
    add_draft_device(database, "A01", 9, 10, "Synthetic A01")
    project = add_draft_device(database, "B01", 6, 7, "Synthetic B01")
    return database, project


def _mutate(path: Path, edit):
    workbook = load_workbook(path)
    try:
        sheet = next(sheet for sheet in workbook.worksheets if "设备 ID" == sheet["H1"].value)
        edit(sheet)
        workbook.save(path)
    finally:
        workbook.close()


def _rename_move_resize(sheet):
    sheet["B2"] = "Renamed\nSynthetic"
    sheet["C2"] = "B01"
    sheet["D2"] = "1-3U"
    sheet["E2"] = 3
    sheet["F2"] = "Room-A"
    # Spreadsheet sorting must not change identities or bind baseline to row number.
    first = [sheet.cell(2, col).value for col in range(1, 10)]
    second = [sheet.cell(3, col).value for col in range(1, 10)]
    for col in range(1, 10):
        sheet.cell(2, col, second[col - 1])
        sheet.cell(3, col, first[col - 1])


@pytest.mark.parametrize("layout", ["source", "standard"])
def test_bound_round_trip_moves_across_rooms_resizes_and_renames(tmp_path: Path, layout: str):
    source, database, project = _bound(tmp_path)
    before_source = source.read_bytes()
    original = tmp_path / "original.xlsx"
    original.write_bytes(before_source)
    export = export_project_workbook(project, tmp_path / "positions.xlsx", layout=layout)
    if layout == "source":
        changed_parts = {"xl/workbook.xml", "xl/_rels/workbook.xml.rels", "xl/styles.xml", "[Content_Types].xml"}
        with ZipFile(source) as before, ZipFile(export) as after:
            assert all(before.read(name) == after.read(name) for name in before.namelist() if name not in changed_parts)
        workbook = load_workbook(export)
        assert workbook["Room-A"].defined_names["LocalOriginal"].attr_text == "'Room-A'!$C$3"
        assert workbook.defined_names["GlobalOriginal"].attr_text == "'Room-Z'!$C$3"
        workbook.close()
    _mutate(export, _rename_move_resize)
    preview = preview_position_table(database, export)
    assert preview.accepted, preview.to_dict()
    assert len(preview.changes) == 1
    assert source.read_bytes() == before_source
    result = apply_position_table(database, export, expected_preview_token=preview.preview_token)
    assert result.status == "applied", result.to_dict()
    assert result.backup_path and Path(result.backup_path).read_bytes() == before_source
    assert original.read_bytes() == before_source
    updated = load_project(database)
    assert {device.device_id for device in updated.devices} == {device.device_id for device in project.devices}
    device_id = preview.changes[0].device_id
    placement = next(p for p in updated.placements if p.device_id == device_id)
    target = next(r for r in updated.racks if r.rack_name == "B01")
    assert (placement.rack_id, placement.start_u, placement.end_u) == (target.rack_id, 1, 3)
    assert next(d for d in updated.devices if d.device_id == device_id).display_text == "Renamed\nSynthetic"
    assert next(m for m in updated.mappings if m.device_id == device_id).sheet_name == "Room-A"
    assert not preview_position_table(database, export).accepted  # revision consumed
    fresh = export_project_workbook(updated, tmp_path / "fresh.xlsx", layout=layout)
    assert preview_position_table(database, fresh).accepted


def test_draft_atomic_batch_round_trip_and_recovery_backup(tmp_path: Path):
    database, project = _draft(tmp_path)
    old_database = database.read_bytes()
    output = export_project_workbook(project, tmp_path / "positions.xlsx")

    def edit(sheet):
        sheet["B2"] = "New A"
        sheet["D2"] = "1-3U"
        sheet["E2"] = 3
        sheet["B3"] = "New B"
        sheet["D3"] = "4U"
        sheet["E3"] = 1

    _mutate(output, edit)
    plan = preview_position_table(database, output)
    assert plan.accepted and len(plan.changes) == 2
    assert database.read_bytes() == old_database
    result = apply_position_table(database, output, expected_preview_token=plan.preview_token)
    assert result.status == "applied"
    assert result.backup_path and Path(result.backup_path).read_bytes() == old_database
    assert {(p.start_u, p.end_u) for p in load_project(database).placements} == {(1, 3), (4, 4)}


def test_project_revision_ignores_retrieval_order(tmp_path: Path):
    database, project = _draft(tmp_path)
    loaded = load_project(database)
    reordered = replace(loaded, racks=list(reversed(loaded.racks)), devices=list(reversed(loaded.devices)),
                        placements=list(reversed(loaded.placements)), mappings=list(reversed(loaded.mappings)))
    assert project_revision(project) == project_revision(loaded) == project_revision(reordered)


@pytest.mark.parametrize("bound", [False, True])
def test_unchanged_export_and_preview_token_are_required(tmp_path: Path, bound: bool):
    if bound:
        source, database, project = _bound(tmp_path)
        source_hash = sha256_file(source)
    else:
        database, project = _draft(tmp_path)
    old = sha256_file(database)
    output = export_project_workbook(project, tmp_path / "positions.xlsx")
    plan = preview_position_table(database, output)
    assert plan.accepted and not plan.changes
    assert apply_position_table(database, output).status == "rejected"
    assert apply_position_table(database, output, expected_preview_token="wrong").status == "rejected"
    assert apply_position_table(database, output, expected_preview_token=plan.preview_token).status == "unchanged"
    assert sha256_file(database) == old
    if bound:
        assert sha256_file(source) == source_hash


_BAD_EDITS = [
    lambda s: setattr(s["B2"], "value", "=1+1"),
    lambda s: setattr(s["D2"], "value", "=1+1"),
    lambda s: setattr(s["A2"], "value", "=1+1"),
    lambda s: setattr(s["H2"], "value", "unknown-device"),
    lambda s: setattr(s["H2"], "value", s["H3"].value),
    lambda s: setattr(s["H2"], "value", None),
    lambda s: setattr(s["I2"], "value", "wrong-original-rack"),
    lambda s: setattr(s["G2"], "value", "已删除"),
    lambda s: setattr(s["E2"], "value", 99),
    lambda s: setattr(s["E2"], "value", True),
    lambda s: setattr(s["D2"], "value", "12-9U"),
    lambda s: setattr(s["D2"], "value", "0U"),
    lambda s: setattr(s["D2"], "value", "13U"),
    lambda s: setattr(s["D2"], "value", "9~10U"),
    lambda s: setattr(s["C2"], "value", "unknown-rack"),
    lambda s: setattr(s["F2"], "value", "unknown-room"),
    lambda s: setattr(s["B2"], "value", ""),
    lambda s: setattr(s["A1"], "value", "改标题"),
    lambda s: s.delete_rows(2),
    lambda s: s.defined_names.clear(),
]


@pytest.mark.parametrize("edit", _BAD_EDITS)
def test_invalid_table_never_partially_changes_draft(tmp_path: Path, edit):
    database, project = _draft(tmp_path)
    before = database.read_bytes()
    output = export_project_workbook(project, tmp_path / "positions.xlsx")
    _mutate(output, edit)
    plan = preview_position_table(database, output)
    assert not plan.accepted, plan.to_dict()
    result = apply_position_table(database, output, expected_preview_token=plan.preview_token)
    assert result.status == "rejected"
    assert database.read_bytes() == before


def test_conflicting_batch_and_stale_project_are_rejected_without_any_write(tmp_path: Path):
    database, project = _draft(tmp_path)
    output = export_project_workbook(project, tmp_path / "positions.xlsx")
    before = database.read_bytes()

    def overlap(sheet):
        sheet["B2"] = "Would otherwise rename"
        sheet["C2"] = "B01"
        sheet["D2"] = "6-7U"

    _mutate(output, overlap)
    plan = preview_position_table(database, output)
    assert not plan.accepted
    assert database.read_bytes() == before
    fresh = export_project_workbook(project, tmp_path / "fresh.xlsx")
    edited = replace(project, devices=[replace(d, display_text="external") for d in project.devices])
    save_project(database, edited)
    assert not preview_position_table(database, fresh).accepted


def test_changed_table_after_preview_rejects_old_token(tmp_path: Path):
    database, project = _draft(tmp_path)
    output = export_project_workbook(project, tmp_path / "positions.xlsx")
    _mutate(output, lambda s: setattr(s["B2"], "value", "First intent"))
    plan = preview_position_table(database, output)
    before = database.read_bytes()
    _mutate(output, lambda s: setattr(s["B2"], "value", "Changed intent"))
    result = apply_position_table(database, output, expected_preview_token=plan.preview_token)
    assert result.status == "rejected"
    assert database.read_bytes() == before


def test_metadata_baseline_is_checked_against_trusted_project(tmp_path: Path):
    from racktool.core.position_table import write_position_metadata

    database, project = _draft(tmp_path)
    output = export_project_workbook(project, tmp_path / "positions.xlsx")
    changed = replace(project, devices=[replace(d, display_text="forged baseline") for d in project.devices])

    def forge(sheet):
        # Keep the real project revision but replace a baseline field in hidden constants.
        import base64
        import json
        import zlib

        sheet.defined_names.clear()
        write_position_metadata(sheet, changed)
        names = sorted(sheet.defined_names)
        encoded = "".join(sheet.defined_names[n].attr_text[1:-1] for n in names)
        metadata = json.loads(zlib.decompress(base64.b64decode(encoded)))
        metadata["project_revision"] = project_revision(project)
        payload = base64.b64encode(zlib.compress(json.dumps(metadata).encode())).decode()
        sheet.defined_names.clear()
        for index, offset in enumerate(range(0, len(payload), 200)):
            sheet.defined_names.add(DefinedName(f"_RackToolPositions_v1_{index:05d}",
                                               attr_text='"' + payload[offset:offset + 200] + '"', hidden=True))

    _mutate(output, forge)
    plan = preview_position_table(database, output)
    assert not plan.accepted
    assert "基线" in plan.conflicts[0].message


@pytest.mark.parametrize("change", ["source-stale", "source-missing", "table-overlap"])
def test_bound_failures_leave_source_and_database_unchanged(tmp_path: Path, change: str):
    source, database, project = _bound(tmp_path)
    output = export_project_workbook(project, tmp_path / "positions.xlsx")
    if change == "source-stale":
        source.write_bytes(source.read_bytes() + b"external change")
    elif change == "source-missing":
        source.unlink()
    else:
        def overlap(sheet):
            sheet["C2"] = "B01"
            sheet["F2"] = "Room-A"
            sheet["D2"] = "6-7U"

        _mutate(output, overlap)
    old_database = database.read_bytes()
    old_source = source.read_bytes() if source.is_file() else None
    plan = preview_position_table(database, output)
    assert not plan.accepted
    assert apply_position_table(database, output, expected_preview_token=plan.preview_token).status == "rejected"
    assert database.read_bytes() == old_database
    assert (source.read_bytes() if source.is_file() else None) == old_source


def test_empty_draft_table_is_usable_without_fake_device_rows(tmp_path: Path):
    database = tmp_path / "draft.sqlite"
    project = create_draft_project(database, count=1, height_u=12)
    output = export_project_workbook(project, tmp_path / "positions.xlsx")
    plan = preview_position_table(database, output)
    assert plan.accepted and not plan.changes
    assert apply_position_table(database, output, expected_preview_token=plan.preview_token).status == "unchanged"


def test_source_export_handles_empty_self_closing_defined_names(tmp_path: Path):
    import re
    from io import BytesIO

    source, _database, _ = _bound(tmp_path)
    output = BytesIO()
    with ZipFile(source) as before, ZipFile(output, "w") as after:
        for name in before.namelist():
            payload = before.read(name)
            if name == "xl/workbook.xml":
                payload = re.sub(rb"<definedNames>.*?</definedNames>", b"<definedNames/>", payload)
            after.writestr(name, payload)
    source.write_bytes(output.getvalue())
    project = import_project(source, tmp_path / "empty-names.sqlite")
    export = export_project_workbook(project, tmp_path / "positions.xlsx", layout="source")
    assert preview_position_table(tmp_path / "empty-names.sqlite", export).accepted


def test_direct_source_export_of_valid_workbook_without_styles(tmp_path: Path):
    """A plain XLSX legally omits the entire styles part and its relationship."""
    import importlib.util
    from xml.etree import ElementTree as ET

    script = Path(__file__).resolve().parents[2] / "scripts" / "verify_skill_workflows.py"
    specification = importlib.util.spec_from_file_location("public_fixture", script)
    assert specification is not None and specification.loader is not None
    fixture = importlib.util.module_from_spec(specification)
    specification.loader.exec_module(fixture)
    source = tmp_path / "no-styles.xlsx"
    fixture.make_workbook(source)
    source_bytes = source.read_bytes()
    database = tmp_path / "no-styles.sqlite"
    project = import_project(source, database)
    exported = export_project_workbook(project, tmp_path / "exported.xlsx", layout="source")
    assert source.read_bytes() == source_bytes
    main = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    relationship_type = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles"
    with ZipFile(source) as before, ZipFile(exported) as after:
        assert "xl/styles.xml" not in before.namelist()
        for name in before.namelist():
            if name.startswith("xl/worksheets/"):
                assert before.read(name) == after.read(name)
        relations = ET.fromstring(after.read("xl/_rels/workbook.xml.rels"))
        styles = [item for item in relations if item.attrib.get("Type") == relationship_type]
        assert len(styles) == 1 and styles[0].attrib["Target"] == "/xl/styles.xml"
        content = ET.fromstring(after.read("[Content_Types].xml"))
        assert len([item for item in content if item.attrib.get("PartName") == "/xl/styles.xml"]) == 1
        indexes = ET.fromstring(after.read("xl/styles.xml")).find(f"{{{main}}}cellXfs")
        assert indexes is not None and len(indexes) > 1
    opened = load_workbook(exported)
    try:
        assert opened["Room-Z"]["C3"].value == "A01"
        assert opened["Room-A"]["C4"].value == 1
        assert opened["Room-Z"]["C3"].style_id == 0
        assert opened["设备位置表"]["A1"].font.bold
        assert opened["设备位置表"]["A1"].fill.fgColor.rgb == "00344054"
    finally:
        opened.close()
    assert preview_position_table(database, exported).accepted
