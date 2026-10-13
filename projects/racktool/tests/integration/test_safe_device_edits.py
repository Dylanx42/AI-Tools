from __future__ import annotations

import json
from dataclasses import replace
from pathlib import Path

import pytest
from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

from racktool.cli.main import main
from racktool.core import service, sync
from racktool.core.service import commit_write_plan, import_project
from racktool.core.sync import plan_device_add, plan_device_rename
from racktool.persistence import load_project
from racktool.profiles import validate_profile_data
from tests.layout_corpus import INVALID_TEXT, LayoutCase, build_layout, device_range


def _fixture(tmp_path: Path, metadata: str | None = None, operation: str = "add"):
    facts = build_layout(LayoutCase("paired-two", "descending", 12, "integer"), tmp_path)
    source = tmp_path / "source.xlsx"
    rack = facts["racks"][0]
    target = device_range(rack, 9, 10) if operation == "add" else facts["devices"][0]["source_range"]
    anchor = target.split(":")[0]
    workbook = load_workbook(source)
    sheet = workbook["Racks"]
    if metadata == "comment":
        sheet[anchor].comment = Comment("SYNTHETIC PROTECTED COMMENT", "fixture")
    elif metadata == "hyperlink":
        sheet[anchor].hyperlink = "https://example.com"
    elif metadata == "validation":
        validation = DataValidation(type="list", formula1='"A,B"')
        validation.add(target)
        sheet.add_data_validation(validation)
    elif metadata == "defined-name":
        workbook.defined_names.add(DefinedName("ProtectedTarget", attr_text=f"'Racks'!{target}"))
    elif metadata == "blank-merge":
        sheet.merge_cells(target)
    workbook.save(source)
    workbook.close()
    database = tmp_path / "project.sqlite"
    project = import_project(source, database)
    return source, database, project, facts


def _fixture_plan(project, facts, operation: str):
    if operation == "add":
        rack_id = next(rack.rack_id for rack in project.racks if rack.rack_name == facts["racks"][0]["name"])
        return plan_device_add(project, rack_id, 9, 10, "NEW")
    device_id = next(device.device_id for device in project.devices if device.display_text == facts["devices"][0]["text"])
    return plan_device_rename(project, device_id, "RENAMED")


@pytest.mark.parametrize("operation", ["add", "rename"])
@pytest.mark.parametrize("metadata,code", [
    ("comment", "unsupported-cell-comment"), ("hyperlink", "unsupported-cell-hyperlink"),
    ("validation", "unsupported-data-validation"), ("defined-name", "unsupported-defined-name"),
])
def test_new_edits_refuse_protected_cell_metadata_without_writes(
    tmp_path: Path, metadata: str, code: str, operation: str,
) -> None:
    source, database, project, _facts = _fixture(tmp_path, metadata, operation)
    source_before, database_before = source.read_bytes(), database.read_bytes()
    plan = _fixture_plan(project, _facts, operation)
    codes = {conflict.code for conflict in plan.conflicts}
    # openpyxl fills a blank linked cell with its URL; the occupancy guard fires first.
    assert code in codes or (metadata == "hyperlink" and operation == "add" and "target-u-occupied" in codes)
    assert commit_write_plan(source, database, plan).status == "rejected"
    assert source.read_bytes() == source_before and database.read_bytes() == database_before


def test_add_refuses_preexisting_blank_merge(tmp_path: Path) -> None:
    source, database, project, _facts = _fixture(tmp_path, "blank-merge")
    source_before, database_before = source.read_bytes(), database.read_bytes()
    plan = _fixture_plan(project, _facts, "add")
    assert "target-merge-conflict" in {conflict.code for conflict in plan.conflicts}
    assert commit_write_plan(source, database, plan).status == "rejected"
    assert source.read_bytes() == source_before and database.read_bytes() == database_before


@pytest.mark.parametrize("operation", ["add", "rename"])
@pytest.mark.parametrize("failure", ["reload", "database"])
def test_add_and_rename_failures_leave_workbook_and_database_intact(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, operation: str, failure: str,
) -> None:
    source, database, project, _facts = _fixture(tmp_path)
    source_before, database_before = source.read_bytes(), database.read_bytes()
    plan = _fixture_plan(project, _facts, operation)
    assert not plan.conflicts

    def fail(*_args: object, **_kwargs: object) -> None:
        raise ValueError("forced verification failure")

    monkeypatch.setattr(sync if failure == "reload" else service,
                        "_validate_written_workbook" if failure == "reload" else "save_project", fail)
    assert commit_write_plan(source, database, plan).status == "failed"
    assert source.read_bytes() == source_before and database.read_bytes() == database_before
    assert not list(tmp_path.glob(".source.tmp-*"))


@pytest.mark.parametrize("operation", ["add", "rename"])
def test_new_edits_reject_stale_or_forged_plans(tmp_path: Path, operation: str) -> None:
    source, database, project, _facts = _fixture(tmp_path)
    source_before, database_before = source.read_bytes(), database.read_bytes()
    plan = _fixture_plan(project, _facts, operation)
    forged = replace(plan, actions=(replace(plan.actions[0], new_range="Z100"),))
    assert commit_write_plan(source, database, forged).status == "rejected"
    assert source.read_bytes() == source_before and database.read_bytes() == database_before
    workbook = load_workbook(source)
    workbook["Preserved"]["A1"] = "EXTERNAL UPDATE"
    workbook.save(source)
    workbook.close()
    changed = source.read_bytes()
    assert commit_write_plan(source, database, plan).status == "rejected"
    assert source.read_bytes() == changed and database.read_bytes() == database_before


def test_cli_preview_commit_literal_text_and_identity_preservation(
    tmp_path: Path, capsys: pytest.CaptureFixture[str],
) -> None:
    source, database, project, _facts = _fixture(tmp_path)
    source_before, database_before = source.read_bytes(), database.read_bytes()
    arguments = ["sync", "add-device", str(source), str(database), project.racks[0].rack_id,
                 "9", "10", "--text", "=SYNTHETIC NEW\nMODEL"]
    assert main(arguments) == 0
    assert json.loads(capsys.readouterr().out)["actions"][0]["operation"] == "add"
    assert source.read_bytes() == source_before and database.read_bytes() == database_before
    assert main([*arguments, "--commit"]) == 0
    added = json.loads(capsys.readouterr().out)
    assert added["status"] == "applied"
    device_id = added["plan"]["actions"][0]["device_id"]
    assert main(["sync", "rename", str(source), str(database), device_id,
                 "--text", "=SYNTHETIC RENAMED\nMODEL", "--commit"]) == 0
    assert json.loads(capsys.readouterr().out)["status"] == "applied"
    assert any(device.device_id == device_id and device.display_text.startswith("=SYNTHETIC RENAMED")
               for device in load_project(database).devices)
    workbook = load_workbook(source)
    try:
        mapping = next(mapping for mapping in load_project(database).mappings if mapping.device_id == device_id)
        assert workbook[mapping.sheet_name][mapping.source_range.split(":")[0]].data_type == "s"
    finally:
        workbook.close()


def test_profile_ignored_source_cells_and_ignored_new_text_are_not_overwritten(tmp_path: Path) -> None:
    facts = build_layout(LayoutCase("paired-two", "descending", 12, "integer"), tmp_path)
    source = tmp_path / "source.xlsx"
    target = device_range(facts["racks"][0], 9, 10).split(":")[0]
    workbook = load_workbook(source)
    workbook["Racks"][target] = "RESERVED"
    workbook.save(source)
    workbook.close()
    profile = validate_profile_data({
        "schema_version": 1, "profile_id": "synthetic-ignore", "name": "Synthetic ignored labels",
        "match": {"sheet_name_regex": "^Racks$", "min_confidence": 1.0},
        "rack": {"title": {"mode": "merged_cell_above_u_axis"}, "height": {"mode": "infer_from_u_axis"}},
        "u_axis": {"direction": "descending", "pairing": "paired", "allowed_heights": [12]},
        "device_area": {"mode": "between_u_axes"}, "text": {"ignore_exact": ["RESERVED"]},
    })
    database = tmp_path / "project.sqlite"
    project = import_project(source, database, profile=profile)
    before_source, before_database = source.read_bytes(), database.read_bytes()
    rack_id = next(rack.rack_id for rack in project.racks if rack.rack_name == facts["racks"][0]["name"])
    plan = plan_device_add(project, rack_id, 9, 10, "NEW")
    assert "target-cell-not-empty" in {conflict.code for conflict in plan.conflicts}
    assert commit_write_plan(source, database, plan).status == "rejected"
    plan = plan_device_rename(project, project.devices[0].device_id, "reserved")
    assert "ignored-device-text" in {conflict.code for conflict in plan.conflicts}
    assert commit_write_plan(source, database, plan).status == "rejected"
    assert source.read_bytes() == before_source and database.read_bytes() == before_database


@pytest.mark.parametrize("text", ["", " \n", "A" * 32_768], ids=["empty", "whitespace", "over-limit"])
def test_invalid_device_text_is_rejected_before_writes(tmp_path: Path, text: str) -> None:
    source, database, project, facts = _fixture(tmp_path)
    original_source, original_database = source.read_bytes(), database.read_bytes()
    rack_id = next(rack.rack_id for rack in project.racks if rack.rack_name == facts["racks"][0]["name"])
    assert plan_device_add(project, rack_id, 9, 10, text).conflicts
    assert plan_device_rename(project, project.devices[0].device_id, text).conflicts
    assert source.read_bytes() == original_source and database.read_bytes() == original_database


def test_noop_rename_creates_no_workbook_or_database_changes(tmp_path: Path) -> None:
    source, database, project, _facts = _fixture(tmp_path)
    before_source, before_database = source.read_bytes(), database.read_bytes()
    plan = plan_device_rename(project, project.devices[0].device_id, project.devices[0].display_text)
    assert not plan.actions and not plan.conflicts
    assert commit_write_plan(source, database, plan).status == "applied"
    assert source.read_bytes() == before_source and database.read_bytes() == before_database


def test_huge_invalid_u_ranges_stop_before_occupancy_iteration(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    source, database, project, facts = _fixture(tmp_path)
    original_source, original_database = source.read_bytes(), database.read_bytes()
    rack_id = next(rack.rack_id for rack in project.racks if rack.rack_name == facts["racks"][0]["name"])

    def occupancy_must_not_run(*_args: object):
        raise AssertionError("Invalid U ranges must stop before occupancy iteration")

    monkeypatch.setattr(sync, "_occupied_units", occupancy_must_not_run)
    assert "u-out-of-bounds" in {conflict.code for conflict in plan_device_add(project, rack_id, 1, 10**12, "NEW").conflicts}
    device = next(device for device in project.devices if device.display_text == facts["devices"][0]["text"])
    assert "u-out-of-bounds" in {conflict.code for conflict in sync.plan_device_move(project, device.device_id, rack_id, 1, 10**12).conflicts}
    assert source.read_bytes() == original_source and database.read_bytes() == original_database


@pytest.mark.parametrize("text", INVALID_TEXT, ids=["nul", "vertical-tab", "xml-noncharacter", "lone-surrogate"])
def test_bound_workbook_edits_reject_invalid_xml_text_without_changes(tmp_path: Path, text: str) -> None:
    source, database, project, facts = _fixture(tmp_path)
    original_source, original_database = source.read_bytes(), database.read_bytes()
    rack_id = next(rack.rack_id for rack in project.racks if rack.rack_name == facts["racks"][0]["name"])
    for plan in (plan_device_add(project, rack_id, 9, 10, text),
                 plan_device_rename(project, project.devices[0].device_id, text)):
        assert "invalid-device-text" in {conflict.code for conflict in plan.conflicts}
        assert commit_write_plan(source, database, plan).status == "rejected"
    assert source.read_bytes() == original_source and database.read_bytes() == original_database


def test_forged_multi_action_plan_stops_before_expanding_huge_invalid_ranges(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch,
) -> None:
    source, database, project, facts = _fixture(tmp_path)
    before_source, before_database = source.read_bytes(), database.read_bytes()
    original = _fixture_plan(project, facts, "add")
    invalid = replace(original.actions[0], device_id="SYNTHETIC-FORGED-ID", start_u=1, end_u=10**12)
    forged = replace(original, actions=(*original.actions, invalid))
    original_u_range = sync._u_range

    def refuse_huge_iteration(start: int, end: int):
        if max(start, end) > 100:
            raise AssertionError("Unsafe plan must reject bounds before pairwise overlap iteration")
        return original_u_range(start, end)

    monkeypatch.setattr(sync, "_u_range", refuse_huge_iteration)
    assert commit_write_plan(source, database, forged).status == "rejected"
    assert source.read_bytes() == before_source and database.read_bytes() == before_database
