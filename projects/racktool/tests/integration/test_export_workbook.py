from __future__ import annotations

from dataclasses import replace
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from racktool.cli.main import main
from racktool.core.export import export_project_workbook
from racktool.gui.session import GuiSession


def _fill_descending_axis(
    sheet: Worksheet,
    column: int,
    start_row: int,
    height: int,
) -> None:
    for offset, u_number in enumerate(range(height, 0, -1)):
        sheet.cell(start_row + offset, column, u_number)


def _make_layout(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "一期机柜"
    sheet.merge_cells("A1:C1")
    sheet["A1"] = "RACK-A"
    _fill_descending_axis(sheet, 1, 2, 12)
    _fill_descending_axis(sheet, 3, 2, 12)
    sheet.merge_cells("B2:B3")
    sheet["B2"] = "核心交换机\nS5735-L48T4XE\n带外管理"
    sheet["B6"] = "防火墙"

    sheet.merge_cells("E1:G1")
    sheet["E1"] = "RACK-B"
    _fill_descending_axis(sheet, 5, 2, 12)
    _fill_descending_axis(sheet, 7, 2, 12)
    sheet["F4"] = "应用服务器"
    workbook.save(path)


def _make_multirow_layout(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "机柜布局"

    def add_rack(name: str, title_row: int, start_column: int) -> None:
        sheet.merge_cells(
            start_row=title_row,
            start_column=start_column,
            end_row=title_row,
            end_column=start_column + 2,
        )
        sheet.cell(title_row, start_column, name)
        _fill_descending_axis(sheet, start_column, title_row + 1, 6)
        _fill_descending_axis(sheet, start_column + 2, title_row + 1, 6)

    for index in range(6):
        add_rack(f"A{index + 1:02d}", 1, index * 4 + 1)
    for index in range(3):
        add_rack(f"B{index + 1:02d}", 10, index * 4 + 1)
    workbook.save(path)


def test_export_contains_rack_diagram_and_filterable_device_positions(
    tmp_path: Path,
) -> None:
    source = tmp_path / "layout.xlsx"
    _make_layout(source)
    source_bytes = source.read_bytes()
    session = GuiSession.open_workbook(source)

    output = session.export_xlsx(tmp_path / "racktool-export")

    assert output == tmp_path / "racktool-export.xlsx"
    assert output.is_file()
    assert source.read_bytes() == source_bytes
    workbook = load_workbook(output, read_only=False, data_only=False)
    try:
        assert workbook.sheetnames == ["机柜图", "设备位置表"]
        diagram = workbook["机柜图"]
        positions = workbook["设备位置表"]
        diagram_values = {
            cell.value
            for row in diagram.iter_rows()
            for cell in row
            if cell.value is not None
        }
        assert {"RACK-A", "RACK-B"}.issubset(diagram_values)
        assert "核心交换机\nS5735-L48T4XE\n带外管理" in diagram_values
        assert any(
            cell.value == "核心交换机\nS5735-L48T4XE\n带外管理"
            and cell.alignment.wrap_text
            for row in diagram.iter_rows()
            for cell in row
        )

        headers = [positions.cell(4, column).value for column in range(1, 12)]
        assert headers == [
            "序号",
            "设备",
            "机柜",
            "起始 U",
            "结束 U",
            "高度（U）",
            "工作表",
            "来源单元格",
            "状态",
            "设备 ID",
            "机柜 ID",
        ]
        assert positions.freeze_panes == "A5"
        assert positions.column_dimensions["J"].hidden
        assert positions.column_dimensions["K"].hidden
        assert "DevicePositionsTable" in positions.tables
        rows = list(positions.iter_rows(min_row=5, max_col=11, values_only=True))
        assert [row[1] for row in rows] == [
            "核心交换机\nS5735-L48T4XE\n带外管理",
            "防火墙",
            "应用服务器",
        ]
        assert rows[0][2:9] == ("RACK-A", 11, 12, 2, "一期机柜", "B2:B3", "在位")
        assert all(
            cell.data_type != "f"
            for sheet in workbook.worksheets
            for row in sheet.iter_rows()
            for cell in row
        )
    finally:
        workbook.close()


def test_export_preserves_source_rack_rows_instead_of_wrapping_every_four(
    tmp_path: Path,
) -> None:
    source = tmp_path / "multirow.xlsx"
    output = tmp_path / "multirow-export.xlsx"
    _make_multirow_layout(source)
    session = GuiSession.open_workbook(source)

    session.export_xlsx(output)

    workbook = load_workbook(output, read_only=False, data_only=False)
    try:
        diagram = workbook["机柜图"]
        rack_names = {
            *(f"A{index:02d}" for index in range(1, 7)),
            *(f"B{index:02d}" for index in range(1, 4)),
        }
        coordinates = {
            str(cell.value): (cell.row, cell.column)
            for row in diagram.iter_rows()
            for cell in row
            if cell.value in rack_names
        }
        top_rows = {coordinates[f"A{index:02d}"][0] for index in range(1, 7)}
        bottom_rows = {coordinates[f"B{index:02d}"][0] for index in range(1, 4)}

        assert len(top_rows) == 1
        assert len(bottom_rows) == 1
        assert top_rows != bottom_rows
        assert [coordinates[f"A{index:02d}"][1] for index in range(1, 7)] == [
            1,
            6,
            11,
            16,
            21,
            26,
        ]
        assert "每排机柜按源工作表中的行列位置排列" in str(diagram["A2"].value)
    finally:
        workbook.close()


def test_export_preserves_source_workbook_sheet_order(tmp_path: Path) -> None:
    source = tmp_path / "sheet-order.xlsx"
    output = tmp_path / "sheet-order-export.xlsx"
    workbook = Workbook()
    first = workbook.active
    assert first is not None
    first.title = "Z-先显示"
    first.merge_cells("A1:C1")
    first["A1"] = "RACK-Z"
    _fill_descending_axis(first, 1, 2, 6)
    _fill_descending_axis(first, 3, 2, 6)
    second = workbook.create_sheet("A-后显示")
    second.merge_cells("A1:C1")
    second["A1"] = "RACK-A"
    _fill_descending_axis(second, 1, 2, 6)
    _fill_descending_axis(second, 3, 2, 6)
    workbook.save(source)
    session = GuiSession.open_workbook(source)

    session.export_xlsx(output)

    exported = load_workbook(output, read_only=False, data_only=False)
    try:
        diagram = exported["机柜图"]
        group_rows = {
            str(cell.value): cell.row
            for row in diagram.iter_rows()
            for cell in row
            if str(cell.value).startswith("来源工作表：")
        }
        assert group_rows["来源工作表：Z-先显示"] < group_rows["来源工作表：A-后显示"]
    finally:
        exported.close()


def test_export_draws_complete_borders_around_merged_titles_and_devices(
    tmp_path: Path,
) -> None:
    source = tmp_path / "layout.xlsx"
    output = tmp_path / "bordered-export.xlsx"
    _make_layout(source)
    session = GuiSession.open_workbook(source)

    session.export_xlsx(output)

    workbook = load_workbook(output, read_only=False, data_only=False)
    try:
        diagram = workbook["机柜图"]

        def merged_range_for(value: str):
            return next(
                merged
                for merged in diagram.merged_cells.ranges
                if diagram.cell(merged.min_row, merged.min_col).value == value
            )

        def assert_complete_medium_outline(value: str) -> None:
            merged = merged_range_for(value)
            for column in range(merged.min_col, merged.max_col + 1):
                assert diagram.cell(merged.min_row, column).border.top.style == "medium"
                assert diagram.cell(merged.max_row, column).border.bottom.style == "medium"
            for row in range(merged.min_row, merged.max_row + 1):
                assert diagram.cell(row, merged.min_col).border.left.style == "medium"
                assert diagram.cell(row, merged.max_col).border.right.style == "medium"

        assert_complete_medium_outline("RACK-A")
        assert_complete_medium_outline("核心交换机\nS5735-L48T4XE\n带外管理")
    finally:
        workbook.close()


def test_export_refuses_source_overwrite_and_preserves_existing_target_on_failure(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    source = tmp_path / "layout.xlsx"
    _make_layout(source)
    session = GuiSession.open_workbook(source)
    source_bytes = source.read_bytes()

    with pytest.raises(ValueError, match="不能覆盖当前源工作簿"):
        export_project_workbook(session.project, source)
    assert source.read_bytes() == source_bytes

    target = tmp_path / "existing.xlsx"
    target.write_bytes(b"existing export stays intact")
    with pytest.raises(FileExistsError, match="文件已存在"):
        export_project_workbook(session.project, target)
    assert target.read_bytes() == b"existing export stays intact"

    def fail_validation(*_args: object, **_kwargs: object) -> None:
        raise ValueError("forced validation failure")

    monkeypatch.setattr("racktool.core.export._validate_export", fail_validation)
    with pytest.raises(ValueError, match="forced validation failure"):
        export_project_workbook(session.project, target, overwrite=True)

    assert target.read_bytes() == b"existing export stays intact"
    assert not list(tmp_path.glob(".existing.tmp-*.xlsx"))


def test_export_treats_formula_like_device_text_as_literal(tmp_path: Path) -> None:
    source = tmp_path / "layout.xlsx"
    output = tmp_path / "literal-text.xlsx"
    _make_layout(source)
    session = GuiSession.open_workbook(source)
    selected = next(device for device in session.project.devices if device.display_text == "防火墙")
    session.project = replace(
        session.project,
        devices=[
            replace(device, display_text="=2+2\n这只是设备文字")
            if device.device_id == selected.device_id
            else device
            for device in session.project.devices
        ],
    )

    export_project_workbook(session.project, output)

    workbook = load_workbook(output, read_only=False, data_only=False)
    try:
        matching = [
            cell
            for sheet in workbook.worksheets
            for row in sheet.iter_rows()
            for cell in row
            if cell.value == "=2+2\n这只是设备文字"
        ]
        assert matching
        assert all(cell.data_type == "s" for cell in matching)
    finally:
        workbook.close()


def test_export_excludes_unsynced_moves_and_says_so(tmp_path: Path) -> None:
    source = tmp_path / "layout.xlsx"
    output = tmp_path / "saved-state-only.xlsx"
    _make_layout(source)
    source_bytes = source.read_bytes()
    session = GuiSession.open_workbook(source)
    device = next(device for device in session.project.devices if device.display_text == "防火墙")
    target = next(rack for rack in session.project.racks if rack.rack_name == "RACK-B")

    plan = session.stage_move(device.device_id, target.rack_id, 1, 1)
    assert not plan.conflicts
    session.export_xlsx(output)

    workbook = load_workbook(output, read_only=False, data_only=False)
    try:
        positions = workbook["设备位置表"]
        row = next(
            values
            for values in positions.iter_rows(min_row=5, max_col=9, values_only=True)
            if values[1] == "防火墙"
        )
        assert row[2:5] == ("RACK-A", 8, 8)
    finally:
        workbook.close()
    assert "不含 1 项待同步更改" in session.status_message
    assert source.read_bytes() == source_bytes


def test_export_cli_uses_saved_project_state(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    source = tmp_path / "layout.xlsx"
    output = tmp_path / "cli-export.xlsx"
    _make_layout(source)
    session = GuiSession.open_workbook(source)

    assert main(["export", str(session.database_path), str(output)]) == 0

    assert output.is_file()
    assert '"status": "exported"' in capsys.readouterr().out
