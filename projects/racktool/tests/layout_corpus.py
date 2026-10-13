"""Synthetic layouts with construction facts independent of RackCore detection."""

from __future__ import annotations

import json
from dataclasses import asdict, dataclass
from itertools import product
from pathlib import Path
from typing import Any

import yaml
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

FAMILIES = ("paired-one", "paired-two", "shared-axis", "edge-axis", "two-bands", "fixed-profile")
HEIGHTS = (12, 42, 47, 53)
LABELS = ("integer", "ascii-u", "fullwidth-u")
INVALID_TEXT = ("\x00SYNTHETIC", "\x0BSYNTHETIC", "\ufffeSYNTHETIC", "\ud800SYNTHETIC")


@dataclass(frozen=True)
class LayoutCase:
    family: str
    direction: str
    height: int
    label: str

    @property
    def case_id(self) -> str:
        return f"{self.family}-{self.direction}-{self.height}-{self.label}"


LAYOUT_CASES = tuple(LayoutCase(*values) for values in product(
    FAMILIES, ("ascending", "descending"), HEIGHTS, LABELS,
))


@dataclass(frozen=True)
class DraftCase:
    count: int
    height: int
    density: int
    text_mode: str

    @property
    def case_id(self) -> str:
        return f"draft-{self.count}-{self.height}-{self.density}-{self.text_mode}"


DRAFT_CASES = tuple(DraftCase(*values) for values in product(
    (1, 2, 3, 6, 18), HEIGHTS, (0, 25, 50), ("multiline", "formula-text"),
))


def rectangle(columns: list[int], rows: list[int]) -> str:
    first = f"{get_column_letter(min(columns))}{min(rows)}"
    last = f"{get_column_letter(max(columns))}{max(rows)}"
    return first if first == last else f"{first}:{last}"


def device_range(rack: dict[str, Any], low: int, high: int) -> str:
    return rectangle(rack["device_columns"], [rack["u_to_row"][u] for u in range(low, high + 1)])


def build_layout(case: LayoutCase, directory: Path) -> dict[str, Any]:
    """Write fixture and known facts; never use the parser to manufacture expectations."""
    directory.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Racks"
    origin_column = 1 + LABELS.index(case.label) * 2
    title_row = 2 + HEIGHTS.index(case.height)
    sheet.cell(1, 20, "SYNTHETIC ONLY / " + case.case_id)
    racks: list[dict[str, Any]] = []
    devices: list[dict[str, Any]] = []
    edge = Side(style="thin", color="667788")
    border = Border(left=edge, right=edge, top=edge, bottom=edge)
    bands = 2 if case.family == "two-bands" else 1
    count = 1 if case.family == "fixed-profile" else 2
    width = 2 if case.family in {"paired-two", "fixed-profile"} else 1
    profile: dict[str, Any] | None = None

    for band in range(bands):
        top = title_row + band * (case.height + 5)
        body = top + (3 if case.family == "fixed-profile" else 1)
        for index in range(count):
            if case.family == "shared-axis":
                left = origin_column + index * 2
            else:
                left = origin_column + index * (width + 4)
            columns = list(range(left + 1, left + width + 1))
            right = None if case.family == "edge-axis" else columns[-1] + 1
            title_left = columns[0] if case.family == "fixed-profile" else left
            title_right = columns[-1] if case.family in {"shared-axis", "edge-axis", "fixed-profile"} else right
            assert title_right is not None
            sheet.merge_cells(start_row=top, start_column=title_left, end_row=top, end_column=title_right)
            name = f"{chr(65 + band)}{index + 1:02d}"
            title = sheet.cell(top, title_left, name)
            title.fill = PatternFill("solid", fgColor="F6C88F")
            title.font = Font(bold=True)
            u_to_row = {
                u: body + (u - 1 if case.direction == "ascending" else case.height - u)
                for u in range(1, case.height + 1)
            }
            rack = {
                "name": name, "height": case.height, "sheet": "Racks",
                "left": left, "right": right, "device_columns": columns,
                "title_range": rectangle([title_left, title_right], [top]), "u_to_row": u_to_row,
            }
            racks.append(rack)
            for u, row in u_to_row.items():
                value: int | str = u if case.label == "integer" else (
                    f"{u}U" if case.label == "ascii-u" else f"{u:02d}Ｕ"
                )
                for column in (left, right):
                    if column is not None:
                        axis = sheet.cell(row, column, value)
                        axis.fill = PatternFill("solid", fgColor="FCE0B8")
                        axis.border = border
                for column in columns:
                    cell = sheet.cell(row, column)
                    cell.border = border
                    cell.alignment = Alignment(wrap_text=True, vertical="center")
                sheet.row_dimensions[row].height = 20 + LABELS.index(case.label)
            for column in columns:
                sheet.column_dimensions[get_column_letter(column)].width = 24 / width
            for low, high, suffix in ((1, 2, "2U"), (4, 4, "1U")):
                text = f"SYNTH-{name}-{suffix}\n型号 / 管理口"
                if case.label == "fullwidth-u" and low == 4:
                    text = "=" + text
                source_range = device_range(rack, low, high)
                if ":" in source_range:
                    sheet.merge_cells(source_range)
                rows = [u_to_row[u] for u in range(low, high + 1)]
                anchor = sheet.cell(min(rows), columns[0], text)
                anchor.data_type = "s"
                anchor.font = Font(name="Arial", size=10, bold=True)
                anchor.fill = PatternFill("solid", fgColor="D9EAF7")
                devices.append({
                    "text": text, "rack": name, "low": low, "high": high, "source_range": source_range,
                })
            if case.family == "fixed-profile":
                profile = {
                    "schema_version": 1, "profile_id": "synthetic-fixed-layout",
                    "name": "Synthetic fixed-coordinate layout",
                    "match": {"sheet_name_regex": "^Racks$", "min_confidence": 1.0},
                    "rack": {"title": {"mode": "fixed_range", "range": rack["title_range"]},
                             "height": {"mode": "infer_from_u_axis"}},
                    "u_axis": {
                        "direction": case.direction, "pairing": "paired", "max_missing_rows": 0,
                        "allowed_heights": [case.height], "left_column": get_column_letter(left),
                        "right_column": get_column_letter(right), "start_row": body,
                        "end_row": body + case.height - 1,
                    },
                    "device_area": {"mode": "fixed_range", "range": rectangle(columns, [body, body + case.height - 1])},
                }

    notes = workbook.create_sheet("Preserved")
    notes.sheet_state = "hidden"
    notes["A1"] = "SYNTHETIC UNRELATED TEXT"
    notes["A2"].comment = Comment("SYNTHETIC COMMENT", "fixture")
    notes["A3"] = "SYNTHETIC LINK"
    notes["A3"].hyperlink = "https://example.com"
    notes["D3"] = "=1+2"
    workbook.defined_names.add(DefinedName("SyntheticKeep", attr_text="'Preserved'!$A$1"))
    source = directory / "source.xlsx"
    workbook.save(source)
    workbook.close()
    facts = {"case": asdict(case), "synthetic": True, "racks": racks, "devices": devices}
    (directory / "expected.json").write_text(json.dumps(facts, ensure_ascii=False, indent=2), encoding="utf-8")
    if profile is not None:
        (directory / "profile.yaml").write_text(yaml.safe_dump(profile, allow_unicode=True, sort_keys=False), encoding="utf-8")
    return facts
