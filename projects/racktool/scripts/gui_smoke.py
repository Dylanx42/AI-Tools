"""Exercise Qt bindings and CLI behavior using disposable synthetic data."""

from __future__ import annotations

import hashlib
import json
import os
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import Workbook

from racktool.gui.session import GuiSession
from racktool.gui.window import CockpitWindow, create_app


def main() -> None:
    os.environ["QT_QPA_PLATFORM"] = "offscreen"
    with tempfile.TemporaryDirectory(prefix="racktool-gui-") as directory:
        root = Path(directory)
        os.environ["XDG_CACHE_HOME"] = directory
        os.environ["XDG_RUNTIME_DIR"] = directory
        path = root / "synthetic.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        assert sheet is not None
        sheet.title = "Synthetic 12U"
        sheet.merge_cells("A1:C1")
        sheet["A1"] = "SYNTHETIC-RACK"
        for row, u in enumerate(range(12, 0, -1), start=2):
            sheet.cell(row, 1, u)
            sheet.cell(row, 3, u)
        sheet["B2"] = "SYNTHETIC-A"
        sheet["B5"] = "SYNTHETIC-B"
        workbook.save(path)
        original_hash = hashlib.sha256(path.read_bytes()).hexdigest()
        for command in ("inspect", "analyze"):
            arguments = [sys.executable, "-m", "racktool", command, str(path)]
            first = subprocess.check_output(arguments)
            assert first == subprocess.check_output(arguments), f"Non-deterministic {command}"
            assert isinstance(json.loads(first), dict)
        app = create_app([])
        cockpit = CockpitWindow()
        session = GuiSession.open_workbook(path)
        cockpit.load_session(session)
        cockpit.widget().show()
        app.processEvents()
        assert cockpit.device_table.rowCount() == 2
        assert cockpit.rack_table.rowCount() == 1
        assert cockpit.mapping_table.rowCount() > 0
        cockpit.rack_table.selectRow(0)
        cockpit.device_table.selectRow(0)
        app.processEvents()
        assert cockpit.occupancy_table.rowCount() == 12
        assert cockpit.rack_box.currentData() == session.project.racks[0].rack_id
        assert cockpit.start_u.value() in (9, 12)
        cockpit.widget().close()
        app.processEvents()
        assert hashlib.sha256(path.read_bytes()).hexdigest() == original_hash
    print("PASS: deterministic CLI and Qt workbook/table/selection bindings on synthetic data")
    print("This offscreen smoke is not a Golden or desktop/Excel/WPS acceptance result.")


if __name__ == "__main__":
    main()
